# -*- coding: utf-8 -*-
"""
後加工工賃 3ソース金額整合性検証スクリプト

① 東レ送付CSV (RVSHEETyyyymm.csv)      : 東レ検収データ。発注No.=契約NO、金額列を集計
     入庫場所 A010=国分工場 / A010P=湖南工場 (工場ごとに絞り込む)
     H列が「-」の行は取消行 → 金額の符号を反転して集計 (T行小計と同じ扱い)
② 工場ごとにフォーマットが異なる:
   国分: 国分工場/長岡後加工賃明細/後加工工賃明細*.xlsx 「東レまとめ」シート
         依頼NO = A列 & "-" & B列 / C列 = 契約NO / AA列 = 加工賃合計
   湖南: \\192.168.0.101\共有フォルダ\湖南工場\湖南共有\002  加工G\000  後加工業務\2 後加工試算 (KONAN_SHISAN_ROOT)
           直下の「月度加工賃試算.xlsm」= 当月作業中ファイル / 「yyyy年度試算　湖南 / m月度加工賃試算.xlsm」= 確定済み各月(年度=暦年)
           シート内「yyyy年m月度」で月を判定し、同一年月が複数あれば更新日時最新を採用
         「東レT.V.C」「東レY.S」「東レW.E..」3シート
         A列 = 加工依頼No.(依頼NO) / C列 = 契約No. / 「加工種類別 加工賃」合計列(V列) = 加工賃
③ アラジン 依頼NO別問合せ*.xlsx
   国分: 国分工場/アラジン (●自動検証内)
   湖南: \\192.168.0.101\共有フォルダ\湖南工場\湖南共有\生産管理システム\管理システム\●DATA\月次実績
         (生産管理システムの月次実績出力フォルダを直接参照。KONAN_ALADDIN_DIR)
     項目「加工金額」の「--合計--」列 = 依頼NO別加工金額
     シート内「対象年月」で月を判定し、同一年月に複数ファイルがあれば _yyyymmdd_hhmmss 最新を採用
     湖南は東レ以外の得意先が混在するため 得意先=049006(東ﾚ自材部) のみ集計
手動判定 ●自動検証/手動判定.csv (工場,対象月,契約NO,正とする側,理由)
     当月・当工場の検証Aの不一致に「①正(手動)」「②正(手動)」を付けて要確認から外す(対象月は必須。他月には効かない)
前月過不足 ●自動検証/前月過不足.csv (工場,対象月,契約NO,依頼NO,①東レ金額,②長岡金額,理由)
     前月の工賃差検証で次月修正依頼済みだが当月①に未反映の差額を検証A「前月過不足」として報告計上に含める
     (検証A: ①列空・②列=プラス差。マイナス側は①取消行で前月調整に出るがプラス側は①に現れない)
(湖南のみ) 月次処理ファイル (KONAN_MONTHLY_ROOT/yyyy年/yyyy年mm月度 湖南工場 月次処理ファイル.xlsm)
     契約NO等の明細が無いため「月次検証」「集計表」シートの東レ合計だけを総額照合 (検証C)
(国分のみ) ②「東レまとめ」の内部整合性 (検証D)
     まとめは 東レT/東レV.C/東レY/東レW.E の各行を参照し G/Z/AA を自行で再計算する構造。
     参照式のずれ・元シートのデータ行の取込漏れ・G/AA列の式崩れ(空欄/直接入力)・行別AA金額差 を検出

検証:
  A … ① vs ② 契約NO突合 (前月調整 / 翌月記載 の自動判定を含む)
  B … ② vs ③ 依頼NO突合 (複数月累計での月ずれ解消・枝番統合を含む)
  C … (湖南) 月次処理ファイル東レ合計 vs 本検証総額
  D … (国分) ②まとめ vs 元4シート

突合キー:
  ①↔② … 契約NO   (①は「191-352R」形式、②は「191352R」形式 → ハイフン除去して比較)
  ②↔③ … 依頼NO   (「C7-52」形式。全角英数字は半角へ正規化して比較)

使い方:
  python .cursor/skills/kouchin-consistency-check/scripts/verify_kouchin.py                        # 国分工場 (各フォルダの対象月ファイルを自動検出)
  python .cursor/skills/kouchin-consistency-check/scripts/verify_kouchin.py --factory konan        # 湖南工場
  python .cursor/skills/kouchin-consistency-check/scripts/verify_kouchin.py --toray <csv> --nagaoka <xlsx|xlsm> --aladdin <xlsx>
  python .cursor/skills/kouchin-consistency-check/scripts/verify_kouchin.py --tol 0.5              # 許容差(円)を変更
"""
import argparse
import csv
import io
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

try:
    from python_calamine import CalamineWorkbook
except ImportError:
    sys.exit('python-calamine が必要です:  pip install python-calamine')

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.styles.fonts import DEFAULT_FONT
    from openpyxl.utils import get_column_letter
    from openpyxl.writer.theme import theme_xml
except ImportError:
    sys.exit('openpyxl が必要です:  pip install openpyxl')

# ---- フォント: BIZ UDPゴシックを全面適用 ----
# DEFAULT_FONT書き換えで無指定セル(原本コピー含む)全部に適用される。scheme=None必須
# (minorのままだとテーマ側のMS Pゴシックが優先される)
FONT_NAME = 'BIZ UDPゴシック'
FONT_NUM = 'BIZ UDゴシック'    # 等幅版: 数字の字幅が揃うため金額列に使用
FONT_SIZE = 10.5
DEFAULT_FONT.name = FONT_NAME
DEFAULT_FONT.size = FONT_SIZE
DEFAULT_FONT.scheme = None
# テーマの既定フォントも差し替え (手入力時の既定・name無しフォントレコードの保険)
BIZ_THEME = theme_xml.replace('<a:latin typeface="Calibri"/>', f'<a:latin typeface="{FONT_NAME}"/>')
BIZ_THEME = BIZ_THEME.replace('<a:latin typeface="Cambria"/>', f'<a:latin typeface="{FONT_NAME}"/>')
BIZ_THEME = re.sub(r'(<a:font script="Jpan" typeface=")[^"]*(")', rf'\g<1>{FONT_NAME}\g<2>', BIZ_THEME)


def F(**kw):
    """BIZ UDPゴシック固定のFontファクトリ"""
    kw.setdefault('name', FONT_NAME)
    kw.setdefault('size', FONT_SIZE)
    return Font(**kw)

# スクリプト位置 (.cursor/skills/<skill>/scripts/) から ●自動検証 フォルダを解決
DEFAULT_BASE = Path(__file__).resolve().parents[4]
# 湖南②③は●自動検証の外 (湖南共有) を直接参照するため絶対パス
# ②: 後加工試算ルート。直下の「月度加工賃試算.xlsm」が当月作業中ファイル、「yyyy年度試算　湖南\m月度加工賃試算.xlsm」が確定済み各月
KONAN_SHISAN_ROOT = r'\\192.168.0.101\共有フォルダ\湖南工場\湖南共有\002  加工G\000  後加工業務\2 後加工試算'
# ③: 生産管理システム(アラジン)の月次実績出力フォルダ
KONAN_ALADDIN_DIR = r'\\192.168.0.101\共有フォルダ\湖南工場\湖南共有\生産管理システム\管理システム\●DATA\月次実績'
# 月次処理ファイル(検証C): 年フォルダ「yyyy年」配下に「yyyy年mm月度 湖南工場 月次処理ファイル.xlsm」
KONAN_MONTHLY_ROOT = r'\\192.168.0.101\共有フォルダ\湖南工場\湖南共有\生産管理システム\アラジンオフィスシステムデータ\月次実績表\0 東レ月次処理ファイル'


def resolve_dir(base: Path, d) -> Path:
    """プロファイルのフォルダ指定を解決する。絶対パス(UNC含む)はそのまま、相対なら base 配下"""
    p = Path(d)
    return p if p.is_absolute() else base / p

KEIYAKU_CSV_PAT = re.compile(r'^\d{3}-\d{3}[A-Z0-9]$')   # ①の発注No. 例: 191-352R
KEIYAKU_PAT = re.compile(r'^[0-9A-Z]{7}$')               # 契約NO = ハイフン無し7桁英数字 例: 191352R


def norm(v):
    """全角→半角(NFKC)正規化、空白除去(内部空白含む)、大文字化。数値は整数文字列化。"""
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(v))).upper()


def latest(folder: Path, pattern: str) -> Path:
    files = [f for f in folder.glob(pattern) if not f.name.startswith('~$')]
    if not files:
        sys.exit(f'ファイルが見つかりません: {folder}\\{pattern}')
    return max(files, key=lambda f: f.stat().st_mtime)


GATSUDO_PAT = re.compile(r'(\d{4})\s*年\s*(\d{1,2})\s*月度')
RVSHEET_PAT = re.compile(r'RVSHEET(\d{4})(\d{2})', re.IGNORECASE)


def find_toray_csv(folder: Path) -> Path:
    """①をファイル名 RVSHEETyyyymm の年月最大で選択 (更新日時より堅牢)"""
    files = [(m, f) for f in folder.glob('RVSHEET*.csv') if not f.name.startswith('~$')
             and (m := RVSHEET_PAT.search(f.name))]
    if not files:
        sys.exit(f'RVSHEETyyyymm.csv が見つかりません: {folder}')
    return max(files, key=lambda t: (int(t[0].group(1)), int(t[0].group(2))))[1]


def target_ym_of_csv(path: Path):
    """①ファイル名から対象年月 (yyyy, mm) を取得"""
    m = RVSHEET_PAT.search(path.name)
    return (int(m.group(1)), int(m.group(2))) if m else None


def find_nagaoka_files(folder: Path, target_ym):
    """②のファイル名「yyyy年m月度」から当月・過去月・翌月以降のファイル一覧を特定する。

    当月 = target_ym(①CSVの対象月)と同じ年月のファイル。翌月ファイルが先行作成されていても
    誤選択しない。同一年月に複数版(V01/V02等)がある場合は更新日時が最新のものを採用。
    戻り値: (当月ファイル,
             [(ラベル, 過去月ファイル), ...] 新しい月順,
             [(ラベル, 翌月以降ファイル), ...] 近い月順)
    """
    dated = defaultdict(list)
    for f in folder.glob('後加工工賃明細*.xlsx'):
        if f.name.startswith('~$'):
            continue
        ym = ym_from_filename(f)
        if ym:
            dated[ym].append(f)
    if not dated:
        sys.exit(f'「yyyy年m月度」を含む後加工工賃明細*.xlsx が見つかりません: {folder}')
    return _pick_by_ym(dated, target_ym, folder)


def ym_from_filename(path: Path):
    """ファイル名の「yyyy年m月度」から (yyyy, m) を返す。無ければ None"""
    m = GATSUDO_PAT.search(norm(path.name))
    return (int(m.group(1)), int(m.group(2))) if m else None


PREV_MONTH_WINDOW = 6   # 過去月明細として読む範囲(当月から何か月前まで)。年を含まない依頼NO(C8-1等)の前年分との衝突を防ぐため12未満にする


def _pick_by_ym(dated, target_ym, folder):
    """年月→ファイル群 の辞書から 当月/過去月/翌月以降 を振り分ける (同一年月は更新日時最新を採用)。
    過去月は直近 PREV_MONTH_WINDOW か月に限る (前月調整の判定と検証Bの累計比較に使う範囲)"""
    if target_ym is None:
        cur_ym = max(dated)
    elif target_ym in dated:
        cur_ym = target_ym
    else:
        sys.exit(f'②に {target_ym[0]}年{target_ym[1]}月度 のファイルがありません: {folder}\n'
                 f'  存在する月度: {", ".join(f"{y}年{m}月度" for y, m in sorted(dated))}\n'
                 f'  別ファイルを使う場合は --nagaoka で明示指定してください')
    newest = lambda ym: max(dated[ym], key=lambda f: f.stat().st_mtime)
    midx = lambda ym: ym[0] * 12 + ym[1]
    current = newest(cur_ym)
    prevs = [(f'{y}年{m}月度', newest((y, m))) for (y, m) in sorted(dated, reverse=True)
             if 0 < midx(cur_ym) - midx((y, m)) <= PREV_MONTH_WINDOW]
    nexts = [(f'{y}年{m}月度', newest((y, m))) for (y, m) in sorted(dated) if (y, m) > cur_ym]
    return current, prevs, nexts


# ---- 湖南工場 ② 加工賃試算xlsm ----
SHISAN_SHEETS = ('東レT.V.C', '東レY.S', '東レW.E..')


def shisan_ym(path: Path):
    """湖南②(m月度加工賃試算.xlsm)の対象年月を東レシート先頭の「yyyy年m月度」から取得 (ファイル名に年が無い)"""
    wb = CalamineWorkbook.from_path(str(path))
    for name in SHISAN_SHEETS:
        if name not in wb.sheet_names:
            continue
        for r in wb.get_sheet_by_name(name).to_python()[:3]:
            for v in r:
                if isinstance(v, str) and (m := GATSUDO_PAT.search(norm(v))):
                    return (int(m.group(1)), int(m.group(2)))
    return None


_NENDO_DIR_PAT = re.compile(r'^(\d{4})年度試算')   # 「2026年度試算　湖南」(暦年。H30年度等の旧フォルダは対象外)


def _shisan_candidates(root: Path, target_ym):
    """湖南②の候補ファイルを集める。
    ・root 直下の「月度加工賃試算.xlsm」= 当月の作業中ファイル (月末処理後に年度フォルダへ m月度 として格納される)
    ・root/「yyyy年度試算　湖南」/「m月度加工賃試算.xlsm」= 確定済みの各月 (年度=暦年)。
      対象年の前後1年のフォルダも見る (1月の過去月 / 12月の翌月が別フォルダになるため)。
    対象年が不明なら年の新しい2フォルダを見る。temp・バックアップ等の他フォルダは見ない"""
    def xls(folder):
        return [f for f in folder.glob('*加工賃試算*.xls*') if not f.name.startswith('~$')]
    files = xls(root)
    year_dirs = {}
    for d in root.iterdir():
        if d.is_dir() and (m := _NENDO_DIR_PAT.match(norm(d.name))):
            year_dirs[int(m.group(1))] = d
    years = {target_ym[0] - 1, target_ym[0], target_ym[0] + 1} if target_ym else set(sorted(year_dirs)[-2:])
    for y in sorted(years):
        if y in year_dirs:
            files += xls(year_dirs[y])
    return files


def find_shisan_files(folder: Path, target_ym):
    """湖南②「m月度加工賃試算.xlsm」を、シート内の年月で 当月/過去月/翌月以降 に振り分ける。
    folder = 後加工試算ルート (作業中ファイル + 年度フォルダを走査)。同一年月が複数あれば更新日時最新を採用"""
    dated = defaultdict(list)
    for f in _shisan_candidates(folder, target_ym):
        ym = shisan_ym(f)
        if ym:
            dated[ym].append(f)
    if not dated:
        sys.exit(f'東レシートに「yyyy年m月度」を持つ *加工賃試算*.xlsm が見つかりません: {folder} '
                 f'(直下および「yyyy年度試算　湖南」フォルダ)')
    return _pick_by_ym(dated, target_ym, folder)


def read_shisan(path: Path):
    """湖南② 加工賃試算xlsmの東レ3シート(東レT.V.C/東レY.S/東レW.E..)を合算して読む。
    A列=加工依頼No.(依頼NO) / C列=契約No. / 「加工種類別 加工賃」の合計列(V列)=加工賃。
    ヘッダー行の3行下が合計行、4行下からデータ行。合計行とデータ行合計を検算する。
    戻り値は read_nagaoka と同じ: (依頼NO->金額, 契約NO->金額, 契約NO->依頼NO集合, 形式不正リスト, 警告)"""
    wb = CalamineWorkbook.from_path(str(path))
    by_irai, by_keiyaku, keiyaku_to_irai = defaultdict(float), defaultdict(float), defaultdict(set)
    bad, warnings = [], []
    found = 0
    for name in SHISAN_SHEETS:
        if name not in wb.sheet_names:
            warnings.append(f'②{path.name}: シート「{name}」が無いためスキップしました')
            continue
        found += 1
        rows = wb.get_sheet_by_name(name).to_python()
        hdr = next((i for i, r in enumerate(rows[:20]) if r and '加工依頼' in norm(r[0])), None)
        if hdr is None or len(rows[hdr]) < 3 or '契約' not in norm(rows[hdr][2]):
            sys.exit(f'②「{name}」のヘッダー(A列「加工依頼No.」/C列「契約No.」)が見つかりません: {path}')
        amt_col = next((j for j, v in enumerate(rows[hdr]) if '加工賃' in norm(v) and '加工種類別' in norm(v)), None)
        if amt_col is None:
            sys.exit(f'②「{name}」に「加工種類別 加工賃」列が見つかりません: {path}')
        total_row = rows[hdr + 3] if len(rows) > hdr + 3 else []
        sheet_total = total_row[amt_col] if len(total_row) > amt_col and isinstance(total_row[amt_col], (int, float)) else None
        sheet_sum = 0.0
        for i, r in enumerate(rows[hdr + 4:], start=hdr + 5):
            irai = norm(r[0]) if len(r) > 0 else ''
            c = norm(r[2]).replace('-', '') if len(r) > 2 else ''
            amt = r[amt_col] if len(r) > amt_col else None
            if not irai and not c:
                continue
            if not isinstance(amt, (int, float)):
                if norm(amt):
                    warnings.append(f'②{path.name}「{name}」{i}行目: 依頼NO {irai} 契約NO {c} の加工賃「{amt}」が数値でないため除外')
                continue
            sheet_sum += amt
            if irai:
                by_irai[irai] += amt
            if KEIYAKU_PAT.match(c):
                by_keiyaku[c] += amt
                if irai:
                    keiyaku_to_irai[c].add(irai)
            elif c and c != '0':
                bad.append((f'「{name}」{i}', c, amt))
            elif not c and abs(amt) > 0.5:
                bad.append((f'「{name}」{i}', '(契約NO空欄)', amt))
        if sheet_total is not None and abs(sheet_total - sheet_sum) > 0.5:
            warnings.append(f'②{path.name}「{name}」: 合計行({sheet_total:,.0f}円)とデータ行合計({sheet_sum:,.0f}円)が'
                            f' {sheet_sum - sheet_total:+,.0f}円 不一致。シート内の集計式を確認')
    if not found:
        sys.exit(f'②に東レシート({"/".join(SHISAN_SHEETS)})が1つもありません: {path}')
    if not by_irai and not by_keiyaku:
        sys.exit(f'②の東レシートからデータ行を1件も取り込めませんでした: {path}')
    return dict(by_irai), dict(by_keiyaku), dict(keiyaku_to_irai), bad, warnings


def find_monthly_file(folder: Path, target_ym):
    """湖南 月次処理ファイル (yyyy年mm月度 湖南工場 月次処理ファイル.xlsm) を対象月で選択。無ければ None。
    folder 直下と、年フォルダ「yyyy年」(対象年) の両方を見る。同一月度が複数あれば更新日時最新"""
    if not folder.is_dir():
        return None
    folders = [folder]
    if target_ym:
        folders.append(folder / f'{target_ym[0]}年')
    else:
        folders += sorted((d for d in folder.iterdir() if d.is_dir() and re.fullmatch(r'\d{4}年', norm(d.name))), reverse=True)[:1]
    cands = [f for d in folders if d.is_dir() for f in d.glob('*月次処理ファイル*.xls*') if not f.name.startswith('~$')
             and (target_ym is None or ym_from_filename(f) == target_ym)]
    return max(cands, key=lambda f: f.stat().st_mtime) if cands else None


def read_monthly_file(path: Path):
    """湖南 月次処理ファイルの東レ合計値を読む (契約NO等の明細が無いため合計のみ照合)。
    「月次検証」シート東レ行: 月次実績①(Excel) / 月次実績表②(アラジン) / 売上明細③
    「集計表」シート東レ合計行: 加工金額
    戻り値: {項目名: 金額 or None}"""
    wb = CalamineWorkbook.from_path(str(path))
    res = {}
    if '月次検証' in wb.sheet_names:
        rows = wb.get_sheet_by_name('月次検証').to_python()
        hdr = next((i for i, r in enumerate(rows[:5]) if any('月次実績' in norm(v) for v in r)), None)
        toray = next((r for r in rows if r and norm(r[0]) == '東レ'), None)
        if hdr is not None and toray:
            h = [norm(v) for v in rows[hdr]]

            def col(key):
                j = next((j for j, v in enumerate(h) if norm(key) in v), None)
                v = toray[j] if j is not None and len(toray) > j else None
                return v if isinstance(v, (int, float)) else None
            res['月次検証: 月次実績①(Excel)'] = col('月次実績①')
            res['月次検証: 月次実績表②(アラジン)'] = col('月次実績表②')
            res['月次検証: 売上明細③'] = col('売上明細③')
    if '集計表' in wb.sheet_names:
        rows = wb.get_sheet_by_name('集計表').to_python()
        hdr = next((i for i, r in enumerate(rows[:5]) if any(norm(v) == '加工金額' for v in r)), None)
        row = next((r for r in rows if r and norm(r[0]) == '東レ合計'), None)
        if hdr is not None and row:
            j = next(j for j, v in enumerate(rows[hdr]) if norm(v) == '加工金額')
            v = row[j] if len(row) > j else None
            res['集計表: 東レ合計 加工金額'] = v if isinstance(v, (int, float)) else None
    return res


def _hl2_nagaoka(vals, att_irai, att_keiyaku):
    """国分②原本(東レまとめ)の行が要確認キーに該当するか"""
    if len(vals) < 3:
        return False
    irai = f'{norm(vals[0])}-{norm(vals[1])}'
    return irai in att_irai or norm(vals[2]).replace('-', '') in att_keiyaku


def _hl2_shisan(vals, att_irai, att_keiyaku):
    """湖南②原本(加工賃試算 東レシート)の行が要確認キーに該当するか"""
    if len(vals) < 3:
        return False
    return norm(vals[0]) in att_irai or norm(vals[2]).replace('-', '') in att_keiyaku


def _raw2_nagaoka(path):
    return [('東レまとめ', CalamineWorkbook.from_path(str(path)).get_sheet_by_name('東レまとめ').to_python())]


def _raw2_shisan(path):
    wb = CalamineWorkbook.from_path(str(path))
    return [(name, wb.get_sheet_by_name(name).to_python()) for name in SHISAN_SHEETS if name in wb.sheet_names]


# ---- 検証D (国分のみ): ②「東レまとめ」の内部整合性 ----
MATOME_SHEET = '東レまとめ'
MATOME_SRC_SHEETS = ('東レT', '東レV.C', '東レY', '東レW.E')
MATOME_DATA_START = 6
MATOME_TOTAL_LABELS = ('加工量合計', '加工賃合計', '営業入庫量')   # 元シート末尾の合計行(C列ラベル)
_MATOME_REF_PAT = re.compile(r"^=\+?'?(東レ[^'!]+)'?!\$?([A-Z]{1,2})\$?(\d+)$")
_CELL_ROW_PAT = re.compile(r'(?<![A-Za-z!$])\$?[A-Z]{1,2}\$?(\d+)')


def _matome_expected(col, r):
    """元シート・まとめ共通の定型式 (G=単価合計, Z/AA=加工賃, AB..AR=種類別金額, AT..BJ=種類別数量)"""
    if col == 'G':
        return f'=SUM(H{r}:X{r})'
    if col == 'Z':
        return f'=ROUNDUP($D{r}*G{r},0)'
    if col == 'AA':
        return f'=ROUNDUP($E{r}*G{r},0)'
    return None


def check_matome(path: Path, tol):
    """国分② 「東レまとめ」シートの内部整合性を検証する (検証D)。

    まとめは元4シート(東レT/V.C/Y/W.E)の各行を =東レT!A6 のように列ごとに参照し、
    G(単価合計)・Z・AA(加工賃) は自行で再計算する構造。元シート側で G を直接入力で
    上書きしたり式を消したりすると、まとめの再計算値と元シートの値が食い違う。
    (2026年8月度からは G も =東レT!G7 の参照式に変更されており、この形も定型として受け付ける)
      1. まとめの参照式: 行内でシート・行・列がそろっているか / 自行式(G/Z/AA)が他行を参照していないか
      2. 元シートのデータ行(依頼NO or 契約NOあり)がすべてまとめに取り込まれているか
      3. 元シート・まとめの G/AA が定型式か (空欄・直接入力・別式を検出)
      4. 行別に まとめAA と 元シートAA を突合し、シート別合計も比較する
    戻り値: dict(rows=[(場所, 依頼NO, 契約NO, まとめAA, 元AA, 差, 判定, 内容)],
                 totals=[(元シート, まとめ計, 元シート計, 差)], n_map=まとめ取込行数, skipped=None|理由)
    判定: 金額差 / 参照ずれ / 未取込 (要修正) , 式異常 (空欄・別式) , 直接入力 (値で上書き・注意)
    """
    try:
        fwb = load_workbook(str(path), data_only=False)
    except Exception as e:  # noqa: BLE001 - スタイル破損等で読めないファイルは検証Dだけスキップ
        return {'rows': [], 'totals': [], 'n_map': 0, 'skipped': f'openpyxlで数式を読めませんでした: {e}'}
    if MATOME_SHEET not in fwb.sheetnames or not all(s in fwb.sheetnames for s in MATOME_SRC_SHEETS):
        return {'rows': [], 'totals': [], 'n_map': 0,
                'skipped': f'シート構成が想定外です(必要: {MATOME_SHEET}, {", ".join(MATOME_SRC_SHEETS)})'}
    cwb = CalamineWorkbook.from_path(str(path))
    vals = {s: cwb.get_sheet_by_name(s).to_python() for s in (MATOME_SHEET,) + MATOME_SRC_SHEETS}

    def val(sheet, r, c):
        rows = vals[sheet]
        if 0 < r <= len(rows) and 0 < c <= len(rows[r - 1]):
            return rows[r - 1][c - 1]
        return None

    def num(v):
        return v if isinstance(v, (int, float)) else 0

    def ident(sheet, r):
        """元シート行の (依頼NO, 契約NO)。両方空ならデータ行ではない
        (A列に接頭辞「T7」だけ入った空テンプレート行や、C列が合計ラベルの行は除く)"""
        a, b, c = val(sheet, r, 1), val(sheet, r, 2), val(sheet, r, 3)
        irai = f'{norm(a)}-{norm(b)}' if norm(a) and norm(b) else ''
        keiyaku = norm(c) if norm(c) not in MATOME_TOTAL_LABELS else ''
        return irai, keiyaku

    ws = fwb[MATOME_SHEET]
    src_cols = [get_column_letter(c) for c in range(1, 25) if get_column_letter(c) != 'G']   # A-F, H-X
    rows_out = []
    mapping = {}
    g_ref_style = False   # まとめの G/Z/AA が元シート参照式 (=東レT!G7) になっているか
    for r in range(MATOME_DATA_START, ws.max_row + 1):
        refs = {}
        for col in src_cols:
            f = ws[f'{col}{r}'].value
            if isinstance(f, str) and (m := _MATOME_REF_PAT.match(f)):
                refs[col] = (m.group(1), m.group(2), int(m.group(3)))
        if not refs:
            continue
        sheets = Counter(s for s, _, _ in refs.values())
        src_rows = Counter(n for _, _, n in refs.values())
        s, n = sheets.most_common(1)[0][0], src_rows.most_common(1)[0][0]
        mapping[r] = (s, n)
        irai, keiyaku = ident(s, n)
        problems = []
        if len(sheets) > 1:
            problems.append('参照シートが混在: ' + ', '.join(f'{c}→{sh}' for c, (sh, _, _) in refs.items() if sh != s))
        if len(src_rows) > 1:
            problems.append('参照行がずれ: ' + ', '.join(f'{c}→{sh}!{sc}{sr}' for c, (sh, sc, sr) in refs.items() if sr != n))
        colmis = [f'{c}→{sc}' for c, (_, sc, _) in refs.items() if sc != c]
        if colmis:
            problems.append('参照列が自列と不一致: ' + ', '.join(colmis))
        if problems:
            rows_out.append((f'{MATOME_SHEET}!{r}', irai, keiyaku, '', '', '', '参照ずれ', ' / '.join(problems)))
        for col in ('G', 'Z', 'AA'):
            f = ws[f'{col}{r}'].value
            exp = _matome_expected(col, r)
            if isinstance(f, str) and (m := _MATOME_REF_PAT.match(f)):
                # 「=東レT!G7」のように元シートの同じ列・同じ行をそのまま参照する形も正 (2026年8月度から採用。
                # 元シートのGを直接入力しても まとめ と食い違わない)。別シート・別行・別列なら参照ずれ
                g_ref_style = True
                if (m.group(1), m.group(2), int(m.group(3))) != (s, col, n):
                    rows_out.append((f'{MATOME_SHEET}!{col}{r}', irai, keiyaku, '', '', '', '参照ずれ',
                                     f'参照先が自行の元シート行({s}!{col}{n})と不一致: {f}'))
            elif isinstance(f, str) and f.startswith('='):
                other = {int(x) for x in _CELL_ROW_PAT.findall(f)} - {r}
                if other:
                    rows_out.append((f'{MATOME_SHEET}!{col}{r}', irai, keiyaku, '', '', '', '参照ずれ',
                                     f'自行式が他行を参照: {f} (期待 {exp})'))
                elif f.replace(' ', '') != exp.replace(' ', ''):
                    rows_out.append((f'{MATOME_SHEET}!{col}{r}', irai, keiyaku, '', '', '', '式異常',
                                     f'定型と異なる式: {f} (期待 {exp})'))
            elif f not in (None, '') and (irai or keiyaku):
                rows_out.append((f'{MATOME_SHEET}!{col}{r}', irai, keiyaku, '', '', '', '直接入力',
                                 f'式ではなく値 {f!r} が入力されている (期待 {exp})。元シートの単価欄と食い違う恐れ'))

    if not mapping:
        # A〜X が値貼り付けのファイル(6月度など)は元シートとの対応が取れない
        return {'rows': [], 'totals': [], 'n_map': 0,
                'skipped': f'「{MATOME_SHEET}」に元シートへの参照式が無い(値貼り付け)ため元シートとの突合ができません'}
    covered = defaultdict(set)
    for r, (s, n) in mapping.items():
        covered[s].add(n)
    totals = []
    for s in MATOME_SRC_SHEETS:
        sws = fwb[s]
        rows = vals[s]
        sum_m = sum_s = 0.0
        data_rows = []
        for i in range(MATOME_DATA_START, len(rows) + 1):
            irai, keiyaku = ident(s, i)
            if not (irai or keiyaku):
                continue
            data_rows.append(i)
            # 3. 元シートの式チェック (G/AA)
            for col in ('G', 'AA'):
                f = sws[f'{col}{i}'].value
                exp = _matome_expected(col, i)
                if f in (None, ''):
                    rows_out.append((f'{s}!{col}{i}', irai, keiyaku, '', num(val(s, i, 27)), '', '式異常',
                                     f'{col}列が空欄 (期待 {exp})。加工賃が0になる'))
                elif isinstance(f, str) and f.startswith('='):
                    if f.replace(' ', '') != exp.replace(' ', ''):
                        rows_out.append((f'{s}!{col}{i}', irai, keiyaku, '', num(val(s, i, 27)), '', '式異常',
                                         f'定型と異なる式: {f} (期待 {exp})'))
                else:
                    why = ('まとめはこの値をそのまま参照するため金額は一致するが、単価欄(H〜X)の合計と食い違う' if g_ref_style
                           else 'まとめは単価欄(H〜X)の合計で再計算するため食い違う恐れ')
                    rows_out.append((f'{s}!{col}{i}', irai, keiyaku, '', num(val(s, i, 27)), '', '直接入力',
                                     f'式ではなく値 {f!r} が入力されている (期待 {exp})。{why}'))
        for i in data_rows:
            if i not in covered[s]:
                irai, keiyaku = ident(s, i)
                aa = num(val(s, i, 27))
                rows_out.append((f'{s}!{i}', irai, keiyaku, '', aa, '', '未取込',
                                 f'元シートのデータ行がまとめに参照されていない (AA {aa:,.0f}円)'))
        # 4. 行別金額突合 (元シートの合計行・空行を参照しているまとめ行は対象外)
        data_set = set(data_rows)
        for r, (ms, n) in sorted(mapping.items()):
            if ms != s or n not in data_set:
                continue
            m_aa, s_aa = num(val(MATOME_SHEET, r, 27)), num(val(s, n, 27))
            sum_m += m_aa
            sum_s += s_aa
            if abs(m_aa - s_aa) > tol:
                irai, keiyaku = ident(s, n)
                g_m, g_s = num(val(MATOME_SHEET, r, 7)), num(val(s, n, 7))
                rows_out.append((f'{MATOME_SHEET}!{r} ↔ {s}!{n}', irai, keiyaku, m_aa, s_aa, m_aa - s_aa, '金額差',
                                 f'まとめAA {m_aa:,.0f} ≠ 元シートAA {s_aa:,.0f} (単価合計G: まとめ {round(g_m, 4):g} / 元 {round(g_s, 4):g})。'
                                 '元シートの単価を丸めるか、まとめの直接入力を式に戻して両者を一致させる'))
        totals.append((s, sum_m, sum_s, sum_m - sum_s))
    return {'rows': rows_out, 'totals': totals, 'n_map': len(mapping), 'skipped': None}


def _factories():
    """工場プロファイル: ②のフォーマット・フォルダ・③の得意先絞り込みが工場ごとに異なる"""
    return {
    'kokubu': {
        'label': '国分工場', 'other': '湖南工場', 'basho': 'A010',
        'dir2': '国分工場/長岡後加工賃明細', 'dir3': '国分工場/アラジン', 'dir_monthly': None,
        'name2': '長岡明細', 'desc2': '東レまとめシート', 'amt2': 'AA',
        'find2': find_nagaoka_files, 'read2': read_nagaoka, 'ym2': ym_from_filename,
        'raw2': _raw2_nagaoka, 'freeze2': 'A6', 'hl2': _hl2_nagaoka,
        'check_d': check_matome,   # ②「東レまとめ」の内部整合性 (元4シートとの参照・式・金額突合)
        'customer3': None,
        'box2': ('② 長岡明細', '東レまとめ・AA列(加工賃)'),
        'key_a': '契約NO = ①発注No.のハイフン除去(例 191-352R→191352R) = ②「東レまとめ」C列。ハイフン無し7桁英数字',
        'key_b': '依頼NO = ②A列&"-"&B列(例 C7-52) = ③依頼NO列。全角英数字は半角に直して比較',
        'take2': '② =「東レまとめ」AA列(同一キー複数行は合算)',
    },
    'konan': {
        'label': '湖南工場', 'other': '国分工場', 'basho': 'A010P',
        # ②は湖南共有の後加工試算ルートを直接参照 (find_shisan_files が直下の作業中ファイルと年度フォルダを走査)
        'dir2': KONAN_SHISAN_ROOT, 'dir_monthly': KONAN_MONTHLY_ROOT,
        # ③は湖南の生産管理システムが月次実績を書き出すフォルダを直接参照する (●自動検証の外・絶対パス)。
        # 同一対象年月のファイルが複数あればファイル名 _yyyymmdd_hhmmss が最新のものを採用 (find_aladdin_all)
        'dir3': KONAN_ALADDIN_DIR,
        'name2': '加工賃試算', 'desc2': '東レT.V.C/東レY.S/東レW.E..シート', 'amt2': '加工賃',
        'find2': find_shisan_files, 'read2': read_shisan, 'ym2': shisan_ym,
        'raw2': _raw2_shisan, 'freeze2': 'A12', 'hl2': _hl2_shisan,
        'check_d': None,         # 湖南の試算xlsmはまとめシート構造ではないため対象外
        'customer3': '049006',   # 東ﾚ(株) 自材部。TPI(049052)等が混在するため絞り込む
        'box2': ('② 加工賃試算', '東レ3シート・加工賃合計(V列)'),
        'key_a': '契約NO = ①発注No.のハイフン除去 = ②加工賃試算の東レ3シート C列「契約No.」。ハイフン無し7桁英数字',
        'key_b': '依頼NO = ②A列「加工依頼No.」(例 C8-1) = ③依頼NO列(得意先 049006 東ﾚ自材部の行のみ)',
        'take2': '② = 東レT.V.C/東レY.S/東レW.E..シートの「加工種類別 加工賃」合計列(V列)を3シート合算',
    },
    }


def aladdin_taisho_ym(path: Path):
    """③のシート内「対象年月 : yyyy年mm月」を (yyyy, mm) で返す。見つからなければ None"""
    rows = CalamineWorkbook.from_path(str(path)).get_sheet_by_index(0).to_python()
    for r in rows[:10]:
        joined = ''.join(str(v) for v in r)
        if '対象年月' in joined:
            m = re.search(r'(\d{4})\s*年\s*(\d{1,2})\s*月', joined)
            if m:
                return (int(m.group(1)), int(m.group(2)))
    return None


def find_aladdin_all(folder: Path):
    """③の全照会ファイルを対象年月(シート内)ごとに整理する。
    同一年月に複数あればファイル名 _yyyymmdd_hhmmss が最新のものを採用。
    戻り値: {(y, m): Path}"""
    files = [f for f in folder.glob('依頼NO別問合せ*.xlsx') if not f.name.startswith('~$')]
    if not files:
        sys.exit(f'依頼NO別問合せ*.xlsx が見つかりません: {folder}')

    def key(f):
        m = re.search(r'_(\d{8})_(\d{6})', f.name)
        return (m.group(1) + m.group(2)) if m else f'0{f.stat().st_mtime:015.0f}'
    result = {}
    for f in sorted(files, key=key):  # 古い順に上書き → 各年月の最新が残る
        ym = aladdin_taisho_ym(f)
        if ym:
            result[ym] = f
    if not result:
        sys.exit(f'③のシート内「対象年月」を読めるファイルがありません: {folder}')
    return result


def find_aladdin(folder: Path, target_ym) -> Path:
    """③を選択する。シート内の対象年月が target_ym(①の対象月)と一致するファイル。
    他月の照会ファイルが混在しても誤選択しない。"""
    amap = find_aladdin_all(folder)
    if target_ym is None:
        return amap[max(amap)]
    if target_ym in amap:
        return amap[target_ym]
    listing = ', '.join(f'{f.name}({y}年{m}月)' for (y, m), f in sorted(amap.items()))
    sys.exit(f'③に対象年月 {target_ym[0]}年{target_ym[1]}月 の依頼NO別問合せファイルがありません: {folder}\n'
             f'  存在するファイル: {listing}\n'
             f'  別ファイルを使う場合は --aladdin で明示指定してください')


MANUAL_FILE = '手動判定.csv'
MANUAL_HEADER = ('工場', '対象月', '契約NO', '正とする側', '理由')
PRIOR_FILE = '前月過不足.csv'
ARCHIVE_DIR = '過去検証結果'  # 最新以外の 検証結果_*.xlsx / .html を退避するフォルダ

PRIOR_HEADER = ('工場', '対象月', '契約NO', '依頼NO', '①東レ金額', '②長岡金額', '理由')


def _parse_ym(text):
    """「2026年8月度」「2026-08」「2026/8」「202608」を (y, m) に。読めなければ None"""
    t = norm(text)
    if (m := GATSUDO_PAT.search(t)) or (m := re.fullmatch(r'(\d{4})[-/.](\d{1,2})', t)) or (m := re.fullmatch(r'(\d{4})(\d{2})', t)):
        return (int(m.group(1)), int(m.group(2)))
    return None


def _read_csv_rows(path: Path):
    """UTF-8(BOM可) → cp932 の順でCSVを読む。空行除去済みの行リストを返す。"""
    raw = path.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('cp932')
    return [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]


def _csv_factory_match(fac_cell, factory_label):
    if not fac_cell:
        return True
    fl = norm(factory_label)
    return fac_cell in (fl, fl.replace('工場', ''))


def _parse_amount_cell(text):
    """金額セルを float に。空欄は None、読めなければ ValueError。"""
    t = (text or '').strip().replace(',', '').replace('円', '')
    if not t:
        return None
    return float(t)


def load_manual_judgments(path: Path, factory_label, target_ym):
    """●自動検証\\手動判定.csv (列: 工場,対象月,契約NO,正とする側,理由) から、対象工場・対象月の行だけを読む。
    「今回に限り①東レ金額を正とする」のような人の判断を、検証Aの不一致に対して月単位で適用するための仕組み。
    正とする側: ①/1/東レ → 1 (②を①金額に修正・東レへは報告しない) / ②/2/長岡 → 2 (東レへ報告)。
    他の月・他工場の行は無視されるので、翌月以降には自動的に効かなくなる。
    UTF-8(BOM可)で読み、失敗したら cp932 (Excelの「CSV」保存)。
    戻り値: ({契約NO: (side, 理由)}, 警告リスト)"""
    result, warnings = {}, []
    if not path.is_file():
        return result, warnings
    rows = _read_csv_rows(path)
    if not rows:
        return result, warnings
    hdr = [norm(c) for c in rows[0]]
    if not {'契約NO', '正とする側'} <= set(hdr):
        warnings.append(f'{path.name}: ヘッダーが想定と異なるため無視しました (必要: {",".join(MANUAL_HEADER)})')
        return result, warnings
    col = {h: i for i, h in enumerate(hdr)}

    def cell(r, name):
        i = col.get(name)
        return r[i].strip() if i is not None and i < len(r) else ''
    for n, r in enumerate(rows[1:], start=2):
        if not _csv_factory_match(norm(cell(r, '工場')), factory_label):
            continue
        ym = _parse_ym(cell(r, '対象月'))
        if ym is None:
            # 対象月は必須 (「今回に限り」の判定を翌月以降に持ち越さないため)
            warnings.append(f'{path.name} {n}行目: 対象月「{cell(r, "対象月")}」を読めないため無視しました (例: 2026年8月度)')
            continue
        if target_ym and ym != target_ym:
            continue
        k = norm(cell(r, '契約NO')).replace('-', '')
        side_txt = norm(cell(r, '正とする側'))   # NFKC で ①→1, ②→2
        side = 1 if side_txt in ('1', '東レ', 'トウレ') else 2 if side_txt in ('2', '長岡', 'ナガオカ') else None
        if not k or side is None:
            warnings.append(f'{path.name} {n}行目: 契約NO「{cell(r, "契約NO")}」正とする側「{cell(r, "正とする側")}」が不正のため無視しました (正とする側は ① か ②)')
            continue
        result[k] = (side, cell(r, '理由'))
    return result, warnings


def load_prior_shortfalls(path: Path, factory_label, target_ym):
    """●自動検証\\前月過不足.csv (列: 工場,対象月,契約NO,依頼NO,①東レ金額,②長岡金額,理由) を読む。

    前月の工賃差検証で東レへ次月修正を依頼した差額のうち、当月①CSVにまだ調整行が来ていないものを
    検証Aの「前月過不足」として載せ、報告する過不足(当月)に計上するための入力。
    対象月=当月(調整を反映すべき月)。①②金額は前月当時の金額で、差額=①-② を報告計上額にする。
    戻り値: ([{契約NO,依頼NO,①,②,差額,理由}], 警告リスト)"""
    result, warnings = [], []
    if not path.is_file():
        return result, warnings
    rows = _read_csv_rows(path)
    if not rows:
        return result, warnings
    hdr = [norm(c) for c in rows[0]]
    # ①東レ金額/②長岡金額 の別名も許容 (NFKC正規化で ①→1, ②→2 になる点に注意)
    aliases = {
        '①東レ金額': '①東レ金額', '1東レ金額': '①東レ金額', '①金額': '①東レ金額', '1金額': '①東レ金額',
        '東レ金額': '①東レ金額', '①': '①東レ金額', '1': '①東レ金額',
        '②長岡金額': '②長岡金額', '2長岡金額': '②長岡金額', '②金額': '②長岡金額', '2金額': '②長岡金額',
        '長岡金額': '②長岡金額', '②': '②長岡金額', '2': '②長岡金額',
    }
    hdr_map = {}
    for i, h in enumerate(hdr):
        hdr_map[aliases.get(h, h)] = i
    if not {'契約NO', '①東レ金額', '②長岡金額'} <= set(hdr_map):
        warnings.append(f'{path.name}: ヘッダーが想定と異なるため無視しました (必要: {",".join(PRIOR_HEADER)})')
        return result, warnings

    def cell(r, name):
        i = hdr_map.get(name)
        return r[i].strip() if i is not None and i < len(r) else ''

    for n, r in enumerate(rows[1:], start=2):
        if not _csv_factory_match(norm(cell(r, '工場')), factory_label):
            continue
        ym = _parse_ym(cell(r, '対象月'))
        if ym is None:
            warnings.append(f'{path.name} {n}行目: 対象月「{cell(r, "対象月")}」を読めないため無視しました (例: 2026年8月度)')
            continue
        if target_ym and ym != target_ym:
            continue
        k = norm(cell(r, '契約NO')).replace('-', '')
        irai = norm(cell(r, '依頼NO'))
        try:
            a1 = _parse_amount_cell(cell(r, '①東レ金額'))
            a2 = _parse_amount_cell(cell(r, '②長岡金額'))
        except ValueError:
            warnings.append(f'{path.name} {n}行目: 金額を数値化できないため無視しました')
            continue
        if not k or a1 is None or a2 is None:
            warnings.append(f'{path.name} {n}行目: 契約NO・①東レ金額・②長岡金額は必須です')
            continue
        result.append({'契約NO': k, '依頼NO': irai, '①': a1, '②': a2, '差額': a1 - a2, '理由': cell(r, '理由')})
    return result, warnings


def read_toray_csv(path: Path, basho='A010'):
    """① 契約NO(ハイフン除去) -> 金額合計。小計行(発注No.=T/GT)や飾り行は形式不一致で除外される。

    A列(入庫場所)はブロック先頭のデータ行にのみ入り、以降の行は空欄で引き継がれる。
    A010=国分工場 / A010P=湖南工場 を basho 引数で絞り込む。
    H列(8列目)が「-」の行は取消(マイナス)行で、金額列は絶対値で入っているため符号を反転して集計する
    (T行の小計はこの符号込みで計算されている)。
    T行(ブロック小計)とデータ行合計を検算し、東レ側が小計から除外した行を検出する。
    戻り値: (契約NO->金額合計, 入庫場所->金額合計, 契約NO->入庫月日集合, 警告リスト, 小計不一致リスト,
             契約NO->取消行リスト[(行番号, 入庫月日, 反転後金額)])
    小計不一致リストの要素: (開始行, T行, ブロック内入庫場所, T小計, データ行合計, 差, 差と同額の契約NO候補[(契約NO, 入庫場所, 金額)])
    """
    try:
        text = path.read_bytes().decode('cp932')
    except UnicodeDecodeError as e:
        sys.exit(f'①CSVがcp932でデコードできません: {path}\n{e}')
    rows = list(csv.reader(io.StringIO(text)))
    # ヘッダー検証: 金額列(10列目)の位置が想定どおりか
    if not any(len(r) > 9 and '金額' in r[9] for r in rows[:3]):
        sys.exit(f'①CSVの10列目に「金額」ヘッダーが見つかりません。フォーマット変更の可能性: {path}')
    result = defaultdict(float)
    by_basho = defaultdict(float)
    nyuko_dates = defaultdict(set)   # 契約NO -> 入庫月日(yymmdd)
    minus_rows = defaultdict(list)   # 契約NO -> [(行番号, 入庫月日, 反転後金額)] 対象入庫場所の取消行
    warnings = []
    ct_errors = []
    current_basho = None
    block_sum = 0.0
    block_start = None
    block_rows = []                  # (契約NO, 入庫場所, 金額) 小計ブロック内のデータ行
    for idx, cols in enumerate(rows, start=1):
        if len(cols) < 12:
            continue
        keiyaku = cols[2].strip()
        if KEIYAKU_CSV_PAT.match(keiyaku):
            a = cols[0].strip()
            if a:
                current_basho = norm(a)
            try:
                amount = float(cols[9].strip().replace(',', ''))
            except ValueError:
                warnings.append(f'①CSV {idx}行目: 発注No.{keiyaku} の金額「{cols[9]}」を数値化できず除外')
                continue
            sign_cell = cols[7].strip()
            if sign_cell == '-':
                amount = -amount
            elif sign_cell:
                warnings.append(f'①CSV {idx}行目: 発注No.{keiyaku} のH列(符号)に想定外の値「{sign_cell}」。プラスとして集計')
            basho_key = current_basho or '(不明)'
            by_basho[basho_key] += amount
            block_sum += amount
            if block_start is None:
                block_start = idx
            k = norm(keiyaku).replace('-', '')
            block_rows.append((k, basho_key, amount))
            if basho_key == basho:
                result[k] += amount
                nyuko_dates[k].add(cols[3].strip())
                if sign_cell == '-':
                    minus_rows[k].append((idx, cols[3].strip(), amount))
        elif keiyaku == 'T':
            # ブロック小計行: 直前までのデータ行合計と照合 (不一致=東レが小計から除外した行がある疑い)
            # 1ブロック内に複数の入庫場所(A010/A010P)が混在するため場所は全て列挙し、
            # 差額と同額の契約NOがあれば「小計から除外された行」の候補として返す
            try:
                t_val = float(cols[9].strip().replace(',', ''))
            except ValueError:
                t_val = None
            if t_val is not None and abs(block_sum - t_val) > 0.5:
                d = block_sum - t_val
                per_k = defaultdict(float)
                basho_of = {}
                for k, b, amt in block_rows:
                    per_k[k] += amt
                    basho_of[k] = b
                candidates = [(k, basho_of[k], s) for k, s in per_k.items() if abs(s - d) <= 0.5]
                bashos = '/'.join(sorted({b for _, b, _ in block_rows}))
                ct_errors.append((block_start or idx, idx, bashos, t_val, block_sum, d, candidates))
            block_sum = 0.0
            block_start = None
            block_rows = []
    if not result:
        sys.exit(f'①CSVから入庫場所 {basho} のデータ行を1件も取り込めませんでした。'
                 f'フォーマット変更または --basho の指定を確認してください: {path}')
    return dict(result), dict(by_basho), dict(nyuko_dates), warnings, ct_errors, dict(minus_rows)


def read_nagaoka(path: Path):
    """② 東レまとめシート。
    戻り値: (依頼NO->AA合計, 契約NO->AA合計, 契約NO->依頼NO集合, 契約NO形式でないC列値のリスト, 警告リスト)
    """
    rows = CalamineWorkbook.from_path(str(path)).get_sheet_by_name('東レまとめ').to_python()
    # ヘッダー検証: 列位置のずれ(列挿入等)を検出
    if len(rows) < 6 or norm(rows[3][26]) != '合計' or '契約' not in norm(rows[2][2]):
        sys.exit(f'②「東レまとめ」の構成が想定と異なります(C3セル「契約No.」/AA4セル「合計」を確認): {path}')
    by_irai = defaultdict(float)
    by_keiyaku = defaultdict(float)
    keiyaku_to_irai = defaultdict(set)
    bad_keiyaku = []
    warnings = []
    for i, r in enumerate(rows[5:], start=6):  # 6行目以降がデータ
        a = norm(r[0]) if len(r) > 0 else ''
        b = norm(r[1]) if len(r) > 1 else ''
        aa = r[26] if len(r) > 26 else None    # AA列 = 加工賃合計
        if not a or not b or a == '0' or b == '0':
            continue
        if not isinstance(aa, (int, float)):
            if aa not in ('', None):
                warnings.append(f'②{path.name} {i}行目: 依頼NO {a}-{b} のAA列「{aa}」が数値でないため除外')
            continue
        irai = f'{a}-{b}'
        by_irai[irai] += aa
        c = norm(r[2]).replace('-', '') if len(r) > 2 else ''
        if KEIYAKU_PAT.match(c):
            by_keiyaku[c] += aa
            keiyaku_to_irai[c].add(irai)
        elif c and c != '0':
            bad_keiyaku.append((i, c, aa))
    if not by_irai:
        sys.exit(f'②「東レまとめ」からデータ行を1件も取り込めませんでした: {path}')
    return dict(by_irai), dict(by_keiyaku), dict(keiyaku_to_irai), bad_keiyaku, warnings


def read_aladdin(path: Path, customer=None):
    """③ 項目=「加工金額」の --合計-- 列。
    customer を指定すると「得意先」列がそのコードの行だけを集計する (湖南は東レ以外の得意先が混在するため)。
    戻り値: (依頼NO->加工金額, 対象年月文字列, 対象年月(y,m) or None)"""
    rows = CalamineWorkbook.from_path(str(path)).get_sheet_by_index(0).to_python()
    taisho = ''
    header_idx = None
    for i, r in enumerate(rows[:10]):
        joined = ''.join(str(v) for v in r)
        if '対象年月' in joined:
            taisho = joined.split(':')[-1].strip()
        if any(norm(v) == '依頼NO' for v in r) and any(norm(v) == '項目' for v in r):
            header_idx = i
            break
    if header_idx is None:
        sys.exit(f'③のヘッダー行(依頼NO/項目)が見つかりません: {path}')
    header = [norm(v) for v in rows[header_idx]]
    try:
        i_irai, i_item, i_total = header.index('依頼NO'), header.index('項目'), header.index('--合計--')
    except ValueError as e:
        sys.exit(f'③のヘッダーに想定列が見つかりません(依頼NO/項目/--合計--): {path}\n{e}')
    i_cust = None
    if customer:
        if '得意先' not in header:
            sys.exit(f'③のヘッダーに「得意先」列が無いため得意先 {customer} で絞り込めません: {path}')
        i_cust = header.index('得意先')
    result = defaultdict(float)
    for r in rows[header_idx + 1:]:
        if len(r) <= i_total or norm(r[i_item]) != '加工金額':
            continue
        if i_cust is not None and norm(r[i_cust]) != norm(customer):
            continue
        v = r[i_total]
        if isinstance(v, (int, float)):
            result[norm(r[i_irai])] += v
    if not result:
        sys.exit(f'③から「加工金額」行を1件も取り込めませんでした'
                 + (f'(得意先 {customer} で絞り込み)' if customer else '')
                 + f'。フォーマット変更の可能性: {path}')
    m = re.search(r'(\d{4})\s*年\s*(\d{1,2})\s*月', taisho)
    taisho_ym = (int(m.group(1)), int(m.group(2))) if m else None
    return dict(result), taisho, taisho_ym


FILL_HEADER = PatternFill('solid', fgColor='4472C4')
FILL_MISMATCH = PatternFill('solid', fgColor='FFC7CE')   # 不一致: 赤
FILL_NEXT = PatternFill('solid', fgColor='F8CBAD')       # 翌月記載(②の記載月誤り): オレンジ
FILL_ONLY = PatternFill('solid', fgColor='FFEB9C')       # 片側のみ: 黄
FILL_PREV = PatternFill('solid', fgColor='BDD7EE')       # 前月調整: 青
FILL_PRIOR = PatternFill('solid', fgColor='DDEBF7')      # 前月過不足(①未反映): 淡青
FILL_RESOLVED = PatternFill('solid', fgColor='E2EFDA')   # 月ずれ解消(累計一致): 緑 (橙との色覚混同を避ける明るめ)
FILL_INFO = PatternFill('solid', fgColor='E7E6E6')       # 形式不正: 灰 (視覚的に後退させる明るめ)
FILL_OK = PatternFill('solid', fgColor='C6EFCE')         # 問題なし表示: 緑
FILL_WARNBOX = PatternFill('solid', fgColor='FDE9E9')    # 警告ボックス: 淡赤
FILL_MANUAL = PatternFill('solid', fgColor='E4DFEC')     # 手動判定(①正/②正): 薄紫

MANUAL_JUDGES = ('①正(手動)', '②正(手動)')

# 判定セルの文字色 (塗りとのコントラストAA以上・色覚多様性でも判別可能)
JUDGE_COLOR = {
    '不一致': '9C0006', '翌月記載': '833C00', '①のみ': '7F5F00', '②のみ': '7F5F00',
    '③のみ': '7F5F00', '前月調整': '1F4E79', '前月過不足': '1F4E79', '月ずれ解消': '375623', '枝番統合一致': '375623', '形式不正': '595959',
    # 検証D (②まとめ整合性)
    '金額差': '9C0006', '参照ずれ': '9C0006', '未取込': '9C0006', '式異常': '833C00', '直接入力': '7F5F00',
    # 手動判定 (手動判定.csv)
    '①正(手動)': '5B2C6F', '②正(手動)': '5B2C6F',
}

THIN = Side(style='thin', color='BFBFBF')
BORDER_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX = Side(style='medium', color='4472C4')


def _judge_rank(judge, diff, amounts):
    """並び順: 不一致(差額大)→翌月記載→片側のみ(金額大)→前月調整→月ずれ解消→形式不正→一致"""
    scale = max((abs(a) for a in amounts if isinstance(a, (int, float))), default=0)
    if judge in ('不一致', '金額差'):
        return (0, -abs(diff) if isinstance(diff, (int, float)) else 0)
    if judge in ('翌月記載', '参照ずれ', '未取込', '式異常'):
        return (1, -scale)
    if judge in ('①のみ', '②のみ', '③のみ', '直接入力'):
        return (2, -scale)
    if judge == '前月調整' or judge == '前月過不足' or judge in MANUAL_JUDGES:
        return (3, -scale)
    if judge in ('月ずれ解消', '枝番統合一致'):
        return (4, -scale)
    if judge == '形式不正':
        return (5, -scale)
    return (6, 0)


def _judge_fill(judge):
    if judge in ('不一致', '金額差', '参照ずれ', '未取込'):
        return FILL_MISMATCH
    if judge in ('翌月記載', '式異常'):
        return FILL_NEXT
    if judge in ('①のみ', '②のみ', '③のみ', '直接入力'):
        return FILL_ONLY
    if judge == '前月調整':
        return FILL_PREV
    if judge == '前月過不足':
        return FILL_PRIOR
    if judge in MANUAL_JUDGES:
        return FILL_MANUAL
    if judge in ('月ずれ解消', '枝番統合一致'):
        return FILL_RESOLVED
    if judge == '形式不正':
        return FILL_INFO
    return None


def _write_sheet(wb, title, headers, rows):
    """検証明細シートを作成。rows = [キー, 金額..., 差額, 判定, 備考]"""
    ws = wb.create_sheet(title)
    n_cols = len(headers)
    note_col = n_cols            # 備考列 (1始まり)
    diff_col = next((i for i, h in enumerate(headers, start=1) if h.startswith('差額')), None)
    rep_col = next((i for i, h in enumerate(headers, start=1) if h == '報告計上額'), None)
    wrap_note = Alignment(wrap_text=True, vertical='center')
    ws.append(headers)
    for c in ws[1]:
        c.font = F(bold=True, color='FFFFFF')
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal='center', vertical='center')
    if rep_col:
        # サマリKPI「報告する過不足」の元になる列。ヘッダーを濃色にして区別する
        hc = ws.cell(row=1, column=rep_col)
        hc.fill = PatternFill('solid', fgColor='1F4E79')
        hc.value = '報告計上額\n(列合計=報告する過不足)'
        hc.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        hc.font = F(bold=True, color='FFFFFF', size=9)
    ws.row_dimensions[1].height = 30
    for row in rows:
        ws.append(row)
        r_i = ws.max_row
        judge = row[-2]
        fill = _judge_fill(judge)
        jcolor = JUDGE_COLOR.get(judge)
        for i, c in enumerate(ws[r_i], start=1):
            if fill:
                c.fill = fill
            if isinstance(c.value, float):
                # 合算で生じる浮動小数誤差(44,808.0000001 / -0.0000001)を消し、整数なら整数表示
                c.value = round(c.value, 2)
                if c.value == int(c.value):
                    c.value = int(c.value)
            if isinstance(c.value, (int, float)):
                # 等幅版フォントで桁を揃え、負数は赤で表示。差額(不一致行)と報告計上額は太字
                c.font = F(name=FONT_NUM, bold=(i == diff_col and judge == '不一致') or i == rep_col)
                c.number_format = '#,##0;[Red]-#,##0' if isinstance(c.value, int) else '#,##0.00;[Red]-#,##0.00'
            if i == n_cols - 1:              # 判定列
                c.alignment = CENTER
                if jcolor:
                    c.font = F(bold=True, color=jcolor)
            if i == note_col:
                c.alignment = wrap_note
        # 備考の長さに応じて行高さを確保 (列幅70 ≒ 全角33文字/行)
        note = str(row[-1] or '')
        lines = max(1, -(-len(note) // 33))
        if lines > 1:
            ws.row_dimensions[r_i].height = 14.5 * lines + 4
    widths = [14] + [16] * (n_cols - 3) + [12, 70]
    for i, wd in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = '1:1'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _write_guide(wb, info):
    """「レポートの見方」ガイドシート (セルベースのフローチャート中心)。文言は工場プロファイル(info)に従う"""
    tol, basho = info['許容差'], info['入庫場所']
    factory, other, name2, amt2 = info['工場'], info['他工場'], info['②名称'], info['②金額列']
    ws = wb.create_sheet('レポートの見方')
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '808080'
    ws.column_dimensions['A'].width = 2
    for i in range(2, 15):  # B〜N の13列グリッド
        ws.column_dimensions[get_column_letter(i)].width = 8.6

    _merge_put(ws, 1, 2, 14, f'レポートの見方【{factory}】', F(bold=True, size=15, color='1F4E79'))
    ws.row_dimensions[1].height = 24
    _merge_put(ws, 2, 2, 14, f'{factory}の後加工工賃の3つのデータソースを突合し、金額の矛盾点を検出した結果ブックです。',
               F(size=10, color='595959'))
    r = 4

    def section(title):
        nonlocal r
        _merge_put(ws, r, 2, 14, title, F(bold=True, size=12, color='1F4E79'),
                   align=Alignment(vertical='center'))
        for cc in range(2, 15):
            cur = ws.cell(row=r, column=cc).border
            ws.cell(row=r, column=cc).border = Border(bottom=Side(style='medium', color='4472C4'),
                                                      top=cur.top, left=cur.left, right=cur.right)
        ws.row_dimensions[r].height = 20
        r += 2

    def down_arrow(col=2, label='▼'):
        nonlocal r
        c = ws.cell(row=r, column=col, value=label)
        c.font = F(size=9, bold=True, color='808080')
        c.alignment = CENTER
        ws.row_dimensions[r].height = 13
        r += 1

    def text_rows(pairs, label_end=4):
        nonlocal r
        for label, text in pairs:
            _merge_put(ws, r, 2, label_end, label, F(bold=True, size=10), align=WRAP_TOP)
            _merge_put(ws, r, label_end + 1, 14, text, F(size=10), align=WRAP_TOP)
            ws.row_dimensions[r].height = 14.5 * max(1, -(-len(text) // 52)) + 2
            r += 1
        r += 1

    # ---- 1. 突合構造図 ----
    section('1. データの突合構造')
    box2_name, box2_sub = info['box2']
    boxes = [
        (2, 4, '① 東レ検収CSV', f'お支払データ ({basho}・契約NO別)', 'DDEBF7', '1F4E79'),
        (7, 9, box2_name, box2_sub, 'E2EFDA', '375623'),
        (12, 14, '③ アラジン', '依頼NO別・加工金額' + (f' (得意先{info["③得意先"]})' if info['③得意先'] else ''), 'FFF2CC', '7F5F00'),
    ]
    for c1, c2, name, sub, bg, fg in boxes:
        _merge_put(ws, r, c1, c2, name, F(bold=True, size=11, color=fg), align=CENTER)
        _merge_put(ws, r + 1, c1, c2, sub, F(size=9, color=fg), align=CENTER)
        _draw_box(ws, r, c1, r + 1, c2, fill=PatternFill('solid', fgColor=bg), color=fg)
    for c1, c2, name, sub in ((5, 6, '検証A', '契約NOで突合'), (10, 11, '検証B', '依頼NOで突合')):
        _merge_put(ws, r, c1, c2, f'◀ {name} ▶', F(bold=True, size=10.5), align=CENTER)
        _merge_put(ws, r + 1, c1, c2, sub, F(size=9, color='595959'), align=CENTER)
    ws.row_dimensions[r].height = 20
    ws.row_dimensions[r + 1].height = 16
    r += 2
    _merge_put(ws, r, 2, 14, f'検証A = 東レからの支払と長岡側計上の照合　／　検証B = 長岡側の2システム間({name2}とアラジン)の照合'
               + ('　／　検証C = 月次処理ファイルの東レ合計との総額照合' if info.get('検証C') else '')
               + (f'　／　検証D = ②「{MATOME_SHEET}」が元4シートと整合しているかの内部検査' if info.get('検証D') else ''),
               F(size=9.5, color='595959'), align=CENTER)
    r += 2

    # ---- 2. 判定と色の意味 (表) ----
    section('2. 判定と色の意味')
    for c1, c2, h in ((2, 3, '判定(色)'), (4, 8, '意味'), (9, 14, 'とるべき対応')):
        _merge_put(ws, r, c1, c2, h, F(bold=True, color='FFFFFF'), fill=FILL_HEADER, align=CENTER)
    ws.row_dimensions[r].height = 18
    r += 1
    wrap_mid = Alignment(wrap_text=True, vertical='center')
    legend_rows = list(LEGEND) + [('一致', None, '595959', f'金額が許容差(±{tol}円)以内で一致。色なし', '対応不要')]
    for judge, fill, fg, desc, action in legend_rows:
        _merge_put(ws, r, 2, 3, judge, F(bold=True, size=10, color=fg), fill=fill, align=CENTER_WRAP)
        _merge_put(ws, r, 4, 8, desc, F(size=10), align=wrap_mid)
        _merge_put(ws, r, 9, 14, action, F(size=10), align=wrap_mid)
        _draw_box(ws, r, 2, r, 14, color='BFBFBF', style='thin')
        ws.row_dimensions[r].height = 28
        r += 1
    r += 1

    # ---- 3. 毎月の確認フロー ----
    section('3. 毎月の確認フロー')
    steps = [
        ('サマリの6枚のカード(検証A/B要確認・過不足・判明済み・警告・検算残差)を確認。警告は最優先で対処', 'DDEBF7', '1F4E79'),
        ('検証A: 赤(不一致)を上から確認。「契約間振替の疑い」の注記はネット差額で評価する', 'FFC7CE', '9C0006'),
        ('検証A: 橙(翌月記載)は②の2ファイルを修正(当月へ追記・翌月から削除) → 修正後に再実行して数字を確定', 'F8CBAD', '833C00'),
        ('検証A: 黄(①のみ)は備考の①入庫日を手がかりに①CSV原本を確認し、②への記載漏れか判断', 'FFEB9C', '7F5F00'),
        ('検証B: 赤・黄を確認。緑(月ずれ解消・枝番統合一致)は飛ばしてよい。「計上月ずれの可能性」の注記が無いものを優先', 'F2F2F2', '333333'),
        (f'青(前月調整)は判明済み。備考の「当時{amt2}」「①との差」を参考確認するだけでよい', 'BDD7EE', '1F4E79'),
        ('淡青(前月過不足)は①列が空・②列=プラス差。プラス側は①に現れないため報告に残す', 'DDEBF7', '1F4E79'),
        (f'「報告メール下書き」の数字と注意書きを確認し、{other}分を追記して送信する', 'C6EFCE', '375623'),
    ]
    if info.get('検証C'):
        steps.insert(5, ('検証C(サマリ下部): 月次処理ファイルの東レ合計と②③①の総額が一致しているか確認。不一致は警告にも出る', 'E2EFDA', '375623'))
    if info.get('検証D'):
        steps.insert(1, (f'検証D(サマリ下部・検証Dシート): ②「{MATOME_SHEET}」の要修正(金額差・参照ずれ・未取込)が0件か確認。'
                         'あれば②を直して再実行 (②総額が誤っていると検証A/Bの結果も信用できない)', 'E4DFEC', '5B2C6F'))
    for i, (text, bg, fg) in enumerate(steps, start=1):
        c = ws.cell(row=r, column=2, value=i)
        c.font, c.alignment, c.fill = F(bold=True, size=12, color='FFFFFF'), CENTER, FILL_HEADER
        _merge_put(ws, r, 3, 14, text, F(size=10.5, color=fg), fill=PatternFill('solid', fgColor=bg),
                   align=Alignment(vertical='center', wrap_text=True, indent=1))
        _draw_box(ws, r, 2, r, 14, color='BFBFBF', style='thin')
        ws.row_dimensions[r].height = 28
        r += 1
        if i < len(steps):
            down_arrow()
    r += 1

    # ---- 4. ①のみ契約NOの自動判定 (分岐図) ----
    section('4. 「①にあるが②当月に無い」契約NOの自動判定')
    _merge_put(ws, r, 2, 8, '①に検収があるのに ②当月まとめに無い契約NO',
               F(bold=True, size=10.5, color='1F4E79'), fill=PatternFill('solid', fgColor='DDEBF7'),
               align=CENTER_WRAP)
    _draw_box(ws, r, 2, r, 8, color='1F4E79')
    ws.row_dimensions[r].height = 24
    r += 1
    down_arrow(col=5)

    def branch(question, res_text, res_fill, res_fg, last=False):
        nonlocal r
        _merge_put(ws, r, 2, 8, question, F(bold=True, size=10.5), align=CENTER)
        _draw_box(ws, r, 2, r, 8, color='4472C4')
        _merge_put(ws, r, 9, 10, 'YES ▶', F(bold=True, size=10, color='375623'), align=CENTER)
        _merge_put(ws, r, 11, 14, res_text, F(bold=True, size=10, color=res_fg), fill=res_fill,
                   align=CENTER_WRAP)
        _draw_box(ws, r, 11, r, 14, color=res_fg, style='thin')
        ws.row_dimensions[r].height = 32
        r += 1
        if not last:
            down_arrow(col=5, label='NO ▼')

    branch('過去月の②まとめに あるか？', '前月調整 (青)\n判明済み・参考確認のみ', FILL_PREV, '1F4E79')
    branch('翌月の②まとめに あるか？', '翌月記載 (橙)\n②2ファイル修正 → 再実行', FILL_NEXT, '833C00', last=True)
    down_arrow(col=5, label='NO ▼')
    _merge_put(ws, r, 2, 8, '①のみ (黄)　②への記載漏れの疑い → 原本で要調査',
               F(bold=True, size=10.5, color='7F5F00'), fill=FILL_ONLY, align=CENTER_WRAP)
    _draw_box(ws, r, 2, r, 8, color='7F5F00', style='thin')
    ws.row_dimensions[r].height = 24
    r += 2

    # ---- 5〜8. テキストセクション ----
    section('5. シート構成')
    sheet_rows = [
        ('サマリ', '全体の件数・総額・報告用内訳・警告。まずここを見る'),
        ('検証A / 検証B', '突合結果の明細。色付き行が確認対象。判定列のオートフィルタで絞り込み可'),
    ]
    if info.get('検証D'):
        sheet_rows.append(('検証D', f'②「{MATOME_SHEET}」の内部検査結果。赤(金額差・参照ずれ・未取込)は②の集計誤りなので必ず直す。'
                                    '橙(式異常)=元シートのG/AA列が空欄や別式。黄(直接入力)=式が値で上書きされた箇所(単価変更時は元シートH〜X列を直すのが正)'))
    sheet_rows += [
        ('報告メール下書き', '東レ宛メール文面に数字を自動反映した下書き。冒頭の注意書きを確認・削除して使う'),
        ('原本コピー', '①②③の元データと過去月・翌月の②。要確認行は薄黄色でハイライト済み。Ctrl+Fで検索して原因調査'),
    ]
    text_rows(sheet_rows)

    section('6. 突合キーとデータ仕様')
    spec_rows = [
        ('検証Aのキー', info['key_a']),
        ('検証Bのキー', info['key_b']),
        ('①の対象範囲', f'入庫場所 {basho}({factory})のみ。他の入庫場所({other}等)は対象外(除外額はサマリに表示)'),
        ('①の取消行', 'H列が「-」の行は取消(マイナス)行。金額列は絶対値なので符号を反転し、契約NOごとに純額で集計する(T小計と同じ扱い)。該当契約NOの備考に行番号と金額を注記'),
        ('金額0円の行', '両側とも0円の一致行、および0円の片側のみ(取消済み・翌月予定の空行など)は無効行として検証A/Bシートに出力しない。除外件数はサマリ下部に表示'),
        ('手動判定', f'●自動検証\\{MANUAL_FILE} (列: 工場,対象月,契約NO,正とする側,理由) に書いた契約NOは、その工場・その月の検証Aの不一致に限り'
                  '「①正(手動)」「②正(手動)」として要確認から外す。①正=②を①金額に修正する(報告計上額は空欄)、②正=東レへ報告する(報告計上額=差額)。'
                  '対象月は必須で、他の月には効かない(今回限りの判断を残さないため)。指定した契約NOが当月の不一致でなければ警告が出る'),
        ('前月過不足', f'●自動検証\\{PRIOR_FILE} (列: 工場,対象月,契約NO,依頼NO,①東レ金額,②長岡金額,理由)。'
                  '前月の工賃差検証で東レへ次月修正を依頼したが当月①CSVに調整行が未反映の差額を登録する。'
                  '検証Aでは①列を空・②列にプラス差(長岡−東レ)を表示し、備考に前月当時①②を記す。'
                  'マイナス側は①取消行で「前月調整」に出るがプラス側は①に現れない。報告計上額=前月の①−②。対象月=調整を反映すべき当月'),
        ('金額の取り方', f'{info["take2"]}　／　③ = 項目「加工金額」行の「--合計--」列'
         + (f'(得意先 {info["③得意先"]} の行のみ)' if info['③得意先'] else '')),
    ]
    if info.get('月次ファイル'):
        spec_rows.append(('月次処理ファイル', f'{info["月次ファイル"]}。契約NO等の明細が無いため「月次検証」シート東レ行(月次実績①Excel/月次実績表②アラジン/売上明細③)と'
                          '「集計表」東レ合計の金額だけを、本検証の②③①総額と照合する(検証C)'))
    if info.get('検証D'):
        spec_rows.append(('②まとめの構造(検証D)', f'「{MATOME_SHEET}」は {"/".join(MATOME_SRC_SHEETS)} の各行を =東レT!A6 のように列ごとに参照し、'
                          'G(単価合計=SUM(H:X))・Z・AA(=ROUNDUP(E×G,0))だけ自行で再計算する。'
                          '検証Dは (1)参照式の行・列・シートずれ (2)元シートのデータ行の取込漏れ (3)G/AA列の式崩れ (4)行別のAA金額差 を検査する'))
    text_rows(spec_rows)

    section('7. 備考欄の自動注記')
    text_rows([
        ('契約間振替の疑い', '検証Aで同一依頼NO内の不一致どうしの差額が相殺関係にある場合に注記。契約NO間の数量振替ミスの典型パターン。ネット差額のみが実質の過不足'),
        ('計上月ずれの可能性', '検証Bで依頼NOの月番号(例 Y7-系)が対象月と異なる場合に注記。過去月依頼の差異は月ずれの可能性が高い'),
        ('①入庫日', '検証Aの①のみ・前月調整の行に①CSVの入庫月日(yymmdd)を自動転記。原因調査の起点に使う'),
        ('①小計から除外の疑い', '①CSVのブロック小計(T行)とデータ行合計の差と同額の契約NOに注記。東レ側で取消等により小計から除外された行の可能性。①のみ行に付いていれば「支払対象外」の疑いが濃い'),
        ('枝番統合一致', '検証Bで③が枝番(例 C8-3-1)に分割計上し、親依頼NOに合算すると②と一致する場合。緑で対象外'),
    ])

    section('8. 報告用内訳の見方 (サマリ)')
    text_rows([
        ('分解式', '① 東レお支払データ = ②当月実売上 + 前月調整分 + 翌月記載 + 当月差異 + 未解明分(①のみ−②のみ)'),
        ('当月差異', '検証Aの不一致の差額合計'),
        ('報告する過不足', '当月差異 + 翌月記載 + 前月過不足 + ①のみ − ②のみ。東レへの月次報告メールで「次月調整をお願いする過不足」として報告する金額。検証Aシートの「報告計上額」列(不一致行=差額①-②、翌月記載・①のみ行=①金額、②のみ行=−②金額、前月過不足行=前月の差額①-②)を合計したものと一致する'),
        ('検算残差', '0であれば区分の内部整合OK。ただし許容差内の微差・形式不正行の金額は分解に含まれないため、0でも警告欄と形式不正行の確認は必要'),
    ])

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


FILL_TRACE = PatternFill('solid', fgColor='FFF2CC')  # 原本コピー内の要確認キー該当行


def _append_raw_sheet(wb, title, rows, freeze='A2', highlight=None):
    """元データのコピーシートを追加 (タイトルは31文字制限に切り詰め)。
    highlight(行の値リスト)->True の行は薄黄でハイライトし、検証明細から追いやすくする"""
    ws = wb.create_sheet(title[:31])
    ws.sheet_properties.tabColor = 'A6A6A6'
    for r in rows:
        vals = list(r)
        ws.append(vals)
        if highlight and highlight(vals):
            for c in ws[ws.max_row]:
                c.fill = FILL_TRACE
    ws.freeze_panes = freeze
    return ws


def _mail_figures(info):
    """メール本文用の自工場数字。

    前月調整分 = −(①に現れた前月調整額) + 前月過不足の差額合計
      例: 13500+4470+1320 + (−20)+(−8899) = 10371
      (前月の工賃差検証で依頼した次月修正のネット。①未反映のプラス側も含む)
    当月調整分 = −(当月差異+翌月記載+①のみ−②のみ)  ※前月過不足は当月に含めない
    """
    adj = round(info.get('A前月調整額', 0) or 0)
    prior = round(info.get('A前月過不足額', 0) or 0)
    next_amt = round(info.get('A翌月記載額', 0) or 0)
    # 当月分のみ (前月過不足は前月調整分側へ)
    tougetsu = round((info.get('報告当月差異', 0) or 0) + next_amt
                     + (info.get('報告①のみ額', 0) or 0) - (info.get('報告②のみ額', 0) or 0))
    n = (info.get('A不一致', 0) + info.get('A手動②正件数', 0) + info.get('A翌月記載件数', 0)
         + info.get('A①のみ', 0) + info.get('A②のみ', 0))
    if abs(tougetsu) < 0.5:
        n = 0
    label = info.get('工場', '')
    short = '国分' if '国分' in label else '湖南' if '湖南' in label else label.replace('工場', '')
    return {
        '工場': label, 'short': short, 'basho': info.get('入庫場所', ''),
        't1': round(info.get('①総額', 0) or 0),
        't2': round(info.get('報告②実売上', 0) or 0),
        'adj_disp': -adj + prior,  # 前月調整(符号反転) + 前月過不足差額
        'tougetsu_disp': -tougetsu, 'tougetsu': tougetsu, 'n': n,
        'source': None,
    }


def _find_peer_result(base: Path, own_label):
    """他工場の最新 検証結果_*.xlsx を探す。国分→湖南 / 湖南→国分。"""
    peer = '湖南工場' if '国分' in own_label else '国分工場' if '湖南' in own_label else None
    if not peer:
        return None
    files = sorted(base.glob(f'検証結果_{peer}_*.xlsx'), key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else None


def _archive_old_results(base: Path, factory_label=None, keep: Path = None):
    """●自動検証直下の 検証結果_*.xlsx / .html のうち、工場ごとの最新以外を 過去検証結果\\ へ移す。

    factory_label 指定時はその工場のみ。keep は残すファイル(今回出力分)。
    戻り値: 移動した Path のリスト。
    """
    import shutil
    arch = base / ARCHIVE_DIR
    arch.mkdir(exist_ok=True)
    moved = []
    # 工場ラベルごとにグルーピング (検証結果_国分工場_YYYYMMDD_HHMMSS.xlsx)
    groups = {}
    for p in base.glob('検証結果_*.xlsx'):
        # 検証結果_<工場>_<timestamp>.xlsx → 工場名は末尾タイムスタンプ2段を除く
        stem = p.stem  # 検証結果_国分工場_20260903_180910
        parts = stem.split('_')
        if len(parts) < 4:
            continue
        # parts[0]=検証結果, parts[-2]=日期, parts[-1]=時刻, 中間が工場
        label = '_'.join(parts[1:-2])  # 国分工場 / 湖南工場
        if factory_label and label != factory_label:
            continue
        groups.setdefault(label, []).append(p)
    for label, files in groups.items():
        files_sorted = sorted(files, key=lambda x: x.stat().st_mtime, reverse=True)
        latest = files_sorted[0]
        if keep and keep.resolve() in {f.resolve() for f in files_sorted}:
            latest = keep.resolve()
            # keep 以外を退避
            olds = [f for f in files_sorted if f.resolve() != latest]
        else:
            olds = files_sorted[1:]
        for f in olds:
            dest = arch / f.name
            if dest.exists():
                dest = arch / f'{f.stem}_dup{f.suffix}'
            shutil.move(str(f), str(dest))
            moved.append(dest)
            html = f.with_suffix('.html')
            if html.is_file():
                hdest = arch / html.name
                if hdest.exists():
                    hdest = arch / f'{html.stem}_dup.html'
                shutil.move(str(html), str(hdest))
                moved.append(hdest)
    # 最新xlsxに紐づかない孤立htmlも退避 (同じ工場プレフィックス)
    for html in base.glob('検証結果_*.html'):
        xlsx = html.with_suffix('.xlsx')
        if xlsx.is_file():
            continue  # 最新側に残っている
        parts = html.stem.split('_')
        if len(parts) < 4:
            continue
        label = '_'.join(parts[1:-2])
        if factory_label and label != factory_label:
            continue
        dest = arch / html.name
        if dest.exists():
            dest = arch / f'{html.stem}_dup.html'
        shutil.move(str(html), str(dest))
        moved.append(dest)
    return moved


def _load_mail_figures_from_result(path: Path):
    """他工場の検証結果Excel(サマリ)からメール用数字を読む。読めなければ None。"""
    try:
        wb = load_workbook(str(path), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        return None
    try:
        if 'サマリ' not in wb.sheetnames:
            return None
        ws = wb['サマリ']
        kv = {}
        title = ''
        kpi_a = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=80, values_only=True), 1):
            vals = [c for c in row if c is not None]
            if not vals:
                continue
            if not title and isinstance(vals[0], str) and '検証レポート' in vals[0]:
                title = vals[0]
            if i == 5 and isinstance(vals[0], str) and vals[0].endswith('件'):
                m = re.search(r'(\d+)', vals[0])
                if m:
                    kpi_a = int(m.group(1))
            if len(vals) >= 2 and isinstance(vals[0], str):
                kv[vals[0].strip()] = vals[1]
    finally:
        wb.close()

    def num(*keys):
        for k, v in kv.items():
            if any(x in k for x in keys) and isinstance(v, (int, float)):
                return round(v)
        return None

    t1 = num('① 東レお支払', '① 東レ検収')
    t2 = num('② 当月実売上')
    if t2 is None:
        t2 = num('② 加工賃試算', '② 長岡明細')
    adj = num('前月以前調整分')
    if adj is None:
        adj = num('　前月調整分')
    tougetsu = num('報告する過不足')
    if None in (t1, t2, adj, tougetsu):
        return None
    label = '湖南工場' if '湖南' in title or '湖南' in path.name else '国分工場'
    short = '国分' if '国分' in label else '湖南'
    basho = 'A010P' if short == '湖南' else 'A010'
    n = 0 if abs(tougetsu) < 0.5 else (kpi_a if kpi_a is not None else 0)
    return {
        '工場': label, 'short': short, 'basho': basho,
        't1': t1, 't2': t2, 'adj_disp': -adj, 'tougetsu_disp': -tougetsu,
        'tougetsu': tougetsu, 'n': n, 'source': path.name,
    }


def _fmt_yen(n):
    return f'{int(round(n)):,}'


def _fmt_paren_yen(n):
    """メール本文の件数横（-3,692）/（ 0）表記"""
    v = int(round(n))
    if v == 0:
        return ' 0'
    return f'{v:,}'


def _prior_mail_note_lines(info):
    """前月過不足(①未反映のプラス差)をメール本文に説明する行リスト。無ければ空。"""
    details = info.get('A前月過不足明細') or []
    if not details:
        return []
    items = '、'.join(
        f'{d["依頼表示"]} {d["契約NO"]} +{_fmt_yen(d["プラス差"])}円' for d in details)
    return [
        '※前月調整分には、前月のプラス差で当月①お支払いデータに含まれない分を含みます。',
        f'　{items}',
        '　（マイナス側の調整は①取消行として現れますが、プラス側は①データに含まれません）',
    ]


def _build_mail_html(info, rows):
    """東レ宛メール本文のHTML。添付の実メール体裁に合わせる (表＋件数サマリ＋締め)。"""
    ym = info['対象月ラベル']
    next_ym = info['翌月ラベル']
    total_t1 = sum(r['t1'] for r in rows)
    total_t2 = sum(r['t2'] for r in rows)
    total_adj = sum(r['adj_disp'] for r in rows)
    total_tg = sum(r['tougetsu_disp'] for r in rows)
    total_n = sum(r['n'] for r in rows)

    td = 'border:1px solid #8FAADC;padding:4px 10px;'
    th = 'border:1px solid #8FAADC;padding:4px 10px;background:#BDD7EE;font-weight:bold;text-align:center;'

    def tr(label, t1, t2, adj, tg, bold=False):
        w = 'font-weight:bold;' if bold else ''
        align_l = 'text-align:center;'
        align_r = 'text-align:right;'
        return (
            f'<tr style="{w}">'
            f'<td style="{td}{align_l}">{label}</td>'
            f'<td style="{td}{align_r}">{_fmt_yen(t1)}</td>'
            f'<td style="{td}{align_r}">{_fmt_yen(t2)}</td>'
            f'<td style="{td}{align_r}">{_fmt_yen(adj)}</td>'
            f'<td style="{td}{align_r}">{_fmt_yen(tg)}</td>'
            f'</tr>'
        )

    header = (
        f'<tr>'
        f'<td style="{th}">拠点</td>'
        f'<td style="{th}">東レ支払いデータ</td>'
        f'<td style="{th}">実売上</td>'
        f'<td style="{th}">前月調整分</td>'
        f'<td style="{th}">当月調整分</td>'
        f'</tr>'
    )
    body_rows = ''.join(
        tr(f"{r['short']} ({r['basho']})" if r['basho'] else r['short'],
           r['t1'], r['t2'], r['adj_disp'], r['tougetsu_disp'])
        for r in rows
    )
    total_row = tr('合計', total_t1, total_t2, total_adj, total_tg, bold=True)

    # 件数行: 添付メールと同じ字下げ・全角スペース
    detail = ''.join(
        f'<div>　{r["工場"]}　　　 {r["n"]}件（{_fmt_paren_yen(r["tougetsu_disp"])}）</div>'
        for r in rows
    )
    highlight = f'{total_n}件の差異（{_fmt_yen(abs(total_tg))}円）過不足'
    prior_note = ''.join(f'<div>{line}</div>' for line in _prior_mail_note_lines(info))

    html = f'''<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>後加工賃支払い明細 {ym} 検証結果</title></head>
<body style="font-family:'Yu Gothic','Meiryo','MS PGothic',sans-serif;font-size:14px;color:#000;line-height:1.7;">
<div>東レ株式会社</div>
<div>トーレペフ事業部 御中</div>
<br>
<div>いつも大変お世話になっております。</div>
<div>長岡産業の難波です。</div>
<br>
<div>掲題の件、下記にご報告申し上げます。</div>
<br>
<table style="border-collapse:collapse;font-size:13px;font-family:'Yu Gothic','Meiryo',sans-serif;">
{header}
{body_rows}
{total_row}
</table>
<br>
<div>合計（税抜き）　{_fmt_yen(total_t1)}円の御社お支払いデータに対し</div>
{detail}
<div>合計　<span style="font-weight:bold;text-decoration:underline;">{highlight}</span>がありました。</div>
{prior_note}<br>
<div>添付ファイルをご確認いただき、次月（{next_ym}）での</div>
<div>ご調整をよろしくお願い申し上げます。</div>
<div>今後とも何卒よろしくお願い申し上げます。</div>
<br>
<div>以上</div>
</body></html>
'''
    return html


def _write_mail(wb, info, out_path=None):
    """東レ宛の月次報告メール下書きシート。
    シート上の表現は添付の実メール体裁に合わせる (挨拶→表→件数サマリ→締め)。
    国分結果では湖南の最新検証結果Excelから数字を取り込み、同内容を .html にも保存する。
    """
    ws = wb.create_sheet('報告メール下書き')
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '70AD47'
    for col, wd in zip('ABCDEF', (22, 18, 14, 14, 14, 8)):
        ws.column_dimensions[col].width = wd

    own = _mail_figures(info)
    peer = None
    peer_note = ''
    if '国分' in info['工場']:
        base = Path(out_path).parent if out_path else Path('.')
        peer_path = _find_peer_result(base, info['工場'])
        if peer_path:
            peer = _load_mail_figures_from_result(peer_path)
            peer_note = (f'湖南工場の数字は {peer_path.name} から取り込み' if peer
                         else f'湖南工場の検証結果 {peer_path.name} を読めませんでした')
        else:
            peer_note = '湖南工場の検証結果Excel (検証結果_湖南工場_*.xlsx) が見つかりません'

    rows = []
    if '国分' in own['工場']:
        rows.append(own)
        rows.append(peer if peer else {
            '工場': '湖南工場', 'short': '湖南', 'basho': 'A010P',
            't1': 0, 't2': 0, 'adj_disp': 0, 'tougetsu_disp': 0, 'tougetsu': 0, 'n': 0, 'source': None,
        })
    else:
        rows.append(own)

    html = _build_mail_html(info, rows)
    html_name = (Path(out_path).stem + '.html') if out_path else '報告メール.html'

    # ---- 案内 (本文の外) ----
    r = 1
    _merge_put(ws, r, 1, 5,
               f'【案内】①下のメール本文を目視確認　②シート下部の「HTMLソース」をコピー、または同フォルダの「{html_name}」をブラウザで開き Ctrl+A→C→Outlookへ貼り付け。数字は自動反映済み。',
               F(size=9, color='C00000'))
    r += 1
    if peer_note:
        _merge_put(ws, r, 1, 5, f'※{peer_note}', F(size=9, color='595959'))
        r += 1
    r += 1
    _merge_put(ws, r, 1, 5, '▼▼ ここから下がメール本文 (添付の実メール体裁) ▼▼',
               F(bold=True, size=10, color='375623'), fill=FILL_OK)
    r += 1
    body_start = r

    # ---- メール本文: 挨拶 ----
    for line in (
        '東レ株式会社',
        'トーレペフ事業部 御中',
        '',
        'いつも大変お世話になっております。',
        '長岡産業の難波です。',
        '',
        '掲題の件、下記にご報告申し上げます。',
        '',
    ):
        _merge_put(ws, r, 1, 5, line, F(size=11))
        r += 1

    # ---- 表 (添付と同じ列) ----
    headers = ['拠点', '東レ支払いデータ', '実売上', '前月調整分', '当月調整分']
    header_fill = PatternFill('solid', fgColor='BDD7EE')  # 添付の薄い青ヘッダー
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = F(bold=True, size=10)
        cell.fill = header_fill
        cell.alignment = CENTER
    table_header_r = r
    r += 1
    data_start = r
    for row in rows:
        label = f"{row['short']} ({row['basho']})" if row['basho'] else row['short']
        ws.cell(row=r, column=1, value=label).alignment = CENTER
        ws.cell(row=r, column=1).font = F(size=10)
        for c, v in enumerate((row['t1'], row['t2'], row['adj_disp'], row['tougetsu_disp']), 2):
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
            cell.font = F(name='BIZ UDGothic', size=10)
        r += 1
    ws.cell(row=r, column=1, value='合計').font = F(bold=True, size=10)
    ws.cell(row=r, column=1).alignment = CENTER
    for c in range(2, 6):
        cell = ws.cell(row=r, column=c,
                       value=f'=SUM({get_column_letter(c)}{data_start}:{get_column_letter(c)}{r - 1})')
        cell.number_format = '#,##0'
        cell.font = F(bold=True, name='BIZ UDGothic', size=10)
        cell.alignment = Alignment(horizontal='right')
    _draw_box(ws, table_header_r, 1, r, 5, color='8FAADC', style='thin')
    r += 2

    # ---- 件数サマリ (添付と同じ文言) ----
    total_t1 = sum(x['t1'] for x in rows)
    total_tg = sum(x['tougetsu_disp'] for x in rows)
    total_n = sum(x['n'] for x in rows)
    _merge_put(ws, r, 1, 5, f'合計（税抜き）　{_fmt_yen(total_t1)}円の御社お支払いデータに対し', F(size=11))
    r += 1
    for row in rows:
        _merge_put(ws, r, 1, 5,
                   f'　{row["工場"]}　　　 {row["n"]}件（{_fmt_paren_yen(row["tougetsu_disp"])}）',
                   F(size=11))
        r += 1
    highlight = f'{total_n}件の差異（{_fmt_yen(abs(total_tg))}円）過不足'
    # 添付メールと同じ文言。強調は太字+下線（HTML側は highlight 部分のみ）
    _merge_put(ws, r, 1, 5, f'合計　{highlight}がありました。',
               F(bold=True, size=11, underline='single'))
    r += 1
    for line in _prior_mail_note_lines(info):
        _merge_put(ws, r, 1, 5, line, F(size=10, color='1F4E79'))
        r += 1
    r += 1
    for line in (
        f'添付ファイルをご確認いただき、次月（{info["翌月ラベル"]}）での',
        'ご調整をよろしくお願い申し上げます。',
        '今後とも何卒よろしくお願い申し上げます。',
        '',
        '以上',
    ):
        _merge_put(ws, r, 1, 5, line, F(size=11))
        r += 1
    body_end = r - 1
    _draw_box(ws, body_start, 1, body_end, 5, color='70AD47', style='thin')

    # ---- HTMLソース (同内容を .html にも保存。シート内でもコピー可能) ----
    r += 2
    _merge_put(ws, r, 1, 5,
               f'▼▼ HTMLソース (「{html_name}」と同一。このセルをコピーして利用可) ▼▼',
               F(bold=True, size=10, color='1F4E79'), fill=PatternFill('solid', fgColor='DDEBF7'))
    r += 1
    # Excelセル上限 32767 文字。通常の報告メールHTMLは十分収まる
    html_cell = _merge_put(ws, r, 1, 5, html, F(name='Consolas', size=8),
                           align=Alignment(wrap_text=True, vertical='top'))
    # 行高: おおむね行数×12pt (上限 409)
    n_lines = html.count('\n') + 1
    ws.row_dimensions[r].height = min(409, max(60, n_lines * 11))
    _draw_box(ws, r, 1, r, 5, color='8FAADC', style='thin')

    html_path = None
    if out_path:
        html_path = Path(out_path).with_suffix('.html')
        html_path.write_text(html, encoding='utf-8')
    info['メールHTML'] = str(html_path) if html_path else None
    info['メール他工場'] = peer_note
    return ws


def _csv_cell(v):
    """CSVセルを数値化できれば数値に、できなければ文字列のまま返す"""
    s = v.strip()
    if re.fullmatch(r'-?\d+(\.\d+)?', s):
        f = float(s)
        return int(f) if f == int(f) else f
    return v


def _merge_put(ws, r, c1, c2, value='', font=None, fill=None, align=None, fmt=None):
    """セル結合して値・書式を設定 (fillは結合範囲全体に適用)"""
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(row=r, column=c1)
    cell.value = value
    if font:
        cell.font = font
    if fill:
        for cc in range(c1, c2 + 1):
            ws.cell(row=r, column=cc).fill = fill
    if align:
        cell.alignment = align
    if fmt:
        cell.number_format = fmt
    return cell


def _draw_box(ws, r1, c1, r2, c2, fill=None, color='4472C4', style='medium'):
    """範囲に外枠罫線(+塗り)を引く。結合セルにも効くよう構成セル全部に設定"""
    side = Side(style=style, color=color)
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            cell = ws.cell(row=rr, column=cc)
            if fill:
                cell.fill = fill
            cur = cell.border
            cell.border = Border(
                top=side if rr == r1 else cur.top,
                bottom=side if rr == r2 else cur.bottom,
                left=side if cc == c1 else cur.left,
                right=side if cc == c2 else cur.right)


CENTER = Alignment(horizontal='center', vertical='center')
CENTER_WRAP = Alignment(horizontal='center', vertical='center', wrap_text=True)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')

LEGEND = (  # (表示名, 塗り, 文字色, 意味, とるべき対応)
    ('不一致', FILL_MISMATCH, '9C0006', '両方に存在するが金額が異なる', '最優先で原因確認。差額計が東レへ報告する「当月差異」'),
    ('翌月記載', FILL_NEXT, '833C00', '①当月検収なのに②は翌月に記載(記載月誤り)', '②2ファイルを修正(当月へ追記・翌月から削除)して再実行'),
    ('①②③のみ', FILL_ONLY, '7F5F00', '片側にしか存在しない', '記載漏れ・計上月ずれを調査(①のみは②への記載漏れ疑い)'),
    ('前月調整', FILL_PREV, '1F4E79', '①のみだが過去月の②に存在(過去依頼の金額修正)', '判明済み。備考の当時金額と差額を参考確認のみ'),
    ('前月過不足', FILL_PRIOR, '1F4E79',
     '前月の工賃差検証で次月修正依頼済みだが当月①に未反映。①列は空、②列=プラス差(長岡−東レ)。マイナス側は①取消行で「前月調整」に出るがプラス側は①に現れない',
     '報告計上額=前月の①−②。東レへ再依頼・反映確認。前月過不足.csv で管理'),
    ('①正(手動)\n②正(手動)', FILL_MANUAL, '5B2C6F', '手動判定.csv で当月・当該契約NOにつき正しい側を指定した不一致',
     '判明済み。①正=②を①金額に修正(東レへは報告しない) / ②正=東レへ報告(報告計上額に含む)'),
    ('月ずれ解消\n枝番統合一致', FILL_RESOLVED, '375623', '検証Bで②③の複数月累計が一致(計上月ずれ)、または③の枝番(例 C8-3-1)を親依頼NOに合算すると②と一致', '対応不要。飛ばしてよい'),
    ('形式不正', FILL_INFO, '595959', '②のC列が契約NO形式でないのに金額がある行', '記入漏れ疑い。②原本の該当行を確認'),
)


def _write_summary(wb, info):
    ws = wb.active
    ws.title = 'サマリ'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '1F4E79'
    for col, wd in zip('ABCDEFGH', (2, 30, 24, 24, 24, 24, 24, 2)):
        ws.column_dimensions[col].width = wd

    name2, amt2 = info['②名称'], info['②金額列']
    # タイトル
    _merge_put(ws, 1, 2, 7, f'後加工工賃 金額整合性検証レポート【{info["工場"]}】', F(bold=True, size=15, color='1F4E79'))
    _merge_put(ws, 2, 2, 7,
               f"対象月: {info['対象月ラベル']}　実行: {info['実行日時']} ({info['実行者']})　確認者:＿＿＿＿＿ 確認日:＿＿＿＿＿",
               F(size=10, color='595959'))
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[3].height = 6

    # ---- KPIカード (4行目〜6行目) ----
    n_req_a = info['A不一致'] + info['A翌月記載件数'] + info['A前月過不足件数'] + info['A①のみ'] + info['A②のみ']
    n_req_b = info['B不一致'] + info['B②のみ'] + info['B③のみ']
    n_manual = info['A手動①正件数'] + info['A手動②正件数']
    n_known = info['A前月調整件数'] + info['A翌月記載件数'] + info['A前月過不足件数'] + n_manual + info['B月ずれ解消'] + info['B枝番統合']
    tougetsu = round(info['報告当月差異'] + info['A翌月記載額'] + info['A前月過不足額']
                     + info['報告①のみ額'] - info['報告②のみ額'])
    n_warn = len(info['警告'])
    zansa = info['報告残差']
    cards = [  # (ラベル, 値, 補足, 状態)  状態: bad/ok/neutral/known
        ('検証A 要確認 (①vs②)', f'{n_req_a} 件',
         f"不一致{info['A不一致']}・翌月記載{info['A翌月記載件数']}・前月過不足{info['A前月過不足件数']}・片側のみ{info['A①のみ']+info['A②のみ']}",
         'bad' if n_req_a else 'ok'),
        ('検証B 要確認 (②vs③)', f'{n_req_b} 件',
         f"不一致{info['B不一致']}・片側のみ{info['B②のみ']+info['B③のみ']}",
         'bad' if n_req_b else 'ok'),
        ('報告する過不足(当月)', f'{tougetsu:+,} 円', '検証A「報告計上額」列の合計', 'neutral'),
        ('判明済み', f'{n_known} 件', '前月調整・翌月記載・前月過不足・手動判定・月ずれ解消・枝番統合', 'known'),
        ('警告', f'{n_warn} 件', 'データ品質の注意', 'bad' if n_warn else 'ok'),
        ('検算残差', f'{zansa:,} 円', '0なら内訳の整合OK', 'bad' if zansa else 'ok'),
    ]
    palette = {'bad': ('FDE9E9', '9C0006'), 'ok': ('C6EFCE', '375623'),
               'neutral': ('DDEBF7', '1F4E79'), 'known': ('E2EFDA', '375623')}
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 14
    for i, (label, value, sub, state) in enumerate(cards):
        col = 2 + i
        bg, fg = palette[state]
        fill = PatternFill('solid', fgColor=bg)
        c = ws.cell(row=4, column=col, value=label)
        c.font, c.alignment = F(size=10, bold=True, color=fg), CENTER
        c = ws.cell(row=5, column=col, value=value)
        c.font, c.alignment = F(name=FONT_NUM, size=17, bold=True, color=fg), CENTER
        c = ws.cell(row=6, column=col, value=sub)
        c.font, c.alignment = F(size=8.5, color=fg), CENTER
        _draw_box(ws, 4, col, 6, col, fill=fill, color=fg, style='thin')
    ws.row_dimensions[7].height = 10

    # ---- 詳細リスト ----
    r = 8

    def section(title):
        nonlocal r
        cell = _merge_put(ws, r, 2, 7, title, F(bold=True, size=12, color='1F4E79'))
        for cc in range(2, 8):
            cur = ws.cell(row=r, column=cc).border
            ws.cell(row=r, column=cc).border = Border(bottom=Side(style='medium', color='4472C4'),
                                                      top=cur.top, left=cur.left, right=cur.right)
        ws.row_dimensions[r].height = 20
        cell.alignment = Alignment(vertical='center')
        r += 1

    def put(label, value, fmt='#,##0', bold=False):
        nonlocal r
        ws.cell(row=r, column=2, value=label).font = F(bold=bold)
        cell = _merge_put(ws, r, 3, 7, value, font=F(name=FONT_NUM, bold=bold),
                          align=Alignment(horizontal='left'))
        if isinstance(value, (int, float)):
            cell.number_format = fmt
        r += 1

    def blank(h=8):
        nonlocal r
        ws.row_dimensions[r].height = h
        r += 1

    def note(text):
        nonlocal r
        _merge_put(ws, r, 2, 7, text, F(size=9, color='595959'))
        r += 1

    section('総額')
    put(f'① 東レ検収 (お支払データ・{info["入庫場所"]})', info['①総額'], '#,##0"円"')
    put(f'② {name2} {amt2}合計', info['②総額'], '#,##0"円"')
    put('③ アラジン 加工金額合計' + (f' (得意先{info["③得意先"]})' if info['③得意先'] else ''), info['③総額'], '#,##0"円"')
    blank()

    section('検証A　① vs ② (契約NO突合)')
    put('共通契約NO', info['A共通'], '#,##0"件"')
    n_manual_a = info['A手動①正件数'] + info['A手動②正件数']
    put('　一致 / 不一致(要確認)', f"{info['A共通'] - info['A不一致'] - n_manual_a} / {info['A不一致']}")
    if n_manual_a:
        put('　手動判定済みの不一致 (対象外)',
            f"{n_manual_a}件 (①正 {info['A手動①正件数']}件 {info['A手動①正額']:+,}円 → ②要修正 / "
            f"②正 {info['A手動②正件数']}件 {info['A手動②正額']:+,}円 → 東レへ報告)")
    put('前月調整 (①のみ・過去月②に存在)', f"{info['A前月調整件数']}件 / {info['A前月調整額']:,.0f}円")
    if info['A前月過不足件数']:
        put('前月過不足 (前月差・当月①未反映)',
            f"{info['A前月過不足件数']}件 / {info['A前月過不足額']:+,}円 ※報告計上・①列空/②=プラス差")
        details = info.get('A前月過不足明細') or []
        if details:
            note('　　' + ' / '.join(
                f'{d["依頼表示"]} {d["契約NO"]} プラス差{d["プラス差"]:+,}円'
                for d in details))
    put('翌月記載 (①のみ・翌月②に存在 → ②要修正)', f"{info['A翌月記載件数']}件 / {info['A翌月記載額']:,.0f}円")
    put('①のみ / ②のみ (金額あり)', f"{info['A①のみ']} / {info['A②のみ']}")
    blank()

    section('検証B　② vs ③ (依頼NO突合)')
    put('共通依頼NO', info['B共通'], '#,##0"件"')
    put('　一致 / 不一致(要確認)', f"{info['B共通'] - info['B不一致'] - info['B月ずれ解消共通'] - info['B枝番統合共通']} / {info['B不一致']}")
    put('月ずれ解消 (複数月累計一致・対象外)', f"{info['B月ずれ解消']}件 (②のみ・③のみ由来を含む)")
    put('枝番統合一致 (③枝番を親に合算して一致・対象外)', f"{info['B枝番統合']}件 (親・枝番の両方を数える)")
    put('②のみ / ③のみ (金額あり)', f"{info['B②のみ']} / {info['B③のみ']}")
    blank()

    # ---- 検証C (湖南のみ): 月次処理ファイルの東レ合計との照合表 ----
    if info.get('検証C'):
        section(f'検証C　月次処理ファイル 東レ合計 vs 本検証の総額')
        note(f'　{info["月次ファイル"]}　※契約NO等の明細が無いため東レ部分の合計金額のみ照合')
        c_judge_color = {'一致': '375623', '内訳で説明可': '1F4E79', '不一致': '9C0006', '要確認': '9C0006', '読取不可': '833C00'}
        c_judge_fill = {'一致': FILL_OK, '内訳で説明可': PatternFill('solid', fgColor='DDEBF7'),
                        '不一致': FILL_MISMATCH, '要確認': FILL_MISMATCH, '読取不可': FILL_NEXT}
        for j, h in enumerate(('項目', '月次処理ﾌｧｲﾙ', '本検証', '差(本検証-月次)', '判定', '備考'), start=2):
            c = ws.cell(row=r, column=j, value=h)
            c.font, c.fill, c.alignment = F(bold=True, color='FFFFFF'), FILL_HEADER, CENTER
        r += 1
        long_notes = []
        for key, mv, ours, d, judge, note_txt in info['検証C']:
            # 備考は短い比較対象だけをセルに入れ、長い説明(内訳式など)は表の下に注記として出す
            short, _, rest = note_txt.partition('。')
            if rest:
                long_notes.append(f'　{key}: {rest}')
            ws.cell(row=r, column=2, value=key).font = F()
            for j, v in ((3, mv), (4, ours), (5, d)):
                c = ws.cell(row=r, column=j, value=v if v is not None else '(読取不可)')
                c.font = F(name=FONT_NUM)
                if isinstance(v, (int, float)):
                    c.number_format = '#,##0.00;[Red]-#,##0.00'
                c.alignment = Alignment(horizontal='right')
            c = ws.cell(row=r, column=6, value=judge)
            c.font, c.alignment = F(bold=True, color=c_judge_color.get(judge, '000000')), CENTER
            c.fill = c_judge_fill.get(judge, FILL_INFO)
            c = ws.cell(row=r, column=7, value=short)
            c.font, c.alignment = F(size=9), Alignment(vertical='center')
            _draw_box(ws, r, 2, r, 7, color='BFBFBF', style='thin')
            r += 1
        for txt in long_notes:
            note(txt)
        blank()

    # ---- 検証D (国分のみ): ②「東レまとめ」内部整合性 ----
    chk_d = info.get('検証D')
    if chk_d:
        section(f'検証D　②「{MATOME_SHEET}」の内部整合性 (元4シートとの参照・式・金額突合)')
        if chk_d['skipped']:
            note(f'　スキップ: {chk_d["skipped"]}')
        else:
            counts = Counter(row[6] for row in chk_d['rows'])
            n_err = counts['金額差'] + counts['参照ずれ'] + counts['未取込']
            put('まとめ取込行数', chk_d['n_map'], '#,##0"行"')
            put('要修正 (金額差 / 参照ずれ / 未取込)',
                f"{n_err}件 ({counts['金額差']} / {counts['参照ずれ']} / {counts['未取込']})", bold=n_err > 0)
            put('注意 (式異常 / 直接入力)', f"{counts['式異常'] + counts['直接入力']}件 ({counts['式異常']} / {counts['直接入力']})")
            for j, h in enumerate(('元シート', 'まとめAA計', '元シートAA計', '差(まとめ-元)', '判定', ''), start=2):
                c = ws.cell(row=r, column=j, value=h)
                c.font, c.fill, c.alignment = F(bold=True, color='FFFFFF'), FILL_HEADER, CENTER
            r += 1
            tot_m = tot_s = 0
            for s, sm, ss, d in chk_d['totals'] + [('合計', sum(t[1] for t in chk_d['totals']),
                                                    sum(t[2] for t in chk_d['totals']), None)]:
                d = sm - ss if d is None else d
                ok = abs(d) <= info['許容差']
                ws.cell(row=r, column=2, value=s).font = F(bold=(s == '合計'))
                for j, v in ((3, sm), (4, ss), (5, d)):
                    c = ws.cell(row=r, column=j, value=round(v))
                    c.font, c.alignment = F(name=FONT_NUM, bold=(s == '合計')), Alignment(horizontal='right')
                    c.number_format = '#,##0;[Red]-#,##0'
                c = ws.cell(row=r, column=6, value='一致' if ok else '不一致')
                c.font, c.alignment = F(bold=True, color='375623' if ok else '9C0006'), CENTER
                c.fill = FILL_OK if ok else FILL_MISMATCH
                _draw_box(ws, r, 2, r, 7, color='BFBFBF', style='thin')
                r += 1
            note('　金額差・参照ずれ・未取込 = ②の集計自体が誤っているため、②を直して再実行する。'
                 '直接入力 = 元シートの単価欄が式でなく値で上書きされておりまとめの再計算と食い違う恐れ')
        blank()

    section('報告用内訳　①お支払データの分解')
    put('① 東レお支払データ', info['①総額'], '#,##0"円"', bold=True)
    put('　② 当月実売上 (契約NOベース)', info['報告②実売上'], '#,##0"円"')
    put('　前月調整分', round(info['A前月調整額']), '#,##0"円"')
    put('　翌月記載 (②修正後は実売上へ)', round(info['A翌月記載額']), '#,##0"円"')
    put('　当月差異 (不一致の差額計)' + (' ※手動②正を含む' if info['A手動②正件数'] else ''), info['報告当月差異'], '#,##0"円"')
    if info['A手動①正件数']:
        put('　手動判定①正 (②要修正・東レへ報告しない)', info['A手動①正額'], '#,##0"円"')
    if info['A前月過不足件数']:
        put('　前月過不足 (①未反映・報告計上)', info['A前月過不足額'], '#,##0"円"')
    put('　未解明 ①のみ / ②のみ(符号逆)', f"{info['報告①のみ額']:,.0f} / {-info['報告②のみ額']:,.0f}")
    put('　検算残差', info['報告残差'], '#,##0"円"', bold=True)
    blank(4)
    put('報告する過不足 (当月)', tougetsu, '#,##0"円"', bold=True)
    note(f"　　= 当月差異 {info['報告当月差異']:,} + 翌月記載 {round(info['A翌月記載額']):,} "
         + (f"+ 前月過不足 {info['A前月過不足額']:,} " if info['A前月過不足件数'] else '')
         + f"+ ①のみ {round(info['報告①のみ額']):,} − ②のみ {round(info['報告②のみ額']):,}")
    note('　　= 検証Aシート「報告計上額」列の合計 (不一致行=差額①-②、翌月記載・①のみ行=①金額、②のみ行=−②金額'
         + ('、前月過不足行=前月の差額①-②' if info['A前月過不足件数'] else '')
         + ')。メールで東レへ報告する過不足の金額')
    if info['A手動①正件数']:
        note(f'　　※手動判定①正 {info["A手動①正件数"]}件 ({info["A手動①正額"]:+,}円) は①が正しいため報告に含めない。'
             f'②{info["②名称"]}を①金額に修正すること')
    if info['A前月過不足件数']:
        note(f'　　※前月過不足 {info["A前月過不足件数"]}件 ({info["A前月過不足額"]:+,}円) は前月の工賃差検証で次月修正依頼済みだが'
             f'当月①CSVに未反映のため報告計上に残している ({PRIOR_FILE})。'
             f'プラス差は①データに含まれないため検証Aでは①列を空・②列にプラス差(長岡−東レ)を表示')
    blank()

    # ---- 警告 (赤帯ボックス) ----
    if info['警告']:
        top = r
        _merge_put(ws, r, 2, 7, f'⚠ 警告　データ品質の注意 ({len(info["警告"])}件)',
                   F(bold=True, size=12, color='9C0006'), fill=FILL_WARNBOX)
        ws.row_dimensions[r].height = 20
        r += 1
        for wmsg in info['警告']:
            cell = _merge_put(ws, r, 2, 7, '・' + wmsg, F(size=10, color='9C0006'),
                              fill=FILL_WARNBOX, align=WRAP_TOP)
            ws.row_dimensions[r].height = 14.5 * max(1, -(-len(wmsg) // 75)) + 3
            r += 1
        _draw_box(ws, top, 2, r - 1, 7, color='9C0006', style='thin')
        blank()

    # ---- 凡例 ----
    section('凡例　明細シートの色分け')
    for judge, fill, fg, desc, action in LEGEND:
        cell = ws.cell(row=r, column=2, value=judge)
        cell.fill = fill
        cell.font = F(bold=True, color=fg)
        cell.alignment = CENTER_WRAP
        _draw_box(ws, r, 2, r, 2, color='BFBFBF', style='thin')
        _merge_put(ws, r, 3, 7, f'{desc}　→ {action}', F(size=10), align=Alignment(vertical='center', wrap_text=True))
        ws.row_dimensions[r].height = 28 if ('\n' in judge or len(desc) + len(action) > 60) else 15
        r += 1
    blank()

    # ---- 使用ファイル ----
    section('使用ファイル・実行条件')
    put('対象工場', info['工場'], None)
    put('① 東レCSV', f"{info['①ファイル']} (入庫場所 {info['入庫場所']} のみ)", None)
    put(f'② {name2}', f"{info['②ファイル']} ({info['②説明']})", None)
    put('　過去月明細', info['前月ファイル'], None)
    put('　翌月明細', info['翌月ファイル'], None)
    if info.get('②フォルダ'):
        put('　②参照フォルダ', info['②フォルダ'], None)
    put('③ アラジン', f"{info['③ファイル']} (対象年月: {info['対象年月']}"
        + (f" / 得意先 {info['③得意先']} のみ" if info['③得意先'] else '') + ')', None)
    if info.get('③フォルダ'):
        put('　③参照フォルダ', info['③フォルダ'], None)
    if info.get('月次ファイル'):
        put('月次処理ファイル', f"{info['月次ファイル']} (東レ合計のみ照合)", None)
        if info.get('月次フォルダ'):
            put('　月次処理ファイル参照フォルダ', info['月次フォルダ'], None)
    put('許容差', f"±{info['許容差']}円", None)
    put('①の対象外入庫場所', info['対象外入庫場所'] or 'なし', None)
    put('①の取消行(H列「-」)', info['①取消行'], None)
    put('金額0円の無効行(出力対象外)', info['0円除外'], None)
    if info['A手動①正件数'] + info['A手動②正件数']:
        put('手動判定', f"{MANUAL_FILE} の {info['対象月ラベル']}・{info['工場']} 行を適用 "
                     f"(①正 {info['A手動①正件数']}件 / ②正 {info['A手動②正件数']}件)", None)

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def write_excel(out_path, info, records_a, records_b, raw_sheets=()):
    wb = Workbook()
    wb.loaded_theme = BIZ_THEME

    # ---- サマリシート ----
    _write_summary(wb, info)

    # ---- レポートの見方 ----
    _write_guide(wb, info)

    # ---- 明細シート ----
    col2 = f'②{info["②名称"]}金額'
    ws_a = _write_sheet(wb, '検証A_契約NO(①vs②)',
                        ['契約NO', '依頼NO', '①東レ金額', col2, '差額(①-②)', '報告計上額', '判定', '備考'], records_a)
    ws_a.sheet_properties.tabColor = 'C00000'
    ws_b = _write_sheet(wb, '検証B_依頼NO(②vs③)',
                        ['依頼NO', col2, '③アラジン金額', '差額(②-③)', '判定', '備考'], records_b)
    ws_b.sheet_properties.tabColor = 'ED7D31'

    # ---- 検証D (国分): ②まとめ整合性の明細 ----
    chk_d = info.get('検証D')
    if chk_d and not chk_d['skipped']:
        rows_d = [tuple(None if v == '' else v for v in row) for row in chk_d['rows']]
        rows_d.sort(key=lambda x: _judge_rank(x[6], x[5], x[3:5]))
        if not rows_d:
            rows_d = [(f'{MATOME_SHEET} 全{chk_d["n_map"]}行', None, None, None, None, None, '問題なし',
                       'まとめの参照式・元シートの式・行別金額がすべて整合しています')]
        ws_d = _write_sheet(wb, '検証D_②まとめ整合性',
                            ['場所 (シート!行)', '依頼NO', '契約NO', 'まとめAA(加工賃)', '元シートAA(加工賃)',
                             '差額(まとめ-元)', '判定', '内容・対処'], rows_d)
        ws_d.column_dimensions['A'].width = 30
        ws_d.sheet_properties.tabColor = '7030A0'

    # ---- 報告メール下書き ----
    _write_mail(wb, info, out_path)

    # ---- 元データのコピーシート ----
    for title, rows, freeze, highlight in raw_sheets:
        _append_raw_sheet(wb, title, rows, freeze, highlight)

    wb.save(out_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description='後加工工賃 3ソース金額整合性検証')
    ap.add_argument('--factory', choices=('kokubu', 'konan'), default='kokubu',
                    help='対象工場 kokubu=国分工場(既定) / konan=湖南工場。②のフォーマットとフォルダが切り替わる')
    ap.add_argument('--base', default=str(DEFAULT_BASE), help='●自動検証フォルダのパス')
    ap.add_argument('--toray', help='① 東レCSVのパス')
    ap.add_argument('--nagaoka', help='② 長岡明細xlsx(国分) / 加工賃試算xlsm(湖南) のパス')
    ap.add_argument('--prev', help='②の前月ファイルのパス (①のみ契約NOの前月調整判定に使用)')
    ap.add_argument('--aladdin', help='③ アラジンxlsxのパス')
    ap.add_argument('--tol', type=float, default=0.5, help='一致とみなす許容差(円) 既定0.5 (1円差も不一致として検出)')
    ap.add_argument('--basho', help='①の対象入庫場所。既定は工場に応じて A010(国分) / A010P(湖南)')
    ap.add_argument('--force', action='store_true', help='①②③の対象月不一致でも警告のみで続行する')
    ap.add_argument('--out', help='結果Excelの出力先 (既定: <base>\\検証結果_<工場>_タイムスタンプ.xlsx)')
    args = ap.parse_args(argv)
    fac = _factories()[args.factory]
    args.basho = norm(args.basho or fac['basho'])
    name2, amt2 = fac['name2'], fac['amt2']

    base = Path(args.base)
    all_warnings = []

    # ---- ファイル選択と対象月の整合チェック (①のRVSHEETyyyymmを基準月とする) ----
    p1 = Path(args.toray) if args.toray else find_toray_csv(base / '東レ送付CSV')
    target_ym = target_ym_of_csv(p1)
    if target_ym is None:
        all_warnings.append(f'①ファイル名から対象年月(RVSHEETyyyymm)を特定できず、月整合チェックをスキップしました: {p1.name}')

    dir2 = resolve_dir(base, fac['dir2'])
    if not dir2.is_dir():
        sys.exit(f'②{name2}のフォルダにアクセスできません: {dir2}')

    def disp2(p: Path) -> str:
        """②ファイルの表示名。参照フォルダ配下ならサブフォルダ付き (湖南: 年度フォルダ/作業中ファイルの区別のため)"""
        try:
            return str(p.relative_to(dir2))
        except ValueError:
            return p.name
    if args.nagaoka:
        p2 = Path(args.nagaoka)
        p2_ym = fac['ym2'](p2)
        if target_ym and p2_ym and p2_ym != target_ym:
            all_warnings.append(f'②の月度({p2_ym[0]}年{p2_ym[1]}月度)が①の対象月({target_ym[0]}年{target_ym[1]}月)と不一致です(明示指定のため続行)')
        # 過去月 = ②の年月より古いもののみ (未来ファイルの混入を防ぐ)
        _, prevs, nexts = fac['find2'](dir2, p2_ym or target_ym)
    else:
        p2, prevs, nexts = fac['find2'](dir2, target_ym)
    if args.prev:
        pf = Path(args.prev)
        pym = fac['ym2'](pf)
        prevs = [(f'{pym[0]}年{pym[1]}月度' if pym else '前月', pf)]

    dir3 = resolve_dir(base, fac['dir3'])
    if not dir3.is_dir():
        sys.exit(f'③アラジンのフォルダにアクセスできません: {dir3}')
    p3 = Path(args.aladdin) if args.aladdin else find_aladdin(dir3, target_ym)

    toray, toray_by_basho, nyuko_dates, w1, ct_errors, minus_rows = read_toray_csv(p1, args.basho)
    nag_irai, nag_keiyaku, keiyaku_to_irai, bad_keiyaku, w2 = fac['read2'](p2)
    all_warnings += w1 + w2

    # ---- 国分: ②「東レまとめ」の内部整合性 (検証D)。まとめが元4シートとずれていれば②総額そのものが疑わしい ----
    check_d = None
    if fac.get('check_d'):
        check_d = fac['check_d'](p2, args.tol)
        if check_d['skipped']:
            all_warnings.append(f'検証D(②まとめ整合性)をスキップしました: {check_d["skipped"]}')
        else:
            n_err = sum(1 for row in check_d['rows'] if row[6] in ('金額差', '参照ずれ', '未取込'))
            if n_err:
                gap = sum(d for _, _, _, d in check_d['totals'])
                all_warnings.append(f'②「東レまとめ」が元4シートとずれています(検証D 要修正 {n_err}件 / 加工賃合計の差 {gap:+,.0f}円)。'
                                    f'②総額の信頼性に関わるため②を修正して再実行すること。詳細は検証Dシート')

    # ---- 湖南: 月次処理ファイル(契約NO明細なし)の東レ合計を読む → 検証Cで総額照合 ----
    monthly_path, monthly = None, {}
    if fac['dir_monthly']:
        monthly_path = find_monthly_file(resolve_dir(base, fac['dir_monthly']), target_ym)
        if monthly_path:
            try:
                monthly = read_monthly_file(monthly_path)
            except Exception as e:  # noqa: BLE001 - 読めなくても本検証は続行
                all_warnings.append(f'月次処理ファイル {monthly_path.name} を読めませんでした: {e}')
        else:
            all_warnings.append(f'{fac["label"]}の月次処理ファイル(対象月 {target_ym})が見つからないため検証Cをスキップしました')
    # ②の依頼NO別金額を年月別に保持 (検証Bの累計比較用)
    nag_irai_by_ym = {}
    if target_ym:
        nag_irai_by_ym[target_ym] = nag_irai

    def label_ym(label):
        m = GATSUDO_PAT.search(label)
        return (int(m.group(1)), int(m.group(2))) if m else None

    # 過去月ファイルを新しい月から順に読み込む (①のみ契約NOの過去月調整判定に使用)
    prev_data = []
    for label, pf in prevs:
        try:
            bi, pk, pim, _, _ = fac['read2'](pf)
            prev_data.append((label, pk, pim))
            if (ym := label_ym(label)):
                nag_irai_by_ym[ym] = bi
        except SystemExit as e:
            all_warnings.append(f'過去月明細 {pf.name} を読めないためスキップしました: {e}')
    # 翌月以降ファイルを近い月から順に読み込む (①のみ契約NOが誤って翌月に記載されていないかの判定に使用)
    next_data = []
    for label, nf in nexts:
        try:
            bi, nk, nim, _, _ = fac['read2'](nf)
            next_data.append((label, nk, nim))
            if (ym := label_ym(label)):
                nag_irai_by_ym[ym] = bi
        except SystemExit as e:
            all_warnings.append(f'翌月明細 {nf.name} を読めないためスキップしました: {e}')
    aladdin, taisho, taisho_ym = read_aladdin(p3, fac['customer3'])
    # 当月照会の取得日が対象月の月中なら実績が途中まで → ②のみ・検証C不一致が大量に出るので先頭で警告
    mfn3 = re.search(r'_(\d{4})(\d{2})(\d{2})_', p3.name)
    if mfn3 and taisho_ym and (int(mfn3.group(1)), int(mfn3.group(2))) == tuple(taisho_ym) and int(mfn3.group(3)) < 28:
        all_warnings.insert(0, f'③当月照会({p3.name})は{int(mfn3.group(2))}月{int(mfn3.group(3))}日時点の途中実績です。'
                               f'月末確定後に取得したファイルを {dir3} に置いて再実行してください'
                               f'(同一対象年月が複数あればファイル名の日時が最新のものを採用します)。'
                               f'途中実績のままでは②のみ・検証C不一致が実態より多く出ます')
    # ③の他月照会ファイルを読み込む (検証Bの累計比較で計上月ずれを自動消込するために使用)
    aladdin_by_ym = {}
    if target_ym:
        aladdin_by_ym[target_ym] = aladdin
        try:
            amap_all = find_aladdin_all(dir3)
        except SystemExit:
            amap_all = {}
        aladdin_other_files = []
        for ym, af in sorted(amap_all.items()):
            if ym == target_ym or (args.aladdin and af == p3):
                continue
            try:
                aladdin_by_ym[ym] = read_aladdin(af, fac['customer3'])[0]
                aladdin_other_files.append(af.name)
            except SystemExit as e:
                all_warnings.append(f'③他月照会 {af.name} を読めないためスキップしました: {e}')
                continue
            # 照会取得日が対象年月の月中の場合、実績が途中までしか入っていない
            mfn = re.search(r'_(\d{4})(\d{2})(\d{2})_', af.name)
            if mfn and (int(mfn.group(1)), int(mfn.group(2))) == ym and int(mfn.group(3)) < 28:
                all_warnings.append(f'③{ym[1]}月照会({af.name})は{int(mfn.group(3))}日時点の途中実績です。'
                                    f'月末確定後に再取得して再実行すると、検証Bの月ずれ解消(累計一致)の精度が上がります')
    else:
        aladdin_other_files = []

    # ③の対象年月が①の対象月と一致するか検証
    if target_ym and taisho_ym and taisho_ym != target_ym:
        msg = (f'③の対象年月({taisho_ym[0]}年{taisho_ym[1]}月)が①の対象月({target_ym[0]}年{target_ym[1]}月)と不一致です。'
               f'アラジンの照会条件を確認してください')
        if not args.force:
            sys.exit(f'中断: {msg}\n  それでも実行する場合は --force を付けてください')
        all_warnings.append(msg + ' (--force指定のため続行)')

    def minus_note(k):
        """契約NOに①の取消行(H列「-」)があれば備考用に返す"""
        rows_m = minus_rows.get(k)
        if not rows_m:
            return ''
        detail = ', '.join(f'{r}行目 {amt:,.0f}円' for r, _, amt in rows_m)
        return f'①に取消行(H列「-」){len(rows_m)}件あり: {detail} → 純額 {toray[k]:,.0f}円'

    # ①CSVのブロック小計(T行)との検算不一致 → 東レが小計から除外した行の疑い
    # 差額と同額の契約NOが対象工場のものなら、検証Aの当該行(①のみ等)の備考にも転記する
    ct_hint = {}
    for start, idx, b, t_val, dsum, d, cands in ct_errors:
        msg = (f'①CSV {start}〜{idx}行目ブロック(入庫場所 {b}): 小計(T={t_val:,.0f}円)とデータ行合計({dsum:,.0f}円)が'
               f' {d:+,.0f}円 不一致。東レ側で集計除外された行(取消等)の疑い')
        if cands:
            msg += ' → 差額と同額の契約NO: ' + ', '.join(f'{k}({kb}) {s:,.0f}円' for k, kb, s in cands)
            for k, kb, s in cands:
                if kb == args.basho:
                    ct_hint[k] = f'①小計(T行)から除外の疑い: {start}〜{idx}行目ブロックの小計差 {d:+,.0f}円 と同額'
        else:
            msg += '。①原本シートの該当行を確認'
        all_warnings.append(msg)

    excluded_basho = ' / '.join(f'{k} {v:,.0f}円' for k, v in sorted(toray_by_basho.items()) if k != args.basho)

    # ---- 検証C (湖南): 月次処理ファイルの東レ合計 vs 本検証の各総額 ----
    # 月次処理ファイルには契約NO等の明細が無いため合計のみ照合する。①との比較は検証A集計後に追加
    check_c = []   # (項目, 月次処理ファイル値, 本検証値, 差(本検証-月次), 判定, 備考)
    if monthly:
        total2_irai = sum(nag_irai.values())
        total3 = sum(aladdin.values())
        for key, ours, ours_label in (
                ('月次検証: 月次実績①(Excel)', total2_irai, f'②{name2} 東レ3シート合計'),
                ('集計表: 東レ合計 加工金額', total2_irai, f'②{name2} 東レ3シート合計'),
                ('月次検証: 月次実績表②(アラジン)', total3, '③アラジン 東レ(049006)合計')):
            mv = monthly.get(key)
            if mv is None:
                check_c.append((key, None, ours, None, '読取不可', f'{ours_label} と比較できず'))
                all_warnings.append(f'月次処理ファイルの「{key}」を読み取れませんでした: {monthly_path.name}')
                continue
            d = ours - mv
            judge = '一致' if abs(d) <= args.tol else '不一致'
            check_c.append((key, mv, ours, d, judge, f'vs {ours_label}'))
            if judge == '不一致':
                all_warnings.append(f'月次処理ファイル「{key}」{mv:,.0f}円 と {ours_label} {ours:,.0f}円 が {d:+,.0f}円 不一致')

    def find_prev(k):
        """契約NOを過去月まとめから検索 (新しい月優先)。戻り値: (ラベル, 当時AA, 依頼NO集合) or None"""
        for label, pk, pim in prev_data:
            if k in pk:
                return label, pk[k], pim.get(k, set())
        return None

    def find_next(k):
        """契約NOを翌月以降のまとめから検索 (近い月優先)。戻り値: (ラベル, 翌月AA, 依頼NO集合) or None"""
        for label, nk, nim in next_data:
            if k in nk:
                return label, nk[k], nim.get(k, set())
        return None

    out = io.StringIO()
    w = lambda s='': print(s, file=out)
    records = []  # 結果CSV行: [検証区分, キー, ①金額, ②金額, ③金額, 差額, 判定, 備考]

    ym_label = f'{target_ym[0]}年{target_ym[1]}月度' if target_ym else '(不明)'
    w('=' * 78)
    w(f'後加工工賃 金額整合性検証レポート【{fac["label"]}】')
    w(f'対象月: {ym_label}')
    w(f'実行日時: {datetime.now():%Y-%m-%d %H:%M:%S}')
    w(f'① 東レCSV   : {p1.name} (入庫場所 {args.basho} のみ)')
    w(f'② {name2}: {disp2(p2)} ({fac["desc2"]})')
    w(f'   過去月明細: {", ".join(disp2(pf) for _, pf in prevs) if prevs else "なし"} (直近{PREV_MONTH_WINDOW}か月まで)')
    w(f'   翌月明細  : {", ".join(disp2(nf) for _, nf in nexts) if nexts else "なし"}')
    if Path(fac['dir2']).is_absolute():
        w(f'   参照フォルダ: {dir2}')
    w(f'③ アラジン  : {p3.name} (対象年月: {taisho}'
      + (f' / 得意先 {fac["customer3"]} のみ' if fac['customer3'] else '') + ')')
    w(f'   他月照会  : {", ".join(aladdin_other_files) if aladdin_other_files else "なし"}')
    if Path(fac['dir3']).is_absolute():
        w(f'   参照フォルダ: {dir3}')
    if monthly_path:
        w(f'月次処理ファイル: {monthly_path.name} (東レ合計のみ照合)')
        if Path(fac['dir_monthly']).is_absolute():
            w(f'   参照フォルダ: {monthly_path.parent}')
    w(f'許容差: ±{args.tol}円')
    if excluded_basho:
        w(f'①の対象外入庫場所: {excluded_basho}')
    n_minus = sum(len(v) for v in minus_rows.values())
    minus_total = sum(a for v in minus_rows.values() for _, _, a in v)
    minus_label = (f'{n_minus}件 {minus_total:,.0f}円 を符号反転して集計 (契約NO: {", ".join(sorted(minus_rows))})'
                   if n_minus else 'なし')
    w(f'①の取消行(H列「-」): {minus_label}')

    # ---- 検証A の判定 (警告ブロックより前に確定させる: 手動判定の不適用も警告に載せるため) ----
    k1, k2 = set(toray), set(nag_keiyaku)
    common = k1 & k2
    # 手動判定 (●自動検証\手動判定.csv): 当月・当工場の契約NOについて「①正」「②正」を人が指定した不一致は要確認から外す
    manual, wm = load_manual_judgments(base / MANUAL_FILE, fac['label'], target_ym)
    all_warnings += wm
    mismatch_a = []     # 要確認の不一致
    manual_a = []       # 手動判定済みの不一致 (k, ①, ②, 差, side, 理由)
    for k in sorted(common):
        d = toray[k] - nag_keiyaku[k]
        if abs(d) > args.tol:
            if k in manual:
                manual_a.append((k, toray[k], nag_keiyaku[k], d, *manual[k]))
            else:
                mismatch_a.append((k, toray[k], nag_keiyaku[k], d))
    manual_done = []    # 指定はあるが当月すでに①②一致 (②を修正済みとみられる) → 警告ではなく情報として表示
    for k in sorted(set(manual) - {m[0] for m in manual_a}):
        if k in common:
            manual_done.append(k)
            continue
        state = '①のみ' if k in k1 else '②のみ' if k in k2 else '①②とも無し'
        all_warnings.append(f'{MANUAL_FILE}: 契約NO {k} は当月の検証Aで不一致ではありません({state})。手動判定は適用されませんでした(契約NOの誤記?)')
    manual_judge = {m[0]: ('①正(手動)' if m[4] == 1 else '②正(手動)') for m in manual_a}
    if manual_a or manual_done:
        w(f'手動判定: {MANUAL_FILE} の {ym_label}・{fac["label"]} 行 '
          + (f'{len(manual_a)}件 適用 ({", ".join(m[0] for m in manual_a)})' if manual_a else '適用なし')
          + (f' / 既に①②一致のため不要 {len(manual_done)}件 ({", ".join(manual_done)})' if manual_done else ''))
    w('=' * 78)
    if all_warnings:
        w()
        w(f'【警告】データ品質の注意 ({len(all_warnings)}件)')
        for wmsg in all_warnings:
            w(f'  ⚠ {wmsg}')

    # ---- 総額サマリ ----
    w()
    w(f'【総額】 ①東レ検収 {sum(toray.values()):>15,.0f} 円 ({len(toray)}契約)')
    w(f'         ②{name2:　<5}{sum(nag_irai.values()):>15,.0f} 円 ({len(nag_irai)}依頼 / {len(nag_keiyaku)}契約)')
    w(f'         ③アラジン {sum(aladdin.values()):>15,.0f} 円 ({len(aladdin)}依頼)')

    # ---- 検証A: ①vs② 契約NO ----

    def manual_note(m):
        k, v1, v2, d, side, reason = m
        act = (f'②を①金額 {v1:,.0f}円 に修正する(差 {d:+,.0f}円)。東レへは報告しない' if side == 1
               else f'①の誤り {d:+,.0f}円 として東レへ報告する(報告計上額に含む)')
        return f'手動判定: 今回({ym_label})に限り{"①東レ" if side == 1 else "②" + name2}を正とする → {act}' + (f' / 理由: {reason}' if reason else '')
    # 同一依頼NO内の不一致どうしで差額が相殺関係にあれば「契約間振替の疑い」を注記
    pair_notes = {}
    by_irai_group = defaultdict(list)
    for k, _, _, d in mismatch_a:
        for irai in keiyaku_to_irai.get(k, ()):
            by_irai_group[irai].append((k, d))
    for irai, items in by_irai_group.items():
        if len(items) < 2:
            continue
        net = sum(d for _, d in items)
        gross = sum(abs(d) for _, d in items)
        if abs(net) < gross * 0.5:
            for k, _ in items:
                pair_notes[k] = (f'契約間振替の疑い: 同一依頼NO {irai} 内の不一致{len(items)}件と相殺関係'
                                 f' (ネット差額 {net:+,.0f}円)')
    # 金額0円の行は無効行(取消済み・翌月予定行など)として Excel には出さず、件数だけ記録する
    zero_excluded = Counter()

    def is_zero(*vals):
        return all(abs(v) <= args.tol for v in vals)

    mismatch_keys = {k for k, *_ in mismatch_a}
    for k in sorted(common):
        if is_zero(toray[k], nag_keiyaku[k]):
            zero_excluded['A一致'] += 1
            continue
        d = toray[k] - nag_keiyaku[k]
        judge = '不一致' if k in mismatch_keys else manual_judge.get(k, '一致')
        extra = (minus_note(k), ct_hint.get(k, '')) if judge == '不一致' else ()
        note_a = ' / '.join(x for x in (pair_notes.get(k, ''), *extra) if x)
        if k in manual_judge:
            note_a = manual_note(next(m for m in manual_a if m[0] == k))
        records.append(['A:①vs②契約NO', k, toray[k], nag_keiyaku[k], '', d, judge, note_a])
    n_common_a = len(common) - zero_excluded['A一致']
    w()
    w('-' * 78)
    w(f'【検証A】①東レCSV vs ②{name2} (契約NO突合)')
    w(f'  共通契約NO {n_common_a} 件中  一致 {n_common_a-len(mismatch_a)-len(manual_a)} 件 / 不一致 {len(mismatch_a)} 件'
      + (f' / 手動判定済み {len(manual_a)} 件' if manual_a else '')
      + (f'  (ほかに①②とも0円の無効行 {zero_excluded["A一致"]}件を除外)' if zero_excluded['A一致'] else ''))
    if mismatch_a:
        w()
        w(f'  ▼金額不一致 ({len(mismatch_a)}件) ※差額の大きい順')
        w(f'  {"契約NO":<10} {"①東レ":>14} {"②" + name2:>14} {"差額(①-②)":>14}')
        for k, v1, v2, d in sorted(mismatch_a, key=lambda x: -abs(x[3])):
            hint = ' / '.join(x for x in (pair_notes.get(k, ''), minus_note(k), ct_hint.get(k, '')) if x)
            note = f'  ({hint})' if hint else ''
            w(f'  {k:<10} {v1:>14,.1f} {v2:>14,.1f} {d:>14,.1f}{note}')
    if manual_a:
        w()
        w(f'  ▼手動判定済みの不一致 ({len(manual_a)}件) ※{MANUAL_FILE} の指定。要確認対象外')
        w(f'  {"契約NO":<10} {"①東レ":>14} {"②" + name2:>14} {"差額(①-②)":>14}  判定')
        for m in manual_a:
            k, v1, v2, d = m[:4]
            w(f'  {k:<10} {v1:>14,.1f} {v2:>14,.1f} {d:>14,.1f}  {manual_judge[k]}  ({manual_note(m)})')
    def report_only(keys, amounts, judge, label, note, col_1, tag, note_fn=None):
        """片側のみのキーを出力。金額0円のキーは無効行として除外(件数のみ記録)し、金額ありのキーのリストを返す。"""
        nonzero = [k for k in keys if abs(amounts[k]) > args.tol]
        n_zero = len(keys) - len(nonzero)
        if nonzero:
            w()
            w(f'  ▼{label} ({len(nonzero)}件 / 計 {sum(amounts[k] for k in nonzero):,.0f}円) {note}')
            for k in nonzero:
                extra = f'  ({note_fn(k)})' if note_fn and note_fn(k) else ''
                w(f'  {k:<10} {amounts[k]:>14,.1f}{extra}')
        if n_zero:
            zero_excluded[f'{tag[0]}{judge}'] += n_zero
            w(f'  (金額0円の{label} {n_zero}件は無効行として除外)')
        for k in nonzero:
            rec = ['', '', '']
            rec[col_1] = amounts[k]
            n = note.lstrip('※')
            if note_fn and note_fn(k):
                n = f'{n} / {note_fn(k)}'
            records.append([tag, k] + rec + ['', judge, n])
        return nonzero

    def nyuko_note(k):
        """契約NOの①入庫月日(+取消行・小計除外の疑い)を備考用に返す"""
        dates = sorted(nyuko_dates.get(k, ()))
        parts = [f'①入庫日 {", ".join(dates)}'] if dates else []
        if minus_note(k):
            parts.append(minus_note(k))
        if k in ct_hint:
            parts.append(ct_hint[k])
        return ' / '.join(parts)

    # ①のみの契約NOを過去月まとめと突合し、「前月調整」と「真の①のみ」に分ける
    only1_keys = sorted(k1 - k2)
    prev_hit = {k: hit for k in only1_keys if (hit := find_prev(k))}
    prev_adjust = sorted(prev_hit)
    only1_rest = [k for k in only1_keys if k not in prev_hit]
    if prev_adjust:
        w()
        w(f'  ▼前月調整とみられる契約NO ({len(prev_adjust)}件 / ①計 {sum(toray[k] for k in prev_adjust):,.0f}円)'
          f' ※②当月に無いが過去月まとめに存在')
        for k in prev_adjust:
            label, prev_aa, irais = prev_hit[k]
            d = toray[k] - prev_aa
            irai = ', '.join(sorted(irais))
            note = f'{label}の依頼NO {irai} / 当時{amt2} {prev_aa:,.0f}円 / ①との差 {d:+,.0f}円'
            if nyuko_note(k):
                note += f' / {nyuko_note(k)}'
            w(f'  {k:<10} {toray[k]:>14,.1f}  ({note})')
            records.append(['A:①vs②契約NO', k, toray[k], '', '', '', '前月調整', note])

    # 前月の工賃差検証で東レへ次月修正を依頼したが、当月①に調整行がまだ来ていない差額
    # (例: 7月差で長岡>東レの契約。マイナス調整は①取消行として前月調整に出るが、プラス側は①に現れない)
    prior_rows, wp = load_prior_shortfalls(base / PRIOR_FILE, fac['label'], target_ym)
    all_warnings += wp
    prior_shortfall = []
    for row in prior_rows:
        k = row['契約NO']
        if k in prev_adjust:
            all_warnings.append(f'{PRIOR_FILE}: 契約NO {k} は既に前月調整(①に調整行あり)のため重複を避けスキップしました')
            continue
        if k in k1 and k in k2:
            all_warnings.append(f'{PRIOR_FILE}: 契約NO {k} は当月①②の両方に存在します。当月の不一致として扱うため前月過不足はスキップしました')
            continue
        if k in k1:
            all_warnings.append(f'{PRIOR_FILE}: 契約NO {k} は当月①に存在します(前月調整判定を確認)。前月過不足はスキップしました')
            continue
        prior_shortfall.append(row)
    if prior_shortfall:
        prior_total = sum(r['差額'] for r in prior_shortfall)
        w()
        w(f'  ▼前月過不足 ({len(prior_shortfall)}件 / 差額①-②計 {prior_total:+,.0f}円)'
          f' ※前月の工賃差検証で次月修正依頼済みだが、当月①CSVに調整行が未反映。報告する過不足に計上')
        for r in prior_shortfall:
            # プラス差(長岡−東レ): マイナス調整は①取消行で「前月調整」に出るが、プラス側は①CSVに現れない
            gap = r['②'] - r['①']  # 例 +8899 / +20
            note = (
                '当月①CSVに未反映'
                + (f'（プラス差 {gap:+,.0f}円は①データに含まれない。マイナス側は取消行で前月調整に出る）'
                   if gap > 0.5 else f'（差額 {r["差額"]:+,.0f}円）')
                + f' / 前月当時① {r["①"]:,.0f}円・② {r["②"]:,.0f}円 / 報告計上(①-②) {r["差額"]:+,.0f}円'
                + (f' / {r["理由"]}' if r['理由'] else '')
            )
            irai_disp = r['依頼NO'] or ''
            w(f'  {r["契約NO"]:<10} {"—":>14} {gap:>14,.1f} {r["差額"]:>14,.1f}'
              + (f'  ({irai_disp})' if irai_disp else '') + f'  ({note})')
            # ①列は空(当月①に無し)、②列にプラス差(長岡−東レ)、差額・報告計上は前月の①−②のまま
            records.append(['A:①vs②契約NO', r['契約NO'], '', gap, '', r['差額'], '前月過不足', note])
            if r['依頼NO']:
                # irai_of が空にならないよう、keiyaku_to_irai に前月依頼NOを補完 (当月に無い契約向け)
                keiyaku_to_irai.setdefault(r['契約NO'], set()).add(r['依頼NO'])

    # 残った①のみを翌月まとめと突合し、「翌月記載(②の記載月誤り)」と「真の①のみ」に分ける
    # ①純額0円(検収後に取消行で相殺)は当月検収なしと同じなので、翌月②に載っていても記載月誤りとはしない(無効行として除外)
    next_hit = {k: hit for k in only1_rest if abs(toray[k]) > args.tol and (hit := find_next(k))}
    next_missrec = sorted(next_hit)
    only1_rest = [k for k in only1_rest if k not in next_hit]
    if next_missrec:
        cur_label = f'{target_ym[1]}月度' if target_ym else '当月'
        w()
        w(f'  ▼翌月の②まとめに記載されている契約NO ({len(next_missrec)}件 / ①計 {sum(toray[k] for k in next_missrec):,.0f}円)'
          f' ※①は{cur_label}検収 → ②の記載月誤り。②2ファイルの修正が必要')
        for k in next_missrec:
            label, next_aa, irais = next_hit[k]
            d = toray[k] - next_aa
            irai = ', '.join(sorted(irais))
            note = (f'{label}の依頼NO {irai} に記載あり({amt2} {next_aa:,.0f}円 / ①との差 {d:+,.0f}円)。'
                    f'本来は{cur_label}の②に記載すべきもの → ②{cur_label}へ追記し{label}から削除する修正が必要')
            if nyuko_note(k):
                note += f' / {nyuko_note(k)}'
            w(f'  {k:<10} {toray[k]:>14,.1f}  ({note})')
            records.append(['A:①vs②契約NO', k, toray[k], '', '', '', '翌月記載', note])
    only1 = report_only(only1_rest, toray, '①のみ', '①のみに存在する契約NO',
                        f'※②{name2}に記載なし(過去月・翌月にも無し)', 0, 'A:①vs②契約NO', note_fn=nyuko_note)
    only2 = report_only(sorted(k2 - k1), nag_keiyaku, '②のみ', '②のみに存在する契約NO', '※①東レCSVに検収なし', 1, 'A:①vs②契約NO')
    if bad_keiyaku:
        # 「9月」「キャンセル」等のメモ書き。契約NO形式(ハイフン無し7桁英数字)でないため突合・Excel出力とも対象外
        counts = Counter(c for _, c, _ in bad_keiyaku)
        summary = ' / '.join(f'"{v}"×{n}' for v, n in counts.most_common())
        w()
        w(f'  (②のC列が契約NO形式でない行 {len(bad_keiyaku)}件は突合・出力対象外: {summary})')
        with_amount = [(row_no, c, aa) for row_no, c, aa in bad_keiyaku if abs(aa) > args.tol]
        if with_amount:
            w(f'  ※うち{len(with_amount)}件は{amt2}列に金額あり。契約NOの記入漏れの可能性:')
            for row_no, c, aa in with_amount:
                w(f'  行{row_no}: "{c}" ({amt2}列 {aa:,.1f}円)')
                records.append(['A:①vs②契約NO', c, '', aa, '', '', '形式不正', f'②{name2} {row_no}行目 C列が契約NO形式でないが金額あり'])

    # ---- 検証B: ②vs③ 依頼NO ----
    def month_note(k):
        """依頼NOの月番号(例 Y7-52 → 7)が対象月と異なる場合の月ずれ注記"""
        m = re.match(r'^[A-Z]+(\d{1,2})-', k)
        if m and target_ym and int(m.group(1)) != target_ym[1]:
            return f'{m.group(1)}月系依頼(計上月ずれの可能性)'
        return ''

    # 累計比較: 依頼NOの②全月度累計と③全月累計が一致すれば「月ずれ解消」(計上月ずれと判明)
    resolved_b = {}

    def cum_parts(k):
        parts2 = [(ym, d[k]) for ym, d in sorted(nag_irai_by_ym.items()) if k in d]
        parts3 = [(ym, d[k]) for ym, d in sorted(aladdin_by_ym.items()) if k in d]
        return parts2, parts3

    def cum_labels(parts2, parts3):
        return ('+'.join(f'{m}月度' for (_, m), _ in parts2),
                '+'.join(f'{m}月' for (_, m), _ in parts3))

    def try_resolve(k):
        if k in resolved_b:
            return True
        parts2, parts3 = cum_parts(k)
        if not parts2 or not parts3 or len(aladdin_by_ym) < 2:
            return False
        cum2 = sum(v for _, v in parts2)
        cum3 = sum(v for _, v in parts3)
        if abs(cum2 - cum3) > args.tol:
            return False
        l2, l3 = cum_labels(parts2, parts3)
        resolved_b[k] = f'累計一致: ②累計 {cum2:,.0f}円({l2}) = ③累計 {cum3:,.0f}円({l3}) → 計上月ずれと判明'
        return True

    def cum_summary(k):
        """非解消項目の備考用に累計内訳を返す (③に実績が全く無ければ空)"""
        parts2, parts3 = cum_parts(k)
        if not parts3 or len(aladdin_by_ym) < 2:
            return ''
        cum2 = sum(v for _, v in parts2)
        cum3 = sum(v for _, v in parts3)
        l2, l3 = cum_labels(parts2, parts3)
        return f'②累計 {cum2:,.0f}円({l2}) / ③累計 {cum3:,.0f}円({l3}) / 累計差 {cum2 - cum3:+,.0f}円'

    def b_note(k):
        return ' / '.join(x for x in (month_note(k), cum_summary(k)) if x)

    i2, i3 = set(nag_irai), set(aladdin)
    common_b = i2 & i3

    # 枝番統合: ③が枝番(例 C8-3-1)に分割計上し②は親依頼NO(C8-3)に一本で計上しているケース。
    # ③の親+枝番の合計が②の親と一致すれば「枝番統合一致」として要確認から外す
    branch_children = defaultdict(list)
    for k in sorted(i3 - i2):
        parent = re.sub(r'-\d+$', '', k)
        if parent != k and parent in i2 and abs(aladdin[k]) > args.tol:
            branch_children[parent].append(k)
    branch_ok = {}      # 親依頼NO -> 注記
    branch_child = {}   # 枝番依頼NO -> 親依頼NO
    for parent, kids in branch_children.items():
        if parent in i3 and abs(nag_irai[parent] - aladdin[parent]) <= args.tol:
            continue    # 親だけで一致済み → 枝番は本当の③のみ
        merged3 = aladdin.get(parent, 0) + sum(aladdin[c] for c in kids)
        if abs(nag_irai[parent] - merged3) <= args.tol:
            parts = ' + '.join([f'{parent} {aladdin.get(parent, 0):,.0f}'] + [f'{c} {aladdin[c]:,.0f}' for c in kids])
            branch_ok[parent] = f'枝番統合一致: ③ {parts} = {merged3:,.0f}円 = ② {nag_irai[parent]:,.0f}円 → ③が枝番に分割計上'
            for c in kids:
                branch_child[c] = parent

    mismatch_b = []
    resolved_rows = []   # (由来, 依頼NO, ②単月, ③単月)
    branch_rows = []     # (依頼NO, ②単月, ③単月)  枝番統合一致 (親・枝番とも)
    for k in sorted(common_b):
        d = nag_irai[k] - aladdin[k]
        if is_zero(nag_irai[k], aladdin[k]):
            zero_excluded['B一致'] += 1
            continue
        if abs(d) <= args.tol:
            records.append(['B:②vs③依頼NO', k, '', nag_irai[k], aladdin[k], d, '一致', ''])
        elif k in branch_ok:
            branch_rows.append((k, nag_irai[k], aladdin[k]))
            records.append(['B:②vs③依頼NO', k, '', nag_irai[k], aladdin[k], d, '枝番統合一致', branch_ok[k]])
        elif try_resolve(k):
            resolved_rows.append(('不一致', k, nag_irai[k], aladdin[k]))
            records.append(['B:②vs③依頼NO', k, '', nag_irai[k], aladdin[k], d, '月ずれ解消', resolved_b[k]])
        else:
            mismatch_b.append((k, nag_irai[k], aladdin[k], d))
            records.append(['B:②vs③依頼NO', k, '', nag_irai[k], aladdin[k], d, '不一致', b_note(k)])
    # ②のみ・③のみも累計一致なら月ずれ解消として除外
    only2b_keys = []
    for k in sorted(i2 - i3):
        if k in branch_ok:
            branch_rows.append((k, nag_irai[k], 0))
            records.append(['B:②vs③依頼NO', k, '', nag_irai[k], '', '', '枝番統合一致', branch_ok[k]])
        elif abs(nag_irai[k]) > args.tol and try_resolve(k):
            resolved_rows.append(('②のみ', k, nag_irai[k], 0))
            records.append(['B:②vs③依頼NO', k, '', nag_irai[k], '', '', '月ずれ解消', resolved_b[k]])
        else:
            only2b_keys.append(k)
    w()
    w('-' * 78)
    w(f'【検証B】②{name2}({amt2}) vs ③アラジン(加工金額) (依頼NO突合)')
    n_branch_common = sum(1 for k, _, _ in branch_rows if k in common_b)
    n_common_b = len(common_b) - zero_excluded['B一致']
    w(f'  共通依頼NO {n_common_b} 件中  一致 {n_common_b-len(mismatch_b)-sum(1 for s,*_ in resolved_rows if s=="不一致")-n_branch_common} 件'
      f' / 月ずれ解消(累計一致) {sum(1 for s, *_ in resolved_rows if s == "不一致")} 件'
      f' / 枝番統合一致 {n_branch_common} 件'
      f' / 不一致(要確認) {len(mismatch_b)} 件'
      + (f'  (ほかに②③とも0円の無効行 {zero_excluded["B一致"]}件を除外)' if zero_excluded['B一致'] else ''))
    if len(aladdin_by_ym) >= 2:
        yms = ', '.join(f'{y}年{m}月' for y, m in sorted(aladdin_by_ym))
        w(f'  ※累計比較に使用した③照会月: {yms} / ②月度: '
          + ', '.join(f'{y}年{m}月度' for y, m in sorted(nag_irai_by_ym)))
    if mismatch_b:
        n_tsukizure = sum(1 for k, *_ in mismatch_b if month_note(k))
        w()
        w(f'  ▼金額不一致・要確認 ({len(mismatch_b)}件) ※差額の大きい順'
          + (f' / うち過去月系依頼 {n_tsukizure}件(計上月ずれの可能性)' if n_tsukizure else ''))
        w(f'  {"依頼NO":<10} {"②" + name2:>14} {"③アラジン":>14} {"差額(②-③)":>14}')
        for k, v2, v3, d in sorted(mismatch_b, key=lambda x: -abs(x[3])):
            extra = f'  ({month_note(k)})' if month_note(k) else ''
            w(f'  {k:<10} {v2:>14,.1f} {v3:>14,.1f} {d:>14,.1f}{extra}')
    only2b = report_only(only2b_keys, nag_irai, '②のみ', '②のみに存在する依頼NO',
                         '※③アラジンに計上なし(他月にも累計一致なし)', 1, 'B:②vs③依頼NO', note_fn=b_note)

    def note3(k):
        """③のみ依頼NOの注記 (枝番の親依頼NO確認・月ずれ)"""
        parent = re.sub(r'-\d+$', '', k)
        if parent != k and parent in i2:
            merged3 = aladdin.get(parent, 0) + aladdin[k]
            note = (f'枝番? 親依頼NO {parent} は②に存在。'
                    f'③親+枝番計 {merged3:,.0f}円 / ②親 {nag_irai[parent]:,.0f}円 で突合してみること')
        else:
            note = f'②{name2}に依頼NOなし'
        if month_note(k):
            note += f' / {month_note(k)}'
        return note

    only3b_all = sorted(i3 - i2)
    only3b = []
    n_zero3 = 0
    for k in only3b_all:
        if abs(aladdin[k]) <= args.tol:
            n_zero3 += 1
            zero_excluded['B③のみ'] += 1
        elif k in branch_child:
            branch_rows.append((k, 0, aladdin[k]))
            records.append(['B:②vs③依頼NO', k, '', '', aladdin[k], '', '枝番統合一致',
                            f'親依頼NO {branch_child[k]} に合算して②と一致 / {branch_ok[branch_child[k]]}'])
        elif try_resolve(k):
            resolved_rows.append(('③のみ', k, 0, aladdin[k]))
            records.append(['B:②vs③依頼NO', k, '', '', aladdin[k], '', '月ずれ解消', resolved_b[k]])
        else:
            only3b.append(k)
            records.append(['B:②vs③依頼NO', k, '', '', aladdin[k], '', '③のみ', note3(k)])
    if only3b:
        w()
        w(f'  ▼③のみに存在する依頼NO ({len(only3b)}件 / 計 {sum(aladdin[k] for k in only3b):,.0f}円)'
          f' ※②に記載なし(他月にも累計一致なし)')
        for k in only3b:
            w(f'  {k:<10} {aladdin[k]:>14,.1f}  ({note3(k)})')
    if n_zero3:
        w(f'  (金額0円の③のみ依頼NO {n_zero3}件は無効行として除外)')

    if resolved_rows:
        w()
        w(f'  ▼累計一致で計上月ずれと判明した依頼NO ({len(resolved_rows)}件) ※要確認対象外'
          f' (由来: 不一致{sum(1 for s, *_ in resolved_rows if s == "不一致")}'
          f' / ②のみ{sum(1 for s, *_ in resolved_rows if s == "②のみ")}'
          f' / ③のみ{sum(1 for s, *_ in resolved_rows if s == "③のみ")})')
        w(f'  {"依頼NO":<10} {"②当月":>14} {"③当月":>14}')
        for _, k, v2, v3 in sorted(resolved_rows, key=lambda x: x[1]):
            w(f'  {k:<10} {v2:>14,.1f} {v3:>14,.1f}  ({resolved_b[k]})')
    if branch_rows:
        w()
        w(f'  ▼③の枝番を親依頼NOに合算すると②と一致した依頼NO ({len(branch_rows)}件 / 親{len(branch_ok)}件) ※要確認対象外')
        w(f'  {"依頼NO":<10} {"②当月":>14} {"③当月":>14}')
        for k, v2, v3 in sorted(branch_rows):
            w(f'  {k:<10} {v2:>14,.1f} {v3:>14,.1f}  ({branch_ok.get(k) or "親依頼NO " + branch_child[k] + " に合算"})')

    # ---- 報告用内訳 (東レ宛メール報告の元数字) ----
    total1 = sum(toray.values())
    total2 = sum(nag_keiyaku.values())
    adj_total = sum(toray[k] for k in prev_adjust)
    next_total = sum(toray[k] for k in next_missrec)
    # 手動判定: ②正 → 東レの誤りなので当月差異(報告)に含める / ①正 → ②を直す側なので報告に含めず別枠で分解
    manual2 = [m for m in manual_a if m[4] == 2]
    manual1 = [m for m in manual_a if m[4] == 1]
    manual1_total = sum(m[3] for m in manual1)
    prior_sf_total = sum(r['差額'] for r in prior_shortfall)
    diff_total = sum(d for _, _, _, d in mismatch_a) + sum(m[3] for m in manual2)
    n_diff = len(mismatch_a) + len(manual2)
    only1_total = sum(toray[k] for k in only1)
    only2_total = sum(nag_keiyaku[k] for k in only2)
    residual = total1 - (total2 + adj_total + next_total + diff_total + manual1_total + only1_total - only2_total)
    w()
    w('-' * 78)
    w(f'【報告用内訳】(①お支払データ = ②実売上 + 前月調整 + 翌月記載 + 当月差異 + 手動①正 + 未解明 の分解)')
    w(f'  ①東レお支払データ({args.basho})  {total1:>15,.0f} 円')
    w(f'  ②当月実売上({name2}計)    {total2:>15,.0f} 円')
    w(f'  前月以前調整分                {adj_total:>15,.0f} 円 ({len(prev_adjust)}件)')
    w(f'  翌月記載(②の記載月誤り)      {next_total:>15,.0f} 円 ({len(next_missrec)}件) ※②修正後は②実売上に入る')
    w(f'  当月差異(金額不一致の差計)    {diff_total:>15,.0f} 円 ({n_diff}件)'
      + (f' ※うち手動判定②正 {len(manual2)}件 {sum(m[3] for m in manual2):,.0f}円' if manual2 else ''))
    if manual1:
        w(f'  手動判定①正(②要修正)        {manual1_total:>15,.0f} 円 ({len(manual1)}件) ※東レへは報告しない。②修正後は②実売上に入る')
    if prior_shortfall:
        w(f'  ※前月過不足(①未反映・報告計上) {prior_sf_total:>10,.0f} 円 ({len(prior_shortfall)}件)'
          f' ← ①には含まれないため上の分解式の外。報告する過不足(当月)に加算')
    w(f'  未解明(①のみ 前月翌月無し)   {only1_total:>15,.0f} 円 ({len(only1)}件)')
    w(f'  未解明(②のみ ①検収なし)     {-only2_total:>15,.0f} 円 ({len(only2)}件)')
    w(f'  検算残差                      {residual:>15,.0f} 円 (0なら分解が完全)')
    if next_missrec:
        w()
        w(f'  ※翌月記載 {len(next_missrec)}件は②の2ファイル修正(当月へ追記・翌月から削除)後に再実行し、数字を確定させること')

    # ---- 検証C (湖南): 月次処理ファイルの東レ合計との照合 ----
    if monthly:
        # ①東レCSV vs 売上明細③(長岡が東レへ請求した売上): 差は報告用内訳(前月調整+翌月記載+当月差異+手動①正+未解明)で説明されるはず
        explained = adj_total + next_total + diff_total + manual1_total + only1_total - only2_total
        key = '月次検証: 売上明細③'
        mv = monthly.get(key)
        if mv is None:
            check_c.append((key, None, total1, None, '読取不可', '①東レCSV総額 と比較できず'))
        else:
            d = total1 - mv
            if abs(d) <= args.tol:
                judge, note = '一致', 'vs ①東レCSV総額'
            elif abs(d - explained) <= args.tol:
                judge = '内訳で説明可'
                note = (f'vs ①東レCSV総額。差 {d:+,.0f}円 = 前月調整 {adj_total:+,.0f} + 翌月記載 {next_total:+,.0f}'
                        f' + 当月差異 {diff_total:+,.0f}' + (f' + 手動①正 {manual1_total:+,.0f}' if manual1 else '')
                        + f' + ①のみ {only1_total:+,.0f} − ②のみ {only2_total:,.0f}')
            else:
                judge = '要確認'
                note = (f'vs ①東レCSV総額。差 {d:+,.0f}円 のうち報告用内訳で説明できるのは {explained:+,.0f}円'
                        f' (未説明 {d - explained:+,.0f}円)')
            check_c.append((key, mv, total1, d, judge, note))
        w()
        w('-' * 78)
        w(f'【検証C】月次処理ファイル({monthly_path.name}) 東レ合計 vs 本検証の総額')
        w(f'  ※月次処理ファイルには契約NO等の明細が無いため東レ部分の合計金額のみ照合')
        w(f'  {"項目":<32} {"月次処理ﾌｧｲﾙ":>14} {"本検証":>14} {"差":>12}  判定')
        for key, mv, ours, d, judge, note in check_c:
            mv_s = f'{mv:,.2f}' if mv is not None else '(読取不可)'
            d_s = f'{d:+,.2f}' if d is not None else ''
            w(f'  {key:<32} {mv_s:>14} {ours:>14,.2f} {d_s:>12}  {judge}  ({note})')

    # ---- 検証D (国分): ②「東レまとめ」内部整合性 ----
    D_ERR = ('金額差', '参照ずれ', '未取込')
    if check_d and not check_d['skipped']:
        d_rows = check_d['rows']
        d_counts = Counter(row[6] for row in d_rows)
        n_d_err = sum(d_counts[j] for j in D_ERR)
        w()
        w('-' * 78)
        w(f'【検証D】②{name2}「{MATOME_SHEET}」の内部整合性 (元4シートとの参照・式・金額突合)')
        w(f'  まとめ取込 {check_d["n_map"]}行 / 要修正 {n_d_err}件 (' + ' / '.join(f'{j}{d_counts[j]}' for j in D_ERR)
          + f') / 注意 {d_counts["式異常"] + d_counts["直接入力"]}件 (式異常{d_counts["式異常"]} / 直接入力{d_counts["直接入力"]})')
        w(f'  {"元シート":<8} {"まとめAA計":>14} {"元シートAA計":>14} {"差":>10}')
        for s, sm, ss, d in check_d['totals']:
            w(f'  {s:<8} {sm:>14,.0f} {ss:>14,.0f} {d:>+10,.0f}' + ('' if abs(d) <= args.tol else '  ← 不一致'))
        tot_m, tot_s = sum(t[1] for t in check_d['totals']), sum(t[2] for t in check_d['totals'])
        w(f'  {"合計":<8} {tot_m:>14,.0f} {tot_s:>14,.0f} {tot_m - tot_s:>+10,.0f}')
        shown = [row for row in d_rows if row[6] in D_ERR] + [row for row in d_rows if row[6] not in D_ERR]
        if shown:
            w()
            for place, irai, keiyaku, m_aa, s_aa, d, judge, desc in shown[:30]:
                w(f'  [{judge}] {place:<28} {irai:<8} {keiyaku:<8} {desc}')
            if len(shown) > 30:
                w(f'  ... ほか {len(shown) - 30}件 (検証Dシート参照)')
        else:
            w('  問題なし: まとめの参照式・元シートの式・行別金額がすべて整合')

    # ---- 結論 ----
    w()
    w('=' * 78)
    n_issues = (len(mismatch_a) + len(mismatch_b) + len(next_missrec) + len(prior_shortfall)
                + len(only1) + len(only2) + len(only2b) + len(only3b))
    w(f'【結論】要確認項目 合計 {n_issues} 件'
      f' (A不一致{len(mismatch_a)} / A翌月記載{len(next_missrec)} / A前月過不足{len(prior_shortfall)}'
      f' / A片側のみ{len(only1)+len(only2)}'
      f' / B不一致{len(mismatch_b)} / B片側のみ{len(only2b)+len(only3b)})')
    w(f'       ほかに判明済み: 前月調整 {len(prev_adjust)}件({adj_total:,.0f}円)'
      + (f' / A手動判定 {len(manual_a)}件(①正{len(manual1)}・②正{len(manual2)})' if manual_a else '')
      + f' / B月ずれ解消(累計一致) {len(resolved_rows)}件 / B枝番統合一致 {len(branch_rows)}件')
    zero_parts_a = [f'{k[1:]} {n}件' for k, n in zero_excluded.items() if k.startswith('A') and n]
    zero_parts_b = [f'{k[1:]} {n}件' for k, n in zero_excluded.items() if k.startswith('B') and n]
    zero_label = ((f'検証A: {"・".join(zero_parts_a)}' if zero_parts_a else '')
                  + (' / ' if zero_parts_a and zero_parts_b else '')
                  + (f'検証B: {"・".join(zero_parts_b)}' if zero_parts_b else '')) or 'なし'
    w(f'       金額0円の無効行(Excel出力対象外): {zero_label}')
    if check_c:
        n_c_ng = sum(1 for *_, judge, _ in check_c if judge in ('不一致', '要確認', '読取不可'))
        w(f'       検証C(月次処理ファイル 東レ合計): {len(check_c)}項目中 要確認 {n_c_ng}件')
    if check_d and not check_d['skipped']:
        n_d_err = sum(1 for row in check_d['rows'] if row[6] in D_ERR)
        n_d_note = len(check_d['rows']) - n_d_err
        w(f'       検証D(②まとめ整合性): 要修正 {n_d_err}件 / 注意 {n_d_note}件')
    if all_warnings:
        w(f'       ⚠ データ品質の警告 {len(all_warnings)} 件あり (冒頭の【警告】参照)')
    w('=' * 78)

    report = out.getvalue()
    print(report)

    # ---- 結果Excel出力 ----
    # records: [検証区分, キー, ①, ②, ③, 差額, 判定, 備考] → シート別の列構成に変換して並べ替え
    # 検証Aには②東レまとめから引いた依頼NO列を付与 (1契約NOが複数依頼NOに属す場合はカンマ区切り)
    # 前月調整行は過去月まとめの依頼NOを「(7月)」のような月ラベル付きで表示
    def _irai_month_tag(irai):
        """依頼NOの月番号から「(7月)Y7-51」形式にする (前月調整の表記に合わせる)"""
        m = re.match(r'^[A-Z]+(\d{1,2})-', irai)
        return f'({int(m.group(1))}月){irai}' if m else irai

    def irai_of(r):
        hit = None
        if r[6] == '前月調整' and r[1] in prev_hit:
            hit = prev_hit[r[1]]
        elif r[6] == '翌月記載' and r[1] in next_hit:
            hit = next_hit[r[1]]
        if hit:
            label, _, irais = hit
            m = GATSUDO_PAT.search(label)
            tag = f'({m.group(2)}月)' if m else '(他月)'
            return ', '.join(f'{tag}{x}' for x in sorted(irais))
        if r[6] == '前月過不足':
            # 依頼NO自体の月番号で「(7月)Y7-51」のように前月調整と同じ表記にする
            return ', '.join(_irai_month_tag(x) for x in sorted(keiyaku_to_irai.get(r[1], [])))
        return ', '.join(sorted(keiyaku_to_irai.get(r[1], [])))

    def report_amt(r):
        """検証Aの各行が「報告する過不足(当月)」にいくら計上されるか。
        不一致/②正(手動)/前月過不足=差額(①-②) / 翌月記載・①のみ=①金額 / ②のみ=-②金額 / それ以外=空欄。
        この列の合計 = 当月差異 + 翌月記載 + 前月過不足 + ①のみ − ②のみ = メールで報告する過不足"""
        judge, a1, a2, diff = r[6], r[2], r[3], r[5]
        if judge in ('不一致', '②正(手動)', '前月過不足'):
            return diff
        if judge == '①正(手動)':
            return None    # ①が正 → ②を直す。東レへ報告する過不足には含めない
        if judge in ('翌月記載', '①のみ'):
            return a1
        if judge == '②のみ':
            return -a2 if isinstance(a2, (int, float)) else None
        return None

    records_a = [[r[1], irai_of(r), r[2], r[3], r[5], report_amt(r), r[6], r[7]]
                 for r in records if r[0].startswith('A')]
    records_b = [[r[1], r[3], r[4], r[5], r[6], r[7]] for r in records if r[0].startswith('B')]
    records_a.sort(key=lambda r: _judge_rank(r[6], r[4], r[2:4]) + (r[0],))
    records_b.sort(key=lambda r: _judge_rank(r[4], r[3], r[1:3]) + (r[0],))

    if target_ym:
        ny, nm = (target_ym[0] + 1, 1) if target_ym[1] == 12 else (target_ym[0], target_ym[1] + 1)
        next_label = f'{ny}年{nm}月度'
    else:
        next_label = '翌月'
    try:
        executor = os.getlogin()
    except OSError:
        executor = os.environ.get('USERNAME', '(不明)')
    info = {
        '工場': fac['label'], '他工場': fac['other'],
        '②名称': name2, '②説明': fac['desc2'], '②金額列': amt2,
        '③得意先': fac['customer3'],
        'box2': fac['box2'], 'key_a': fac['key_a'], 'key_b': fac['key_b'], 'take2': fac['take2'],
        '月次ファイル': monthly_path.name if monthly_path else None,
        '月次フォルダ': str(monthly_path.parent) if monthly_path and Path(fac['dir_monthly']).is_absolute() else None,
        '検証C': check_c,
        '検証D': check_d,
        '実行日時': f'{datetime.now():%Y-%m-%d %H:%M:%S}',
        '実行者': executor,
        '対象月ラベル': ym_label, '翌月ラベル': next_label,
        '警告': all_warnings,
        '①ファイル': p1.name, '②ファイル': disp2(p2), '③ファイル': p3.name,
        '②フォルダ': str(dir2) if Path(fac['dir2']).is_absolute() else None,
        '③フォルダ': str(dir3) if Path(fac['dir3']).is_absolute() else None,
        '前月ファイル': ', '.join(disp2(pf) for _, pf in prevs) if prevs else 'なし',
        '翌月ファイル': ', '.join(disp2(nf) for _, nf in nexts) if nexts else 'なし',
        '対象年月': taisho, '許容差': args.tol,
        '入庫場所': args.basho, '対象外入庫場所': excluded_basho, '①取消行': minus_label, '0円除外': zero_label,
        '①総額': round(sum(toray.values())),
        '②総額': round(sum(nag_irai.values())),
        '③総額': round(sum(aladdin.values())),
        'A共通': len(common), 'A不一致': len(mismatch_a),
        'A手動①正件数': len(manual1), 'A手動①正額': round(manual1_total),
        'A手動②正件数': len(manual2), 'A手動②正額': round(sum(m[3] for m in manual2)),
        'A①のみ': len(only1), 'A②のみ': len(only2),
        'A前月調整件数': len(prev_adjust), 'A前月調整額': adj_total,
        'A前月過不足件数': len(prior_shortfall), 'A前月過不足額': round(prior_sf_total),
        # メール注記用: プラス差(長岡−東レ)と依頼NO表示
        'A前月過不足明細': [
            {'契約NO': r['契約NO'], '依頼NO': r['依頼NO'],
             'プラス差': round(r['②'] - r['①']), '差額': round(r['差額']),
             '依頼表示': _irai_month_tag(r['依頼NO']) if r['依頼NO'] else r['契約NO']}
            for r in prior_shortfall
        ],
        'A翌月記載件数': len(next_missrec), 'A翌月記載額': next_total,
        '報告②実売上': round(total2), '報告当月差異': round(diff_total),
        '報告①のみ額': only1_total, '報告②のみ額': only2_total,
        '報告残差': round(residual),
        'B共通': len(common_b), 'B不一致': len(mismatch_b),
        'B月ずれ解消': len(resolved_rows),
        'B月ずれ解消共通': sum(1 for s, *_ in resolved_rows if s == '不一致'),
        'B枝番統合': len(branch_rows),
        'B枝番統合共通': n_branch_common,
        'B②のみ': len(only2b), 'B③のみ': len(only3b),
        '要確認': n_issues,
    }
    # ---- 元データのコピーシートを収集 ----
    # ①CSV(引用符内改行をcsvモジュールで正しく解釈)、②当月まとめ、前月調整の根拠になった過去月まとめ、③アラジン
    csv_text = p1.read_bytes().decode('cp932')
    csv_rows = [[_csv_cell(c) for c in row] for row in csv.reader(io.StringIO(csv_text))]

    # 検証明細で色が付いたキーに対応する原本行をハイライトするための集合と判定関数
    att_keiyaku = {rec[0] for rec in records_a
                   if rec[6] in ('不一致', '翌月記載', '①のみ', '②のみ', '前月調整', '前月過不足', *MANUAL_JUDGES)}
    att_irai = {rec[0] for rec in records_b if rec[4] in ('不一致', '②のみ', '③のみ')}

    def hl_csv(vals):
        return any(isinstance(v, str) and norm(v).replace('-', '') in att_keiyaku for v in vals[:6])

    def hl_2(vals):
        return fac['hl2'](vals, att_irai, att_keiyaku)

    def hl_aladdin(vals):
        return any(isinstance(v, str) and norm(v) in att_irai for v in vals[:6])

    raw_sheets = [('①東レCSV原本', csv_rows, 'A2', hl_csv)]
    for sheet_name, rows in fac['raw2'](p2):
        raw_sheets.append((f'②{sheet_name}(当月)', rows, fac['freeze2'], hl_2))
    hit_labels = ({label for label, _, _ in prev_hit.values()}
                  | {label for label, _, _ in next_hit.values()})
    for label, pf in list(prevs) + list(nexts):
        if label in hit_labels:
            for sheet_name, rows in fac['raw2'](pf):
                raw_sheets.append((f'②{sheet_name}({label})', rows, fac['freeze2'], hl_2))
    raw_sheets.append(('③アラジン原本', CalamineWorkbook.from_path(str(p3)).get_sheet_by_index(0).to_python(), 'A5', hl_aladdin))

    out_path = Path(args.out) if args.out else base / f'検証結果_{fac["label"]}_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
    write_excel(out_path, info, records_a, records_b, raw_sheets)
    print(f'結果Excelを保存しました: {out_path}')
    if info.get('メールHTML'):
        print(f'報告メールHTMLを保存しました: {info["メールHTML"]}')
        if info.get('メール他工場'):
            print(f'  ({info["メール他工場"]})')
    # 最新以外の検証結果を 過去検証結果\ へ退避 (xlsx と対の html)
    moved = _archive_old_results(base, fac['label'], keep=out_path)
    if moved:
        print(f'過去の検証結果 {len(moved)} 件を {ARCHIVE_DIR}\\ へ移動しました')


if __name__ == '__main__':
    main()
