import csv
import io
import sys
from pathlib import Path

from openpyxl import Workbook

SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_DIR / 'scripts'
SKILL_MD = SKILL_DIR / 'SKILL.md'
SCRIPT = SCRIPTS / 'verify_kouchin.py'

sys.path.insert(0, str(SCRIPTS))


def write_toray_csv(path: Path, data_rows, t_amount=None):
    """①CSV (cp932)。data_rows = (入庫場所, 発注No., 入庫月日, H列符号, 金額)"""
    header = ['入庫場所', '品種', '発注No.', '入庫月日', '入庫No.', '単位',
              '契約数量', '品名', '単価', '金額', '消費税等', '支払総額']
    lines = [header]
    for basho, keiyaku, date, sign, amount in data_rows:
        lines.append([basho, '', keiyaku, date, '1', '', '1', sign, '', str(amount), '', ''])
    if t_amount is not None:
        lines.append(['', '', 'T', '', '', '', '', '', '', str(t_amount), '', ''])
    buf = io.StringIO()
    csv.writer(buf).writerows(lines)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buf.getvalue().encode('cp932'))


def make_nagaoka(path: Path, rows):
    """② 東レまとめ。rows = (A列, B列, 契約NO, AA金額)。実ファイル同様 6行目からデータ。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '東レまとめ'
    ws.cell(1, 1, '後加工工賃明細')
    ws.cell(2, 1, '東レまとめ')
    ws.cell(3, 3, '契約No.')
    ws.cell(4, 27, '合計')
    ws.cell(5, 1, '連番')
    for i, (a, b, c, aa) in enumerate(rows, start=6):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
        ws.cell(i, 27, aa)
    for name in ('東レT', '東レV.C', '東レY', '東レW.E'):
        wb.create_sheet(name)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_aladdin(path: Path, ym='2026年08月', items=None, with_customer=False):
    """③ 依頼NO別問合せ。items = (依頼NO, 金額[, 得意先])"""
    wb = Workbook()
    ws = wb.active
    ws['A1'] = f'対象年月 : {ym}'
    headers = ['依頼NO', '項目', '--合計--']
    if with_customer:
        headers.append('得意先')
    for col, h in enumerate(headers, start=1):
        ws.cell(3, col, h)
    for i, item in enumerate(items or [], start=4):
        ws.cell(i, 1, item[0])
        ws.cell(i, 2, '加工金額')
        ws.cell(i, 3, item[1])
        if with_customer:
            ws.cell(i, 4, item[2])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
