import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))

from write_excel import archive_old_trends, write_trend_book


def _minimal_payload():
    months = [(2026, 7), (2026, 8)]
    kokubu = [
        {
            'ym': (2026, 7), 'path': Path('k7.xlsx'),
            'kinds': {
                'スライス1': {'wage': 100, 'qty': 10},
                'スリット': {'wage': 50, 'qty': 5},
            },
            'sheets': {'東レT': {'wage': 150, 'qty': 15}}, 'warnings': [],
        },
        {
            'ym': (2026, 8), 'path': Path('k8.xlsx'),
            'kinds': {
                'スライス1': {'wage': 180, 'qty': 18},
                'スリット': {'wage': 55, 'qty': 5},
            },
            'sheets': {'東レT': {'wage': 235, 'qty': 23}}, 'warnings': [],
        },
    ]
    konan = [
        {
            'ym': (2026, 7), 'path': Path('n7.xlsx'),
            'kinds': {'スライス': {'wage': 200, 'qty': 20}, 'スリット': {'wage': 40, 'qty': 4}},
            'sheets': {'東レT.V.C': {'wage': 240, 'qty': 24}}, 'warnings': [],
        },
        {
            'ym': (2026, 8), 'path': Path('n8.xlsx'),
            'kinds': {'スライス': {'wage': 120, 'qty': 12}, 'スリット': {'wage': 42, 'qty': 4}},
            'sheets': {'東レT.V.C': {'wage': 162, 'qty': 16}}, 'warnings': [],
        },
    ]
    return {
        'months': months,
        'kokubu': kokubu,
        'konan': konan,
        'warnings': ['テスト警告'],
        'files': {
            '国分': {(2026, 7): Path('k7.xlsx'), (2026, 8): Path('k8.xlsx')},
            '湖南': {(2026, 7): Path('n7.xlsx'), (2026, 8): Path('n8.xlsx')},
        },
    }


def test_write_creates_expected_sheets(tmp_path):
    out = tmp_path / '月トレンド_test.xlsx'
    write_trend_book(out, _minimal_payload())
    from openpyxl import load_workbook
    wb = load_workbook(out)
    for name in ['サマリ', '工程比較', '負荷シフト', '付録_国分', '付録_湖南', '付録_シート区分']:
        assert name in wb.sheetnames
    # スパゲッティ全工程グラフは工程比較にのみ（注目工程×量/賃）
    assert wb['工程比較']._charts
    assert not wb['付録_国分']._charts
    # チャートデザイン検証: 系列は国分/湖南の2本のみ、値ラベルのみ、マーカーあり
    chart0 = wb['工程比較']._charts[0]
    assert len(chart0.series) == 2
    ser0 = chart0.series[0]
    assert ser0.smooth is False
    assert getattr(ser0.graphicalProperties.line.solidFill, 'srgbClr', None) == 'C0504D'
    assert ser0.marker is not None and ser0.marker.symbol == 'circle'
    ser1 = chart0.series[1]
    assert getattr(ser1.graphicalProperties.line.solidFill, 'srgbClr', None) == '9BBB59'
    assert chart0.dataLabels is not None and chart0.dataLabels.showVal is True
    assert chart0.dataLabels.showSerName is False
    assert chart0.dataLabels.showCatName is False
    assert chart0.dataLabels.numFmt == '#,##0'
    assert chart0.legend.position == 'r'
    # タイトルとプロットが重ならないよう overlay=False
    assert chart0.title.overlay is False
    chart1 = wb['工程比較']._charts[1]
    assert chart1.dataLabels.numFmt == '#,##0'
    # ヘッダーは「区分」（「工場」を系列名にしない）
    assert wb['工程比較']['A7'].value == '区分'
    # グラフ同士が上下に重ならないこと（アンカー行の間隔が最低16行以上あること）
    chart_rows = [c.anchor._from.row for c in wb['工程比較']._charts]
    for i in range(len(chart_rows) - 1):
        assert chart_rows[i + 1] - chart_rows[i] >= 16, f"Chart overlap: row {chart_rows[i]} vs {chart_rows[i+1]}"
    # 単位 km / 千円 が工程比較に含まれること
    cmp_texts = [
        str(row[0])
        for row in wb['工程比較'].iter_rows(max_row=40, max_col=1, values_only=True)
        if row[0]
    ]
    assert any('km' in t for t in cmp_texts)
    assert any('千円' in t for t in cmp_texts)
    texts = [
        str(row[0])
        for row in wb['サマリ'].iter_rows(max_row=40, max_col=1, values_only=True)
        if row[0]
    ]
    flat = ' '.join(texts)
    assert '工程比較' in flat
    assert 'スライス' in flat or '振替' in flat


def test_archive_old(tmp_path):
    old = tmp_path / '月トレンド_20260101_000000.xlsx'
    old.write_bytes(b'PK')
    new = tmp_path / '月トレンド_20260907_120000.xlsx'
    new.write_bytes(b'PK')
    archive_old_trends(tmp_path, keep=new)
    assert not old.exists()
    assert (tmp_path / '過去月トレンド' / old.name).exists()
