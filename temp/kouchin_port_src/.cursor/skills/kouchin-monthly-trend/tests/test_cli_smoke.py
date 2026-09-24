import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))

from test_read_kokubu import _make_kokubu
from test_read_konan import _make_konan


def test_main_both_factories(tmp_path, monkeypatch):
    base = tmp_path
    d2026 = tmp_path / '工賃明細2026年度'
    d2026.mkdir()
    _make_kokubu(d2026 / '後加工工賃明細（2026年7月度)V01.xlsx')
    _make_kokubu(d2026 / '後加工工賃明細（2026年8月度)V01.xlsx')

    shisan = tmp_path / 'shisan'
    nendo = shisan / '2026年度試算　湖南'
    from openpyxl import load_workbook
    p7 = nendo / '7月度加工賃試算.xlsx'
    _make_konan(p7)
    wb = load_workbook(p7)
    for name in wb.sheetnames:
        wb[name]['A1'] = '2026年7月度'
    wb.save(p7)
    _make_konan(nendo / '8月度加工賃試算.xlsx')

    import common
    monkeypatch.setattr(common, 'KONAN_SHISAN_ROOT', str(shisan))
    monkeypatch.setattr(common, 'KOKUBU_MEISAI_DIRS', (str(d2026),))

    from trend_kouchin import main
    out = tmp_path / 'out.xlsx'
    rc = main(['--base', str(base), '--months', '2', '--out', str(out)])
    assert rc == 0
    assert out.exists()
