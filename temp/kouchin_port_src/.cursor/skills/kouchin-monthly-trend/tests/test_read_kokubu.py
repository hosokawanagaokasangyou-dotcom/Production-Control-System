import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))

from read_kokubu import read_kokubu_month


def _make_kokubu(path: Path):
    """実ファイルに近い: 加工依頼行の次が工程名、その下にデータ。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '東レまとめ'
    ws.cell(3, 1, '加工依頼No.')
    ws.cell(3, 3, '契約No.')
    ws.cell(3, 5, '入庫数量')
    ws.cell(3, 27, '加工種類別 加工賃')
    ws.cell(3, 46, '加工種類別 入庫量')
    # 工程名行
    ws.cell(4, 7, '合計')
    ws.cell(4, 8, 'スリット')
    ws.cell(4, 9, 'カット')
    ws.cell(4, 27, '合計')
    ws.cell(4, 28, 'スリット')
    ws.cell(4, 29, 'カット')
    ws.cell(4, 46, 'スリット')
    ws.cell(4, 47, 'カット')
    # データ
    ws.cell(6, 1, 'T8')
    ws.cell(6, 2, 1)
    ws.cell(6, 3, '191352R')
    ws.cell(6, 5, 15)
    ws.cell(6, 27, 1500)
    ws.cell(6, 28, 1000)
    ws.cell(6, 29, 500)
    ws.cell(6, 46, 10)
    ws.cell(6, 47, 5)
    ws.cell(7, 3, '加工賃合計')
    ws.cell(7, 27, 9999)

    t = wb.create_sheet('東レT')
    t.cell(6, 1, 'T8')
    t.cell(6, 2, 1)
    t.cell(6, 3, '191352R')
    t.cell(6, 5, 15)
    t.cell(6, 27, 1500)
    for name in ('東レV.C', '東レY', '東レW.E'):
        wb.create_sheet(name)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_read_kokubu_kinds_and_sheets(tmp_path):
    path = tmp_path / '後加工工賃明細（2026年8月度)V01.xlsx'
    _make_kokubu(path)
    agg = read_kokubu_month(path)
    assert agg['ym'] == (2026, 8)
    assert agg['kinds']['スリット']['wage'] == 1000
    assert agg['kinds']['スリット']['qty'] == 10
    assert agg['kinds']['カット']['wage'] == 500
    assert agg['kinds']['カット']['qty'] == 5
    assert agg['sheets']['東レT']['wage'] == 1500
    assert agg['sheets']['東レT']['qty'] == 15
