import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))

from read_konan import read_konan_month


def _add_toray_sheet(wb, title, ym_label, slit_wage, slit_qty, inbound):
    if title in wb.sheetnames:
        ws = wb[title]
    else:
        ws = wb.create_sheet(title)
    ws['A1'] = ym_label
    # hdr row 8
    ws.cell(8, 1, '加工依頼No.')
    ws.cell(8, 3, '契約No.')
    ws.cell(8, 4, '入庫数量')
    ws.cell(8, 5, '加工内容')
    ws.cell(8, 22, '加工種類別 加工賃')
    ws.cell(8, 38, '加工種類別 加工量')
    # name row 9
    ws.cell(9, 6, '合計')
    ws.cell(9, 7, 'スリット')
    ws.cell(9, 22, '合計')
    ws.cell(9, 23, 'スリット')
    ws.cell(9, 38, '合計')
    ws.cell(9, 39, 'スリット')
    # unit / total placeholder rows 10-11
    ws.cell(10, 7, 10)
    ws.cell(11, 22, 0)
    # data row 12
    ws.cell(12, 1, 'C8-1')
    ws.cell(12, 3, '191352R')
    ws.cell(12, 4, inbound)
    ws.cell(12, 5, 'スリット')
    ws.cell(12, 22, slit_wage)
    ws.cell(12, 23, slit_wage)
    ws.cell(12, 38, slit_qty)
    ws.cell(12, 39, slit_qty)


def _make_konan(path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    _add_toray_sheet(wb, '東レT.V.C', '2026年8月度', 100, 10, 10)
    _add_toray_sheet(wb, '東レY.S', '2026年8月度', 50, 5, 5)
    _add_toray_sheet(wb, '東レW.E..', '2026年8月度', 0, 0, 0)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_read_konan_kinds_sum_across_sheets(tmp_path):
    path = tmp_path / '8月度加工賃試算.xlsx'
    _make_konan(path)
    agg = read_konan_month(path)
    assert agg['ym'] == (2026, 8)
    assert agg['kinds']['スリット']['wage'] == 150
    assert agg['kinds']['スリット']['qty'] == 15
    assert agg['sheets']['東レT.V.C']['wage'] == 100
    assert agg['sheets']['東レY.S']['wage'] == 50
    assert agg['sheets']['東レT.V.C']['qty'] == 10
