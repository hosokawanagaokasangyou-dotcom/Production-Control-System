from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.text import CharacterProperties
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.fonts import DEFAULT_FONT
from openpyxl.utils import get_column_letter

from aggregate import (
    build_grouped_cross,
    build_kind_series,
    build_sheet_series,
    factory_month_total,
    pick_focus_processes,
    rank_shift_suspects,
)

FONT_NAME = 'BIZ UDPゴシック'
FONT_NUM = 'BIZ UDゴシック'
FONT_SIZE = 10.5
DEFAULT_FONT.name = FONT_NAME
DEFAULT_FONT.size = FONT_SIZE
DEFAULT_FONT.scheme = None

# 工場色・チャート配色（固定）
COLOR_KOKUBU = 'C0504D'       # 赤茶（国分）
COLOR_KONAN = '9BBB59'        # 若草色（湖南）
COLOR_CHART_TITLE = '1F497D'  # 濃紺（グラフタイトル）
COLOR_CHART_BORDER = 'D9E1F2' # 薄水色（外枠・目盛線）

FILL_WARN = PatternFill('solid', fgColor='FCE4D6')
FILL_SHIFT = PatternFill('solid', fgColor='FFF2CC')
FILL_HEADER = PatternFill('solid', fgColor='D6DCE4')
FILL_KPI = PatternFill('solid', fgColor='E2EFDA')
FILL_SECTION = PatternFill('solid', fgColor='D9E1F2')

_TAB = {
    'サマリ': '1F4E79',
    '工程比較': 'C00000',
    '負荷シフト': '7030A0',
    '付録_国分': 'A6A6A6',
    '付録_湖南': 'A6A6A6',
    '付録_シート区分': 'BFBFBF',
}


def _ym_label(ym: tuple[int, int]) -> str:
    return f'{ym[0]}年{ym[1]}月'


def _font(bold=False, num=False, color=None, size=None):
    return Font(
        name=FONT_NUM if num else FONT_NAME,
        size=size or FONT_SIZE,
        bold=bold,
        color=color,
    )


def archive_old_trends(base: Path, keep: Path):
    dest = base / '過去月トレンド'
    dest.mkdir(parents=True, exist_ok=True)
    keep = keep.resolve()
    for f in base.glob('月トレンド_*.xlsx'):
        if f.resolve() == keep:
            continue
        target = dest / f.name
        if target.exists():
            target.unlink()
        shutil.move(str(f), str(target))


def _style_header_row(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = _font(bold=True)
        cell.fill = FILL_HEADER


def _num_cell(cell, v, pct=False, decimals=None):
    cell.value = v
    cell.font = _font(num=True)
    if v is None:
        return
    if pct:
        cell.number_format = '0.0%'
    elif decimals == 1:
        cell.number_format = '#,##0.0'
    else:
        cell.number_format = '#,##0'


def _round_disp(v):
    """表示用の四捨五入（整数）。正負とも 0.5 以上を切り上げ。"""
    if v is None:
        return None
    x = float(v)
    return int(x + 0.5) if x >= 0 else -int(-x + 0.5)


def _to_display_units(months_data: list[dict]) -> None:
    """元データ（量=m・賃=円）を表示単位（km・千円）へ変換し、四捨五入で整数化。"""
    for m in months_data:
        for bucket in (m.get('kinds') or {}, m.get('sheets') or {}):
            for cell in bucket.values():
                if not isinstance(cell, dict):
                    continue
                if cell.get('qty') is not None:
                    cell['qty'] = _round_disp(float(cell['qty']) / 1000.0)
                if cell.get('wage') is not None:
                    cell['wage'] = _round_disp(float(cell['wage']) / 1000.0)


def _add_two_series_chart(ws, title, header_row, data_rows, n_months, anchor, unit='km'):
    """header_row: 月ラベル行, 続く data_rows 行が 国分/湖南（差行は含めない）。
    unit: 'km'（加工量）または '千円'（加工賃）。
    タイトルはプロット領域外（overlay=False）。Y軸上限に余白を取りラベルと重ならないようにする。
    """
    if data_rows < 1 or n_months < 1:
        return
    num_fmt = '#,##0'

    chart = LineChart()
    chart.title = title
    if chart.title and chart.title.text and chart.title.text.rich and chart.title.text.rich.paragraphs:
        p = chart.title.text.rich.paragraphs[0]
        if p.r:
            p.r[0].rPr = CharacterProperties(b=True, sz=1400, solidFill=COLOR_CHART_TITLE)
    chart.title.overlay = False  # タイトルをプロット領域の外に置き、ピークラベルと重ねない

    chart.height = 8.0
    chart.width = 14.5
    chart.legend.position = 'r'
    chart.dispBlanksAs = 'gap'  # 欠損(None)はゼロ埋めせず線を切る

    # 外枠: 角丸・薄い水色
    chart.roundedCorners = True
    chart.graphical_properties = GraphicalProperties()
    chart.graphical_properties.line.solidFill = COLOR_CHART_BORDER
    chart.graphical_properties.line.width = 12000

    # グリッド線: 水平目盛線のみ薄い水色
    if chart.y_axis.majorGridlines is None:
        from openpyxl.chart.axis import ChartLines
        chart.y_axis.majorGridlines = ChartLines()
    chart.y_axis.majorGridlines.spPr = GraphicalProperties()
    chart.y_axis.majorGridlines.spPr.line.solidFill = COLOR_CHART_BORDER

    # 軸線・Y軸単位・目盛書式
    chart.x_axis.spPr = GraphicalProperties()
    chart.x_axis.spPr.line.solidFill = COLOR_CHART_BORDER
    chart.y_axis.title = unit
    chart.y_axis.numFmt = num_fmt

    # データ最大値からY軸上限を少し広げ、上側ラベルの余白を確保
    ymax = None
    for r in range(header_row + 1, header_row + data_rows + 1):
        for c in range(2, 2 + n_months):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                ymax = v if ymax is None else max(ymax, v)
    if ymax is not None and ymax > 0:
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = ymax * 1.25

    # ヘッダー行（区分+月）は系列に含めない。国分/湖南のみ。
    data = Reference(
        ws,
        min_col=1,
        min_row=header_row + 1,
        max_col=1 + n_months,
        max_row=header_row + data_rows,
    )
    cats = Reference(ws, min_col=2, min_row=header_row, max_col=1 + n_months)
    chart.add_data(data, from_rows=True, titles_from_data=True)
    chart.set_categories(cats)

    def _value_only_labels():
        dl = DataLabelList()
        dl.showVal = True
        dl.showCatName = False
        dl.showSerName = False
        dl.showLegendKey = False
        dl.showPercent = False
        dl.dLblPos = 't'
        dl.numFmt = num_fmt
        return dl

    colors = [COLOR_KOKUBU, COLOR_KONAN]
    for i, ser in enumerate(chart.series):
        ser.smooth = False
        color = colors[i] if i < len(colors) else '000000'
        ser.graphicalProperties.line.solidFill = color
        ser.graphicalProperties.line.width = 38000
        ser.marker = Marker(symbol='circle', size=7)
        ser.marker.graphicalProperties.solidFill = color
        ser.marker.graphicalProperties.line.solidFill = color
        ser.dLbls = _value_only_labels()

    chart.dataLabels = _value_only_labels()
    ws.add_chart(chart, anchor)


def _last_comparable_pair(months, kokubu, konan):
    """両工場に kinds がある直近2か月。"""
    kb = {m['ym'] for m in kokubu if m.get('ym') and m.get('kinds')}
    kn = {m['ym'] for m in konan if m.get('ym') and m.get('kinds')}
    both = [ym for ym in months if ym in kb and ym in kn]
    if len(both) < 2:
        return None, None
    return both[-2], both[-1]


def _write_home(ws, payload, gw, gq, focus, shift_rows):
    months = payload['months']
    kokubu = payload.get('kokubu') or []
    konan = payload.get('konan') or []
    warnings = payload.get('warnings') or []
    files = payload.get('files') or {}

    ws['A1'] = '後加工工賃 月次トレンド — ホーム'
    ws['A1'].font = _font(bold=True, size=14)
    if months:
        ws['A2'] = f'対象期間: {_ym_label(months[0])} 〜 {_ym_label(months[-1])}（{len(months)}か月）'
    ws['A3'] = 'まず「工程比較」で注目工程の国分↔湖南を確認。振替は断定せず「疑い」として見ます。'
    ws['A3'].font = _font(color='666666')

    # 警告（先頭・短く）
    ws['A5'] = '【警告】'
    ws['A5'].font = _font(bold=True, color='C00000')
    if warnings:
        for i, w in enumerate(warnings[:8]):
            cell = ws.cell(6 + i, 1, w)
            cell.fill = FILL_WARN
        r = 6 + min(len(warnings), 8)
    else:
        ws['A6'] = '（なし）'
        r = 8

    prev, latest = _last_comparable_pair(months, kokubu, konan)
    r += 1
    ws.cell(r, 1, '【30秒サマリ】').font = _font(bold=True, size=12)
    r += 1
    if latest and prev:
        ws.cell(r, 1, f'比較月: {_ym_label(prev)} → {_ym_label(latest)}（両工場データあり）')
        r += 1
        # KPI
        headers = ['指標', '国分', '湖南', '両工場合計']
        for j, h in enumerate(headers, 1):
            ws.cell(r, j, h)
        _style_header_row(ws, r, 4)
        r += 1
        for metric, label, unit in (('wage', '加工賃', '千円'), ('qty', '加工量', 'km')):
            tk0 = factory_month_total(kokubu, prev, metric)
            tk1 = factory_month_total(kokubu, latest, metric)
            th0 = factory_month_total(konan, prev, metric)
            th1 = factory_month_total(konan, latest, metric)
            ws.cell(r, 1, f'{label}（{_ym_label(latest)}）[{unit}]')
            _num_cell(ws.cell(r, 2), _round_disp(tk1), decimals=0)
            _num_cell(ws.cell(r, 3), _round_disp(th1), decimals=0)
            if tk1 is not None and th1 is not None:
                _num_cell(ws.cell(r, 4), _round_disp(tk1 + th1), decimals=0)
            for c in range(1, 5):
                ws.cell(r, c).fill = FILL_KPI
            r += 1
            # 前月差
            ws.cell(r, 1, f'{label} 前月差[{unit}]')
            dk = (tk1 - tk0) if None not in (tk0, tk1) else None
            dh = (th1 - th0) if None not in (th0, th1) else None
            _num_cell(ws.cell(r, 2), _round_disp(dk), decimals=0)
            _num_cell(ws.cell(r, 3), _round_disp(dh), decimals=0)
            if dk is not None and dh is not None:
                _num_cell(ws.cell(r, 4), _round_disp(dk + dh), decimals=0)
            r += 1
    else:
        ws.cell(r, 1, '両工場で比較できる連続2か月が不足しています')
        r += 1

    r += 1
    shift_unit = 'km' if shift_rows and shift_rows[0].get('metric') == 'qty' else '千円'
    shift_dec = 0
    ws.cell(r, 1, f'【振替・シフト疑い Top】（{shift_unit}ベース前月差。異符号の相殺）').font = _font(bold=True)
    r += 1
    for j, h in enumerate(['工程', '判定', f'国分前月差({shift_unit})', f'湖南前月差({shift_unit})', f'両工場合計差({shift_unit})', f'相殺量({shift_unit})'], 1):
        ws.cell(r, j, h)
    _style_header_row(ws, r, 6)
    r += 1
    shown = [x for x in shift_rows if x['badge'] in ('強い相殺', '弱い相殺')][:5]
    if not shown:
        ws.cell(r, 1, '（該当なし）')
        r += 1
    for row in shown:
        ws.cell(r, 1, row['process'])
        ws.cell(r, 2, row['badge'])
        _num_cell(ws.cell(r, 3), row['d_kokubu'], decimals=shift_dec)
        _num_cell(ws.cell(r, 4), row['d_konan'], decimals=shift_dec)
        _num_cell(ws.cell(r, 5), row['d_combined'], decimals=shift_dec)
        _num_cell(ws.cell(r, 6), row['offset'], decimals=shift_dec)
        for c in range(1, 7):
            ws.cell(r, c).fill = FILL_SHIFT
        r += 1

    r += 1
    ws.cell(r, 1, '【注目工程】→ 次タブ「工程比較」').font = _font(bold=True)
    r += 1
    ws.cell(r, 1, '、'.join(focus) if focus else '（なし）')
    r += 2
    ws.cell(r, 1, '【ナビ】').font = _font(bold=True)
    r += 1
    ws.cell(r, 1, '1. 工程比較 … 注目工程ごとに国分(赤茶)と湖南(若草)を同じグラフで比較')
    r += 1
    ws.cell(r, 1, '2. 負荷シフト … 動いた工程の一覧（前月差・判定）')
    r += 1
    ws.cell(r, 1, '3. 付録_* … 工場別の全工程表（通常は不要）')
    r += 2
    ws.cell(r, 1, '【使用ファイル】').font = _font(bold=True)
    r += 1
    for factory, fmap in files.items():
        for ym, path in sorted(fmap.items()):
            ws.cell(r, 1, f'{factory} {_ym_label(ym)}: {Path(path).name}')
            r += 1
    ws.column_dimensions['A'].width = 42
    for col in 'BCDEF':
        ws.column_dimensions[col].width = 14


def _write_process_compare(ws, months, gw, gq, focus):
    """注目工程ごとに 国分/湖南 2系列グラフ（加工量優先＋加工賃表）。
    グラフが重ならないよう、1グラフあたり約14行の占有を考慮して十分な行間を確保する。
    """
    ws['A1'] = '工程比較（グループ合算・国分↔湖南）'
    ws['A1'].font = _font(bold=True, size=14)
    ws['A2'] = (
        '国分の「スライス1/3」はグループ「スライス」に合算。色は国分=赤茶(#C0504D)・湖南=若草(#9BBB59)で固定。'
        'グラフ系列は国分・湖南のみ（差は表参照）。ラベルは値のみ。欠損月は空欄（ゼロ埋めしない）。'
    )
    ws['A2'].font = _font(color='666666')
    ws['A3'] = f'注目工程: {", ".join(focus)}'
    ws.column_dimensions['A'].width = 16

    row = 5
    n = len(months)
    chart_col = get_column_letter(n + 3)
    fill_diff = PatternFill('solid', fgColor='F2F2F2')

    for proc in focus:
        # 工程見出し帯（薄い青色背景で視覚的ブロックを明確化）
        ws.cell(row, 1, f'■ {proc}').font = _font(bold=True, size=12)
        for c in range(1, 1 + n + 1):
            ws.cell(row, c).fill = FILL_SECTION
        row += 1

        # --- 1. 加工量ブロック ---
        ws.cell(row, 1, '【加工量】（単位: km）').font = _font(bold=True)
        header_q = row + 1
        ws.cell(header_q, 1, '区分')
        for j, ym in enumerate(months, start=2):
            ws.cell(header_q, j, _ym_label(ym))
        _style_header_row(ws, header_q, 1 + n)
        ws.cell(header_q + 1, 1, '国分')
        ws.cell(header_q + 2, 1, '湖南')
        for j, ym in enumerate(months, start=2):
            _num_cell(ws.cell(header_q + 1, j), gq['values']['国分'].get(proc, {}).get(ym), decimals=0)
            _num_cell(ws.cell(header_q + 2, j), gq['values']['湖南'].get(proc, {}).get(ym), decimals=0)
        # 差分行（表専用・グラフ対象外）
        ws.cell(header_q + 3, 1, '差(国分-湖南)').font = _font(color='808080')
        for j, ym in enumerate(months, start=2):
            k = gq['values']['国分'].get(proc, {}).get(ym)
            h = gq['values']['湖南'].get(proc, {}).get(ym)
            diff = (k - h) if k is not None and h is not None else None
            cell = ws.cell(header_q + 3, j)
            _num_cell(cell, _round_disp(diff), decimals=0)
            cell.fill = fill_diff
            cell.font = _font(num=True, color='808080')
        ws.cell(header_q + 3, 1).fill = fill_diff

        # 加工量グラフ（系列は国分・湖南の2行のみ）
        _add_two_series_chart(
            ws, f'{proc}・加工量 (km)', header_q, 2, n,
            f'{chart_col}{header_q}',
            unit='km',
        )

        # --- 2. 加工賃ブロック ---
        wage_start = header_q + 16
        ws.cell(wage_start, 1, '【加工賃】（単位: 千円）').font = _font(bold=True)
        wage_header = wage_start + 1
        ws.cell(wage_header, 1, '区分')
        for j, ym in enumerate(months, start=2):
            ws.cell(wage_header, j, _ym_label(ym))
        _style_header_row(ws, wage_header, 1 + n)
        ws.cell(wage_header + 1, 1, '国分')
        ws.cell(wage_header + 2, 1, '湖南')
        for j, ym in enumerate(months, start=2):
            _num_cell(ws.cell(wage_header + 1, j), gw['values']['国分'].get(proc, {}).get(ym), decimals=0)
            _num_cell(ws.cell(wage_header + 2, j), gw['values']['湖南'].get(proc, {}).get(ym), decimals=0)
        # 差分行（表専用・グラフ対象外）
        ws.cell(wage_header + 3, 1, '差(国分-湖南)').font = _font(color='808080')
        for j, ym in enumerate(months, start=2):
            k = gw['values']['国分'].get(proc, {}).get(ym)
            h = gw['values']['湖南'].get(proc, {}).get(ym)
            diff = (k - h) if k is not None and h is not None else None
            cell = ws.cell(wage_header + 3, j)
            _num_cell(cell, _round_disp(diff), decimals=0)
            cell.fill = fill_diff
            cell.font = _font(num=True, color='808080')
        ws.cell(wage_header + 3, 1).fill = fill_diff

        _add_two_series_chart(
            ws, f'{proc}・加工賃 (千円)', wage_header, 2, n,
            f'{chart_col}{wage_header}',
            unit='千円',
        )

        row = wage_header + 17


def _write_load_shift(ws, shift_qty, shift_wage, prev, latest):
    ws['A1'] = '負荷シフト（振替・シフト疑い一覧）'
    ws['A1'].font = _font(bold=True, size=14)
    ws['A2'] = (
        '判定は仮説です。「強い相殺」= 国分と湖南の前月差が異符号で大きさが近い。'
        '両工場合計差が大きければ需要変動の可能性が高いです。'
    )
    ws['A2'].font = _font(color='666666')
    if prev and latest:
        ws['A3'] = f'比較: {_ym_label(prev)} → {_ym_label(latest)}'
    row = 5
    for title, rows, dec in (('【加工量ベース】（単位: km）', shift_qty, 0), ('【加工賃ベース】（単位: 千円）', shift_wage, 0)):
        ws.cell(row, 1, title).font = _font(bold=True)
        row += 1
        headers = [
            '工程', '判定', '国分(前)', '国分(当)', '国分差',
            '湖南(前)', '湖南(当)', '湖南差', '合計差', '相殺量',
        ]
        for j, h in enumerate(headers, 1):
            ws.cell(row, j, h)
        _style_header_row(ws, row, len(headers))
        row += 1
        for rec in rows:
            ws.cell(row, 1, rec['process'])
            ws.cell(row, 2, rec['badge'])
            vals = [
                rec['kokubu_prev'], rec['kokubu_latest'], rec['d_kokubu'],
                rec['konan_prev'], rec['konan_latest'], rec['d_konan'],
                rec['d_combined'], rec['offset'],
            ]
            for j, v in enumerate(vals, 3):
                _num_cell(ws.cell(row, j), v, decimals=dec)
            if rec['badge'] in ('強い相殺', '弱い相殺'):
                for c in range(1, 11):
                    ws.cell(row, c).fill = FILL_SHIFT
            row += 1
        row += 2
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 12
    for col in range(3, 11):
        ws.column_dimensions[get_column_letter(col)].width = 12


def _write_appendix_factory(ws, title, months_data, months):
    """全工程表。グラフなし（スパゲッティ禁止）。"""
    wage = build_kind_series(months_data, months, 'wage')
    qty = build_kind_series(months_data, months, 'qty')
    procs = sorted(set(wage['processes']) | set(qty['processes']))
    ws['A1'] = title + '（付録・全工程・グラフなし）'
    ws['A1'].font = _font(bold=True)
    ws['A2'] = '主分析は「工程比較」「負荷シフト」を使ってください。'
    ws.cell(4, 1, '加工賃（単位: 千円）')
    ws.cell(5, 1, '工程')
    for j, ym in enumerate(months, start=2):
        ws.cell(5, j, _ym_label(ym))
    _style_header_row(ws, 5, 1 + len(months))
    for i, p in enumerate(procs, start=6):
        ws.cell(i, 1, p)
        for j, ym in enumerate(months, start=2):
            _num_cell(ws.cell(i, j), wage['values'].get(p, {}).get(ym), decimals=0)
    start = 6 + len(procs) + 2
    ws.cell(start, 1, '加工量（単位: km）').font = _font(bold=True)
    ws.cell(start + 1, 1, '工程')
    for j, ym in enumerate(months, start=2):
        ws.cell(start + 1, j, _ym_label(ym))
    _style_header_row(ws, start + 1, 1 + len(months))
    for i, p in enumerate(procs, start=start + 2):
        ws.cell(i, 1, p)
        for j, ym in enumerate(months, start=2):
            _num_cell(ws.cell(i, j), qty['values'].get(p, {}).get(ym), decimals=0)
    ws.column_dimensions['A'].width = 18


def _write_appendix_sheets(ws, months, kokubu, konan):
    ws['A1'] = '付録: シート区分（工場並記）'
    ws['A1'].font = _font(bold=True)
    row = 3
    for label, data in (('国分', kokubu), ('湖南', konan)):
        wage = build_sheet_series(data, months, 'wage')
        ws.cell(row, 1, label).font = _font(bold=True)
        row += 1
        ws.cell(row, 1, 'シート')
        for j, ym in enumerate(months, start=2):
            ws.cell(row, j, _ym_label(ym))
        _style_header_row(ws, row, 1 + len(months))
        row += 1
        for name in wage['sheet_names']:
            ws.cell(row, 1, name)
            for j, ym in enumerate(months, start=2):
                _num_cell(ws.cell(row, j), wage['values'].get(name, {}).get(ym), decimals=0)
            row += 1
        row += 2
    ws.column_dimensions['A'].width = 18


def write_trend_book(out: Path, payload: dict):
    months = payload['months']
    kokubu = payload.get('kokubu') or []
    konan = payload.get('konan') or []
    # 量: m→km、賃: 円→千円（四捨五入で整数）
    _to_display_units(kokubu)
    _to_display_units(konan)

    gw = build_grouped_cross(kokubu, konan, months, 'wage')
    gq = build_grouped_cross(kokubu, konan, months, 'qty')
    focus = pick_focus_processes(gw, gq, max_n=6)
    shift_qty = rank_shift_suspects(gq, 'qty')
    shift_wage = rank_shift_suspects(gw, 'wage')
    # ホーム用は量優先の疑い
    shift_home = shift_qty if shift_qty else shift_wage
    prev, latest = _last_comparable_pair(months, kokubu, konan)

    wb = Workbook()
    home = wb.active
    home.title = 'サマリ'
    home.sheet_properties.tabColor = _TAB['サマリ']
    _write_home(home, payload, gw, gq, focus, shift_home)

    cmp_ws = wb.create_sheet('工程比較')
    cmp_ws.sheet_properties.tabColor = _TAB['工程比較']
    _write_process_compare(cmp_ws, months, gw, gq, focus)

    shift_ws = wb.create_sheet('負荷シフト')
    shift_ws.sheet_properties.tabColor = _TAB['負荷シフト']
    _write_load_shift(shift_ws, shift_qty, shift_wage, prev, latest)

    app_k = wb.create_sheet('付録_国分')
    app_k.sheet_properties.tabColor = _TAB['付録_国分']
    _write_appendix_factory(app_k, '国分 種類別', kokubu, months)

    app_n = wb.create_sheet('付録_湖南')
    app_n.sheet_properties.tabColor = _TAB['付録_湖南']
    _write_appendix_factory(app_n, '湖南 種類別', konan, months)

    app_s = wb.create_sheet('付録_シート区分')
    app_s.sheet_properties.tabColor = _TAB['付録_シート区分']
    _write_appendix_sheets(app_s, months, kokubu, konan)

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
